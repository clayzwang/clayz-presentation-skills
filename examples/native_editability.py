"""Create before/after native-editing samples using the actual public APIs.

This is a future-tool sample, not a revision of any production presentation.
No GUI behavior is inferred from the programmatic node refresh.
"""
import argparse
import copy
import hashlib
import json
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.util import Inches
from packages.layout import add_native_textbox, add_text_shape, write_native_text
from packages.layout import add_rectangular_connector, refresh_rectangular_connector
from packages.layout.native_page import add_native_page
from packages.layout.semantic_table import Column, plan, write


def ref(path):
    return {'path':str(path.resolve()),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size}


def build(directory, font_file):
    directory.mkdir(parents=True, exist_ok=False)
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    s=add_native_page(prs,layout=prs.slide_layouts[6],background='F7F5EF')
    add_native_textbox(s,'原生编辑样例：自然换行与带字节点',.7,.35,12,.65,name='TITLE',size_pt=28,bold=True)
    original='客户将交易、融资与风险管理交给同一服务平台。连续句子由文本框自动换行；扩大宽度后，原生软件重新安排视觉行，不依赖预先插入的回车。'
    body=add_native_textbox(s,original,.7,1.35,4.2,2.75,name='PROSE',size_pt=21,line_spacing=1.15,space_after_pt=8)
    a=add_text_shape(s,'客户工作',.8,4.5,3,1,name='NODE-A',fill='DFECE8',size_pt=21)
    b=add_text_shape(s,'回报与风险',8,4.5,3,1,name='NODE-B',fill='EFE4D2',size_pt=21)
    edge=add_rectangular_connector(s,begin=(a,'right'),end=(b,'left'),name='A-TO-B',color='317A7A',tail={'type':'triangle'})
    a._element.addprevious(edge._element)  # Put the line below its existing nodes.
    add_native_textbox(s,'工具样例；程序改动与refresh不代表应用GUI自动跟随已验证。',.7,6.95,12,.35,name='NOTE',size_pt=12,color='65727A',space_after_pt=0)
    t=add_native_page(prs,layout=prs.slide_layouts[6],background='F7F5EF')
    add_native_textbox(t,'表头与正文有不同原生锚定',.7,.35,12,.65,name='TITLE',size_pt=28,bold=True)
    ts=plan([Column('客户工作',weight=1),Column('回报与条件',weight=2)],
            [['交易与融资服务','服务可收费；实际占用资本时，还需承担资金成本与风险。'],['基金管理','管理资产属于产品；管理服务收费与产品资产本身的投资回报分别理解。']],
            width=11.8,font_path=font_file,size=22,header_size=20,padding=(.18,.12),caption='工具合成内容，非公司事实',comparison='同行读工作与回报条件')
    write(t,ts,.7,1.55,name='NATIVE-TABLE')
    add_native_textbox(t,'测量lines保留为估计；原生cell只写逻辑正文，wrap与倍数行距由原生属性承载。',.7,6.2,12,.8,name='NOTE',size_pt=18,space_after_pt=0)
    before=directory/'before.pptx';prs.save(before)
    before_end=[edge.end_x,edge.end_y]
    # Edit actual content and geometry without adding visual-line breaks.
    longer=original+'\n增加这一真实段落，用段后间距分隔。保留倍数行距，同时把文本框拉宽，字体和单词仍在原对象内。'
    write_native_text(body,longer,size_pt=21,line_spacing=1.15,space_after_pt=8)
    body.width=Inches(6.0);body.height=Inches(3.05)
    write_native_text(a,'更长的客户工作标签',role='node',size_pt=21)
    a.width=Inches(3.8)
    b.left=Inches(8.3);b.top=Inches(4.9);b.width=Inches(3.2)
    before_refresh=[edge.end_x,edge.end_y]
    assert before_refresh==before_end  # Library property edits are not observers.
    refresh_rectangular_connector(edge)
    after_refresh=[edge.end_x,edge.end_y]
    assert edge.end_x==b.left and edge.end_y==b.top+b.height//2
    # Copy one native object that already owns both text and geometry.
    copied=copy.deepcopy(a._element)
    identity=copied.xpath('.//p:cNvPr')[0]
    identity.set('id',str(max(sh.shape_id for sh in s.shapes)+1));identity.set('name','NODE-A-COPY')
    s.shapes._spTree.insert_element_before(copied,'p:extLst')
    clone=next(sh for sh in s.shapes if sh.name=='NODE-A-COPY');clone.left=Inches(.8);clone.top=Inches(5.85);clone.height=Inches(.85)
    after=directory/'after.pptx';prs.save(after)
    snapshots={}
    for name,path in [('before',before),('after',after)]:
        folder=directory/(name+'-xml');folder.mkdir()
        with zipfile.ZipFile(path)as z:
            for part in ['ppt/slides/slide1.xml','ppt/slides/slide2.xml']:
                (folder/Path(part).name).write_bytes(z.read(part))
        reopened=Presentation(path)
        snapshots[name]={'pptx':ref(path),'objects_per_slide':[len(sl.shapes)for sl in reopened.slides],
            'shapes':[{'name':sh.name,'id':sh.shape_id,'text':sh.text if sh.has_text_frame else None,
                       'box_emu':[sh.left,sh.top,sh.width,sh.height]}for sh in reopened.slides[0].shapes]}
    evidence={'scope':'Actual API generation and explicit programmatic edits; not production costs or GUI proof.',
        'before':snapshots['before'],'after':snapshots['after'],
        'node_edit':{'before_endpoint_emu':before_end,'after_node_move_before_refresh_emu':before_refresh,'after_explicit_refresh_emu':after_refresh,
                     'native_bindings_retained':edge._element.xpath('.//a:stCxn/@id')+edge._element.xpath('.//a:endCxn/@id'),
                     'native_end_arrow':edge._element.xpath('.//a:tailEnd/@type'),
                     'copy_is_one_native_object_with_own_text':clone.text==a.text},
        'sample_native_saves':2,'sample_render_attempts_so_far':0,
        'production_R08_costs_added':0,
        'unverified':['PowerPoint/WPS/LibreOffice GUI dragging and automatic endpoint tracking. No installed UNO document-model route used.'],
        'table_planned_estimates':ts}
    (directory/'native-edit-evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'before':ref(before),'after':ref(after),'evidence':ref(directory/'native-edit-evidence.json')},ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--directory',type=Path,required=True);parser.add_argument('--font-file',type=Path,required=True)
    args=parser.parse_args();build(args.directory,args.font_file)
