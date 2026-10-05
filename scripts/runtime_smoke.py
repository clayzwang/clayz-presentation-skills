#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Real LibreOffice smoke: metadata-stamped tables, multiline text and negative chart labels."""
import argparse
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import hashlib
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt
from stamp_pptx_metadata import stamp


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    office=shutil.which('libreoffice') or shutil.which('soffice')
    if not office or not shutil.which('pdftotext'):
        raise RuntimeError('real smoke requires LibreOffice Impress and Poppler pdftotext')
    with tempfile.TemporaryDirectory(prefix='clayz-smoke-') as td:
        root=Path(td);p=Presentation()
        s=p.slides.add_slide(p.slide_layouts[6]);t=s.shapes.add_table(2,2,0,0,Inches(8),Inches(3))
        for i,text in enumerate(['2025','Line one\nLine two','2025','Complete table']):t.table.cell(i//2,i%2).text=text
        s=p.slides.add_slide(p.slide_layouts[6]);d=CategoryChartData();d.categories=['Negative','Positive'];d.add_series('Values',[-11.09,24.01])
        chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,0,0,Inches(8),Inches(5),d).chart
        chart.plots[0].has_data_labels=True
        chart.plots[0].data_labels.show_value=True
        chart.plots[0].data_labels.number_format='0.00'
        s=p.slides.add_slide(p.slide_layouts[6])
        shared=s.shapes.add_textbox(Inches(.6),Inches(.7),Inches(8),Inches(2))
        shared.name='ART::shared-editorial-text'
        shared.text='Editorial heading\nTwo approved paragraphs share one editable text box.'
        for paragraph,size in zip(shared.text_frame.paragraphs,(26,18)):
            paragraph.font.name='Arial';paragraph.font.size=Pt(size)
        peer=s.shapes.add_textbox(Inches(.6),Inches(3),Inches(8),Inches(.7))
        peer.name='ART::peer-heading';peer.text='A peer heading with different emphasis'
        peer.text_frame.paragraphs[0].font.name='Arial';peer.text_frame.paragraphs[0].font.size=Pt(22)
        src=root/'source.pptx';pptx=root/'stamped.pptx';p.save(src);stamp(src,pptx,{'ClayzVersion':'runtime-smoke'})
        reopened=Presentation(pptx)
        shared=next(shape for shape in reopened.slides[2].shapes if shape.name=='ART::shared-editorial-text')
        if len(shared.text_frame.paragraphs)!=2 or shared.text_frame.paragraphs[0].font.size==shared.text_frame.paragraphs[1].font.size:
            raise RuntimeError('combined native paragraphs or independent styles were lost')
        cmd=[office,'-env:UserInstallation='+ (root/'office-profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(root),str(pptx)]
        proc=subprocess.run(cmd,check=True,capture_output=True,text=True,timeout=90)
        pdf=root/'stamped.pdf'
        if not pdf.is_file():raise RuntimeError('Office returned without an actual PDF: '+proc.stdout+proc.stderr)
        text=subprocess.run(['pdftotext',str(pdf),'-'],check=True,capture_output=True,text=True,timeout=30).stdout
        for value in ['Line one','Line two','Complete table','-11.09','24.01',
                      'Editorial heading','Two approved paragraphs share one editable text box.',
                      'A peer heading with different emphasis']:
            if value not in text:raise RuntimeError('rendered text/negative label missing: '+value)
        result={'status':'passed','office_version':subprocess.run([office,'--version'],capture_output=True,text=True,check=True).stdout.strip(),
                'pptx_sha256':hashlib.sha256(pptx.read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),
                'checks':['metadata-stamped package opens','integrated table and multiline text render','negative chart label retains sign',
                          'combined editable paragraphs and independent text styles survive reopening and rendering'],
                'limitations':['Not a pixel geometry verdict.','Not STKaiti or PowerPoint/WPS native acceptance.']}
        args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(result,indent=2)+'\n')
        print(json.dumps(result))

if __name__=='__main__':main()
