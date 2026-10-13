"""Small authoring operations; no wrapping engine or quality validator.

Art chooses the changed objects and candidate lines. Estimates remain estimates;
unchanged objects retain their authored lines. Native reopen checks remain due.
"""
from copy import deepcopy
import math
import re


def _number(value, label, *, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(label + ' must be a finite number')
    if value < 0 or positive and value == 0:
        raise ValueError(label + ' is outside its domain')
    return value


def configured_size(config, role, requested_pt=None):
    """Resolve an actual size or reject an explicit request below task policy."""
    typography = config['theme']['typography']
    key = {'title': 'title_minimum_pt', 'body': 'body_minimum_pt',
           'annotation': 'footnote_minimum_pt', 'page-number': 'page_number_minimum_pt'}[role]
    minimum = _number(typography[key], key, positive=True)
    if role != 'page-number':
        minimum = max(minimum, _number(typography['minimum_audience_text_pt'],
                                      'minimum_audience_text_pt', positive=True))
    value = minimum if requested_pt is None else _number(requested_pt, 'requested_pt', positive=True)
    if value < minimum:
        raise ValueError(f'{role} {value}pt is below resolved task minimum {minimum}pt')
    return value


def set_native_margins(target, margins_inch):
    """Set a python-pptx TextFrame OR _Cell on its effective native properties.

    Passing cell.text_frame changes bodyPr, not the cell's tcPr margins. Pass
    the cell itself for table padding. No font, alignment or anchor is changed.
    """
    from pptx.util import Inches
    if len(margins_inch) != 4:
        raise ValueError('margins must be left, top, right, bottom')
    values = [_number(v, 'margin') for v in margins_inch]
    # Cell has _tc; TextFrame has _txBody. A Shape must not silently stand in.
    if not (hasattr(target, '_tc') or hasattr(target, '_txBody')):
        raise TypeError('Pass the native cell or text frame, not its shape')
    for side, value in zip(('left', 'top', 'right', 'bottom'), values):
        setattr(target, 'margin_' + side, Inches(value))
    return {side: getattr(target, 'margin_' + side).inches
            for side in ('left', 'top', 'right', 'bottom')}


def revise_lines(specs, replacements):
    """Return a copy changing only (slide_id, element_id) -> authored lines.

    No global reflow, dictionary or numeral rule is inferred. Whitespace-only
    line changes are supported; substantive wording belongs in art_content.
    """
    result = deepcopy(specs)
    index = {}
    for page in result:
        for element in page['elements']:
            key = (page['slide_id'], element['element_id'])
            if key in index:
                raise ValueError('Ambiguous element identity: ' + str(key))
            index[key] = element
    for key, lines in replacements.items():
        if key not in index or index[key].get('kind') != 'text':
            raise ValueError('Unknown or non-text target: ' + str(key))
        if not isinstance(lines, list) or not lines or any(not isinstance(s, str) or not s.strip() for s in lines):
            raise ValueError('Candidate lines must be nonempty strings')
        compact = lambda s: re.sub(r'\s+', '', s)
        old = index[key].get('text', ''.join(index[key]['lines']))
        if compact(''.join(lines)) != compact(old):
            raise ValueError('Line-only revision changed visible wording: ' + str(key))
        index[key]['lines'] = list(lines)
    return result


def line_capacity(widths_pt, usable_width_pt, reserve_ratio, *, evidence_kind):
    """Describe candidate line widths before mutation; never claim ink/QA pass.

    widths_pt must come from the candidate's actual font/run styles. The caller
    names whether this is an estimate or a native metric observation. A manual
    line is assessed like any other line, including a protected numeral range.
    """
    if evidence_kind not in {'estimate', 'native-font-metric'}:
        raise ValueError('Explicit measurement provenance required')
    usable = _number(usable_width_pt, 'usable_width_pt', positive=True)
    reserve = _number(reserve_ratio, 'reserve_ratio')
    if reserve >= 1:
        raise ValueError('reserve_ratio must be below 1')
    widths = [_number(w, 'line width') for w in widths_pt]
    if not widths:
        raise ValueError('At least one measured candidate line is required')
    return {'evidence_kind': evidence_kind, 'usable_width_pt': usable,
            'configured_reserve_ratio': reserve,
            'line_total_reserves': [1 - w / usable for w in widths],
            'lines_below_requested_reserve': [i for i, w in enumerate(widths) if w > usable * (1-reserve)],
            'lines_exceeding_usable_width': [i for i, w in enumerate(widths) if w > usable],
            'quality_status': 'unverified',
            'limitations': ['Horizontal capacity only; check height, padding and visible ink separately.',
                           'Numeric continuity does not establish fit or readable hierarchy.',
                           'Estimate cannot replace reopened native evidence.']}


def set_paragraph_edge_space(target, left_pt, right_pt):
    """Author explicit extra space inside effective native margins.

    Pass Cell or TextFrame. Paragraph marL/marR are absolute, not cumulative;
    existing indentation/bullets require author review. This consumes space,
    never changes the bounding box, font, body/cell padding or QA tolerance.
    Select this only after deciding the paragraph's indentation and alignment.
    """
    from pptx.util import Pt
    left=_number(left_pt,'left_pt');right=_number(right_pt,'right_pt')
    if not (hasattr(target,'_tc') or hasattr(target,'_txBody')):
        raise TypeError('Pass the native cell or text frame')
    tf=target.text_frame if hasattr(target,'_tc') else target
    for paragraph in tf.paragraphs:
        ppr=paragraph._p.get_or_add_pPr()
        ppr.set('marL',str(Pt(left)));ppr.set('marR',str(Pt(right)))
    return {'left_pt':left,'right_pt':right,'paragraph_count':len(tf.paragraphs)}


def native_capacity(target, width_pt, height_pt, line_widths_pt, content_height_pt,
                    reserve_ratio, *, evidence_kind):
    """Read effective properties before writing or after native reopening.

    The caller supplies cell/shape dimensions and style-specific metrics in pt.
    Reports both axes and consumes paragraph edge space; no pass inference.
    Multiple paragraphs conservatively use the largest left/right additions.
    Absolute glyph position/ink and visible relations still need native review.
    """
    if not (hasattr(target,'_tc') or hasattr(target,'_txBody')):
        raise TypeError('Pass the native cell or text frame')
    tf=target.text_frame if hasattr(target,'_tc') else target
    margins={s:getattr(target,'margin_'+s).pt for s in ('left','top','right','bottom')}
    def edge(p,side):
        ppr=p._p.pPr
        return int(ppr.get(side,'0'))/12700 if ppr is not None else 0
    extras={s:max([edge(p,tag) for p in tf.paragraphs],default=0) for s,tag in [('left','marL'),('right','marR')]}
    width=_number(width_pt,'width_pt',positive=True)-margins['left']-margins['right']-extras['left']-extras['right']
    height=_number(height_pt,'height_pt',positive=True)-margins['top']-margins['bottom']
    result=line_capacity(line_widths_pt,width,reserve_ratio,evidence_kind=evidence_kind)
    _number(height,'usable_height_pt',positive=True)
    content_height=_number(content_height_pt,'content_height_pt')
    result.update(effective_margins_pt=margins,paragraph_edge_space_pt=extras,
                  usable_height_pt=height,content_height_pt=content_height,
                  vertical_total_reserve=1-content_height/height,
                  height_below_requested_reserve=content_height>height*(1-reserve_ratio),
                  height_exceeds_usable=content_height>height)
    result['limitations']=['Caller dimensions and font/run metrics must describe this object.',
       'Paragraph edges consume horizontal capacity; negative or hanging indentation needs separate review.',
       'Font metric capacity is not glyph ink, absolute inner-edge containment, or audience quality.',
       'Reopen native properties and inspect actual rendered glyphs before approval.']
    return result
