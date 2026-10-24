"""Check every editable PowerPoint shape against the slide canvas.

Unlike PDF text bounds, this catches bars, containers and lines clipped at the
right or bottom edge before a deck is delivered.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation


def audit(path: Path, tolerance_pt: float = 0.25) -> dict:
    deck = Presentation(path)
    tolerance = round(tolerance_pt * 12700)
    failures = []
    checked = 0
    for slide_index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            checked += 1
            left, top = shape.left, shape.top
            right, bottom = left + shape.width, top + shape.height
            if (left < -tolerance or top < -tolerance or
                    right > deck.slide_width + tolerance or
                    bottom > deck.slide_height + tolerance):
                failures.append({
                    "slide": slide_index,
                    "shape": shape.name,
                    "bounds_inches": [round(v / 914400, 4) for v in
                                      (left, top, right, bottom)],
                    "canvas_inches": [round(v / 914400, 4) for v in
                                      (deck.slide_width, deck.slide_height)],
                })
    return {"pptx": str(path), "slides": len(deck.slides),
            "shapes_checked": checked, "out_of_canvas": failures,
            "passed": not failures}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--tolerance-pt", type=float, default=0.25)
    args = parser.parse_args()
    result = audit(args.pptx, args.tolerance_pt)
    content = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(content)
    else:
        print(content, end="")
    return 0 if result["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
