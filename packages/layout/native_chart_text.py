"""Explicit chart typefaces for python-pptx authoring.

Call after authoring chart, axis, legend and label styles. This writes the
chosen family; it does not choose a font, change sizes, or certify rendering.
"""


def bind_chart_typeface(chart, family):
    """Bind existing DrawingML run/default-run styles without skipping siblings.

    lxml's live depth-first iterator can lose later style targets while their
    children are removed/replaced. Freeze targets before making any mutation.
    Styles created after this call need another authoring call. Native reopen,
    font-file identity and visible labels remain separate observations.
    """
    from pptx.oxml.xmlchemy import OxmlElement

    if not isinstance(family, str) or not family.strip() or family != family.strip():
        raise ValueError('Pass the exact nonempty task font family')
    root = getattr(chart, '_chartSpace', None)
    if root is None:
        raise TypeError('Pass a python-pptx Chart, not a shape or text frame')
    targets = [element for element in list(root.iter())
               if element.tag in {OxmlElement('a:rPr').tag,
                                  OxmlElement('a:defRPr').tag,
                                  OxmlElement('a:endParaRPr').tag}]
    font_tags = [OxmlElement('a:' + script).tag for script in ('latin', 'ea', 'cs')]
    for target in targets:
        for tag, script in zip(font_tags, ('latin', 'ea', 'cs')):
            for old in list(target):
                if old.tag == tag:
                    target.remove(old)
            element = OxmlElement('a:' + script)
            element.set('typeface', family)
            target.append(element)
    return {'family': family, 'styles_authored': len(targets),
            'scripts_authored': ['latin', 'ea', 'cs'],
            'scope': 'Existing chart text styles only; no font or pixel acceptance.'}


def set_manual_legend_colors(swatch, text_frame, *, mark_color, text_color):
    """Style existing manual legend parts with two explicit color decisions.

    A mark's relative emphasis does not choose the identifying words' color.
    No geometry, font, contrast threshold or automatic palette is supplied.
    Equal colors are allowed when explicitly chosen. Inspect actual rendering
    and visible association separately; this reports authorship, not quality.
    """
    from pptx.dml.color import RGBColor

    def rgb(value):
        if not isinstance(value, str):
            raise ValueError('Pass an explicit six-digit RGB color')
        cleaned = value.lstrip('#')
        if len(cleaned) != 6:
            raise ValueError('Pass an explicit six-digit RGB color')
        return RGBColor.from_string(cleaned.upper())

    mark, text = rgb(mark_color), rgb(text_color)
    if not hasattr(swatch, 'fill') or not hasattr(text_frame, '_txBody'):
        raise TypeError('Pass the native swatch shape and label TextFrame')
    swatch.fill.solid()
    swatch.fill.fore_color.rgb = mark
    count = 0
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = text
            count += 1
    return {'mark_color': str(mark), 'text_color': str(text),
            'label_runs_authored': count,
            'scope': 'Existing manual swatch/label colors only; no legibility or semantic verdict.'}


def set_point_label_wrapping(chart, requests):
    """Set wrap only on explicitly selected, existing custom point labels.

    Each request supplies series_index, point_index, expected_text and
    word_wrap (True/False/None). All selectors and literal identities are
    checked before any mutation. None restores inheritance, not no-wrap.
    Existing text, runs, style/position, number caches and workbooks stay put;
    automatic labels are not converted to custom text. This uses python-pptx's
    existing TextFrame wrapping behavior, not a new layout engine.

    A no-wrap request can still overflow, collide or retain explicit breaks.
    Observe the actual chart and neighbours after native export. No fit,
    automatic label update, semantic correspondence or GUI verdict is returned.
    """
    from pptx.chart.chart import Chart
    from pptx.text.text import TextFrame

    if not isinstance(chart, Chart):
        raise TypeError('Pass a python-pptx Chart')
    if not isinstance(requests, (list, tuple)):
        raise TypeError('Pass explicit point-label requests as a list or tuple')
    prepared, seen = [], set()
    required = {'series_index', 'point_index', 'expected_text', 'word_wrap'}
    for item in requests:
        if not isinstance(item, dict) or set(item) != required:
            raise ValueError('Each request needs exactly series_index, point_index, expected_text, word_wrap')
        si, pi = item['series_index'], item['point_index']
        if type(si) is not int or type(pi) is not int or si < 0 or pi < 0:
            raise ValueError('Indices must be nonnegative integers, excluding booleans')
        if si >= len(chart.series):
            raise ValueError('series_index is outside this chart')
        series = chart.series[si]
        if pi >= len(series.points):
            raise ValueError('point_index is outside this series')
        if (si, pi) in seen:
            raise ValueError('Select each point label only once per call')
        seen.add((si, pi))
        expected, choice = item['expected_text'], item['word_wrap']
        if not isinstance(expected, str) or not expected:
            raise ValueError('expected_text must be the existing nonempty literal label')
        if choice is not None and type(choice) is not bool:
            raise ValueError('word_wrap must be True, False or None')
        # DataLabel.text_frame can add rich text and remove adjacent spPr/txPr.
        # Bind the existing body directly, retaining all current label siblings.
        labels = series._element.xpath('./c:dLbls/c:dLbl')
        matches = [el for el in labels if el.xpath('./c:idx/@val') == [str(pi)]]
        if len(matches) != 1:
            raise ValueError('Selected custom point label must already exist once')
        rich = matches[0].xpath('./c:tx/c:rich')
        if len(rich) != 1 or len(rich[0].xpath('./a:bodyPr')) != 1:
            raise ValueError('Existing custom label needs one rich text body/bodyPr')
        tf = TextFrame(rich[0], chart)
        if tf.text != expected:
            raise ValueError('Selected point label text differs from expected_text')
        prepared.append((si, pi, expected, choice, tf, tf.word_wrap))
    observations = []
    for si, pi, expected, choice, tf, before in prepared:
        tf.word_wrap = choice
        observations.append({'series_index': si, 'point_index': pi,
                             'text': expected, 'previous_word_wrap': before,
                             'authored_word_wrap': tf.word_wrap})
    return {'selected_labels': observations,
            'scope': 'Explicit existing custom point-label wrapping only; no text/style/data/geometry mutation or layout-quality verdict.'}


def set_point_value_label_formats(chart, requests):
    """Format explicitly selected existing value labels without fixed text.

    Each request supplies series_index, point_index, expected_value and
    number_format. Validate the whole request before changing any label.
    Existing numeric labels may have their own font, position or shape style;
    those siblings, point values, caches and workbooks are left unchanged.
    Custom text (rich or referenced) is rejected, never silently discarded.

    sourceLinked=False selects this explicit format, not a literal value.
    This does not enable labels, choose placement/typeface, parse Excel format
    syntax, or certify rendering, workbook editing or future GUI behaviour.
    After changes to the data/meaning, verify labels and nearby marks again.
    """
    import math
    from decimal import Decimal
    from numbers import Real

    from pptx.chart.chart import Chart
    from pptx.oxml.xmlchemy import OxmlElement

    if not isinstance(chart, Chart):
        raise TypeError('Pass a python-pptx Chart')
    if not isinstance(requests, (list, tuple)):
        raise TypeError('Pass explicit point-label format requests')
    required = {'series_index', 'point_index', 'expected_value', 'number_format'}
    prepared, seen = [], set()
    for item in requests:
        if not isinstance(item, dict) or set(item) != required:
            raise ValueError('Each request needs series_index, point_index, expected_value and number_format')
        si, pi = item['series_index'], item['point_index']
        if type(si) is not int or type(pi) is not int or si < 0 or pi < 0:
            raise ValueError('Indices must be nonnegative integers, excluding booleans')
        if si >= len(chart.series):
            raise ValueError('series_index is outside this chart')
        series = chart.series[si]
        if pi >= len(series.points):
            raise ValueError('point_index is outside this series')
        if (si, pi) in seen:
            raise ValueError('Select each point only once per call')
        seen.add((si, pi))
        expected = item['expected_value']
        if isinstance(expected, bool) or not isinstance(expected, (Real, Decimal)):
            raise TypeError('expected_value must be a finite number')
        if not math.isfinite(float(expected)):
            raise ValueError('expected_value must be finite')
        values = getattr(series, 'values', None)
        if values is None or pi >= len(values) or values[pi] is None:
            raise ValueError('Selected series needs an existing numeric point value')
        actual = values[pi]
        if not math.isfinite(float(actual)) or float(actual) != float(expected):
            raise ValueError('Selected numeric value differs from expected_value')
        code = item['number_format']
        if not isinstance(code, str) or not code.strip():
            raise ValueError('Pass an explicit nonempty numeric format code')
        labels = series._element.xpath('./c:dLbls/c:dLbl')
        matches = [el for el in labels if el.xpath('./c:idx/@val') == [str(pi)]]
        if len(matches) != 1:
            raise ValueError('Selected point label must already exist once')
        label = matches[0]
        if label.xpath('./c:tx'):
            raise ValueError('Custom label text requires a separate explicit content decision')
        old = label.xpath('./c:numFmt')
        if len(old) > 1:
            raise ValueError('Existing point label has duplicate number formats')
        prepared.append((si, pi, float(actual), code, label, old))
    observations = []
    for si, pi, actual, code, label, old in prepared:
        for element in old:
            label.remove(element)
        fmt = OxmlElement('c:numFmt')
        fmt.set('formatCode', code)
        fmt.set('sourceLinked', '0')
        label.insert_element_before(fmt, 'c:spPr', 'c:txPr', 'c:dLblPos',
                                    'c:showLegendKey', 'c:showVal', 'c:showCatName',
                                    'c:showSerName', 'c:showPercent',
                                    'c:showBubbleSize', 'c:separator', 'c:extLst')
        observations.append({'series_index': si, 'point_index': pi,
                             'value_read': actual, 'number_format': code,
                             'source_format_linked': False})
    return {'selected_labels': observations,
            'scope': 'Explicit numeric formats on existing value labels only; no custom text or data/geometry edits. Font/style/position siblings preserved; supplied formats may affect rendered appearance. No rendering/GUI verdict.'}
