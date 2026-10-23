"""Caller-selected rectangle connections, native line ends and explicit refresh.

Uses python-pptx's experimental attachment API in a narrow geometry scope.
Callers choose relation endpoints and arrow direction; the creation convenience
has overridable line style defaults. It does not select business meaning or
route around obstacles. Native attachment and programmatic refresh do not prove
visible correspondence or a target application's interactive editing behavior.
"""
from __future__ import annotations

from typing import Any

_SITES = {"top": 0, "left": 1, "bottom": 2, "right": 3}


def _direct_slide_container(shape: Any) -> Any:
    from pptx.oxml.ns import qn

    parent = shape._element.getparent()
    if (parent is None or parent.tag != qn("p:spTree")
            or parent.getparent() is None or parent.getparent().tag != qn("p:cSld")
            or parent.getparent().getparent() is None
            or parent.getparent().getparent().tag != qn("p:sld")):
        raise ValueError("Only direct slide spTree objects are supported; groups are not")
    return parent


def _endpoint(value: Any, container: Any) -> tuple[Any, int]:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.shapes.autoshape import Shape

    if not isinstance(value, tuple) or len(value) != 2:
        raise ValueError("An endpoint must be an explicit (rectangle, site) tuple")
    shape, site = value
    if not isinstance(site, str) or site not in _SITES:
        raise ValueError("Choose site top, left, bottom or right explicitly")
    if not isinstance(shape, Shape) or shape.is_placeholder:
        raise ValueError("Only ordinary native rectangle shapes are supported")
    if _direct_slide_container(shape) is not container:
        raise ValueError("Connector and targets must share the same direct slide spTree")
    try:
        is_rectangle = shape.auto_shape_type == MSO_SHAPE.RECTANGLE
    except ValueError:
        is_rectangle = False
    if not is_rectangle or shape.width <= 0 or shape.height <= 0:
        raise ValueError("Target must be a nonempty native rectangle")
    if shape.rotation % 360:
        raise ValueError("Rotated targets are outside the supported scope")
    xfrm = shape._element.spPr.xfrm
    if xfrm is None or xfrm.flipH or xfrm.flipV:
        raise ValueError("Target requires an explicit unflipped native transform")
    return shape, _SITES[site]


def attach_rectangular_connector(
    connector: Any, *, begin: tuple[Any, str] | None = None,
    end: tuple[Any, str] | None = None,
) -> Any:
    """Attach caller-selected ends of an existing straight connector.

    Example: ``attach_rectangular_connector(line, begin=(a, "right"),
    end=(b, "left"))``. An omitted end remains untouched. Both ends omitted
    is an error. Targets must be ordinary unrotated, unflipped rectangles in
    the connector's direct slide spTree. Nested groups and cross-slide objects
    are refused even when they share an owning slide part or numerical IDs.

    Every input is checked before invoking either library connection method,
    so a rejected end cannot leave the other end partially changed. Valid
    attachment writes shape id/site and snaps the selected stored coordinates
    using the existing library API. It does not monitor later property edits,
    implement routing or move separate arrowheads/labels. Repeating an explicit
    call after a programmatic move can refresh that endpoint; GUI behavior and
    rendered/semantic correspondence require separate observation.
    """
    from pptx.shapes.connector import Connector

    if not isinstance(connector, Connector):
        raise ValueError("Pass an existing native connector")
    container = _direct_slide_container(connector)
    if connector.rotation % 360:
        raise ValueError("A rotated connector is outside the supported scope")
    geom = connector._element.spPr.prstGeom
    if geom is None or geom.get("prst") not in {"line", "straightConnector1"}:
        raise ValueError("Only an existing straight connector is supported")
    if begin is None and end is None:
        raise ValueError("Select at least one endpoint explicitly")
    selected_begin = None if begin is None else _endpoint(begin, container)
    selected_end = None if end is None else _endpoint(end, container)
    if selected_begin is not None:
        connector.begin_connect(*selected_begin)
    if selected_end is not None:
        connector.end_connect(*selected_end)
    return connector


_LINE_END_TYPES = frozenset({"none", "triangle", "stealth", "diamond", "oval", "arrow"})
_LINE_END_SIZES = frozenset({"sm", "med", "lg"})


def add_rectangular_connector(slide: Any, *, begin: tuple[Any, str],
                              end: tuple[Any, str], name: str | None = None,
                              color: str = "667085", width_pt: float = 1.5,
                              head: Any = None, tail: Any = None) -> Any:
    """Create a native straight connector and register both node/site ends.

    Targets share the slide's direct spTree and the existing rectangle scope.
    Arrow ends belong to the line itself. Node movement through python-pptx
    does not automatically recalculate geometry: call refresh afterwards.
    Native IDs/sites are not evidence of GUI dragging behavior.
    """
    from math import isfinite
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    container = slide.shapes._spTree
    selected_begin = _endpoint(begin, container)
    selected_end = _endpoint(end, container)
    _line_end_choice(head); _line_end_choice(tail)
    if isinstance(width_pt, bool) or not isinstance(width_pt, (int, float)) or not isfinite(width_pt) or width_pt <= 0:
        raise ValueError("width_pt must be a positive finite number")
    chosen_color = RGBColor.from_string(color.lstrip("#"))
    a, b = selected_begin[0], selected_end[0]
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      a.left, a.top, b.left+b.width, b.top+b.height)
    if name is not None: line.name = name
    line.line.color.rgb = chosen_color
    line.line.width = Pt(width_pt)
    attach_rectangular_connector(line, begin=begin, end=end)
    if head is not None or tail is not None:
        set_connector_line_ends(line, head=head, tail=tail)
    return line


def _registered_rectangular_endpoints(connector: Any) -> dict[str, Any]:
    """Read and validate a registered straight connector without mutation."""
    from pptx.oxml.ns import qn
    from pptx.shapes.connector import Connector

    if not isinstance(connector, Connector):
        raise ValueError("Pass an existing native connector")
    container = _direct_slide_container(connector)
    if connector.rotation % 360:
        raise ValueError("A rotated connector is outside the supported scope")
    geom = connector._element.spPr.prstGeom
    if geom is None or geom.get("prst") not in {"line", "straightConnector1"}:
        raise ValueError("Only an existing straight connector is supported")
    shapes = {}
    for shape in connector._parent:
        if shape.shape_id in shapes:
            raise ValueError("Ambiguous duplicate native shape IDs on the slide")
        shapes[shape.shape_id] = shape
    nv = connector._element.find(qn("p:nvCxnSpPr"))
    props = None if nv is None else nv.find(qn("p:cNvCxnSpPr"))
    sites = {value:key for key,value in _SITES.items()}
    selected = {}
    for tag, key in (("stCxn", "begin"), ("endCxn", "end")):
        nodes = [] if props is None else props.findall(qn("a:"+tag))
        if len(nodes) != 1:
            raise ValueError("Both native node endpoints must be registered exactly once")
        try:
            node = nodes[0]
            endpoint = (shapes[int(node.get("id"))], sites[int(node.get("idx"))])
        except (ValueError, TypeError, KeyError) as exc:
            raise ValueError("Registered endpoint target/site cannot be resolved") from exc
        _endpoint(endpoint, container)
        selected[key] = endpoint
    return selected


def refresh_rectangular_connector(connector: Any) -> Any:
    """Recompute one path from registered IDs/sites, without choosing a route.

    Explicit programmatic maintenance, not an automatic observer or GUI test.
    Missing/moved-out-of-scope targets are rejected before either end changes.
    """
    selected = _registered_rectangular_endpoints(connector)
    return attach_rectangular_connector(connector, **selected)


def refresh_rectangular_connectors(connectors: Any) -> tuple[Any, ...]:
    """Refresh an explicitly selected same-slide batch after full validation.

    No selected connector changes if any input/registration is rejected.
    Selection is a nonempty list/tuple of distinct ordinary straight connectors
    in one direct slide spTree, with both rectangular targets registered.
    Current node geometry is read, while target IDs/sites, native line ends,
    line style and separate labels are preserved. Nodes are not moved here.

    This does not discover all affected edges, choose different sites, shorten
    or route paths, move independent labels, or install an edit observer.
    The caller selects the actual affected set and inspects the whole graphic.
    Prevalidation does not roll back earlier caller edits or unexpected library
    runtime exceptions; no GUI, rendering or semantic-fit verdict is implied.
    """
    if not isinstance(connectors, (list, tuple)) or not connectors:
        raise ValueError("Pass a nonempty explicit connector list or tuple")
    prepared, seen = [], set()
    container = None
    for connector in connectors:
        selected = _registered_rectangular_endpoints(connector)
        current = _direct_slide_container(connector)
        if container is None:
            container = current
        elif current is not container:
            raise ValueError("Selected connectors must share one direct slide spTree")
        if connector._element in seen:
            raise ValueError("Select each native connector only once")
        seen.add(connector._element)
        prepared.append((connector, selected))
    for connector, selected in prepared:
        attach_rectangular_connector(connector, **selected)
    return tuple(connector for connector, _ in prepared)


def _line_end_choice(value: Any) -> dict[str, str] | str | None:
    from collections.abc import Mapping

    if value is None:
        return None
    if isinstance(value, str):
        if value == "inherit":
            return value
        raise ValueError("Use an explicit type mapping or 'inherit'")
    if not isinstance(value, Mapping) or "type" not in value:
        raise ValueError("A selected line end requires an explicit type")
    if set(value) - {"type", "w", "len"}:
        raise ValueError("Line-end attributes are limited to type, w and len")
    result = dict(value)
    if not isinstance(result["type"], str) or result["type"] not in _LINE_END_TYPES:
        raise ValueError("Unsupported native line-end type")
    for key in ("w", "len"):
        if key in result and (not isinstance(result[key], str) or result[key] not in _LINE_END_SIZES):
            raise ValueError("Line-end w and len must be sm, med or lg")
    return result


def set_connector_line_ends(
    connector: Any, *, head: Any = None, tail: Any = None,
) -> Any:
    """Explicitly edit selected native line-end decorations, not separate shapes.

    Example: ``set_connector_line_ends(line, tail={"type": "triangle",
    "w": "med", "len": "med"})``. Names refer to DrawingML headEnd/tailEnd;
    they do not select a business source/destination or a fixed screen side.
    Only a nonzero straight connector in a direct slide spTree is supported.
    Its native flipH/flipV are preserved; a rotated/grouped connector is refused.

    None leaves that end untouched. A mapping requires type; omitted w/len
    preserve existing explicit attributes or leave them inherited. type='none'
    explicitly requests no decoration. The literal 'inherit' removes only that
    selected end element, allowing theme/style inheritance; it does not promise
    no visible decoration. An existing empty a:ln is not removed.

    Both choices and existing element multiplicity are checked before any
    mutation. Geometry, binding IDs/sites, line width/color, the other end,
    unknown attributes and unrelated XML are preserved. No line, independent
    arrowhead or label is created, selected, migrated, moved or routed. Return
    the existing connector. Reopen/render, GUI behavior and the relation's
    semantic direction remain separate evidence.
    """
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.shapes.connector import Connector

    if not isinstance(connector, Connector):
        raise ValueError("Pass an existing native connector")
    _direct_slide_container(connector)
    if connector.rotation % 360:
        raise ValueError("A rotated connector is outside the supported scope")
    sp_pr = connector._element.spPr
    geom = sp_pr.prstGeom
    if geom is None or geom.get("prst") not in {"line", "straightConnector1"}:
        raise ValueError("Only an existing straight connector is supported")
    if connector.width == 0 and connector.height == 0:
        raise ValueError("A zero-length connector has no supported end direction")
    if head is None and tail is None:
        raise ValueError("Select at least one line end explicitly")
    choices = [("headEnd", _line_end_choice(head)), ("tailEnd", _line_end_choice(tail))]
    lines = sp_pr.findall(qn("a:ln"))
    if len(lines) > 1:
        raise ValueError("Ambiguous native line properties")
    ln = lines[0] if lines else None
    if ln is not None:
        for name, _ in choices:
            if len(ln.findall(qn("a:" + name))) > 1:
                raise ValueError("Ambiguous native line-end properties")

    # All input and existing-node checks above precede get_or_add or removal.
    if ln is None and all(choice in (None, "inherit") for _, choice in choices):
        return connector
    if ln is None:
        ln = sp_pr.get_or_add_ln()
    for name, choice in choices:
        if choice is None:
            continue
        node = ln.find(qn("a:" + name))
        if choice == "inherit":
            if node is not None:
                ln.remove(node)
            continue
        if node is None:
            node = OxmlElement("a:" + name)
            successors = ("a:tailEnd", "a:extLst") if name == "headEnd" else ("a:extLst",)
            ln.insert_element_before(node, *successors)
        for key, value in choice.items():
            node.set(key, value)
    return connector
