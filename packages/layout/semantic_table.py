"""Content-measured native tables. Column roles describe data, not importance.

This is an authoring primitive, not a slide template. Art supplies the intended
comparison, explicit emphasis and available region; overflow returns to Art.
"""
from dataclasses import dataclass
from math import isfinite
import sys
from pathlib import Path

if __package__ in {None, ""}:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from PIL import ImageFont
from packages.layout.text_flow import wrap


@dataclass(frozen=True)
class Column:
    label: str
    role: str = "text"  # text, number, period; never inferred from array position
    weight: float = 1.0


def plan(columns, rows, *, width, font_path, size=21, header_size=18,
         total_rows=(), emphasis=(), max_height=None, padding=(.17, .13),
         caption="", comparison="", protected_terms=None, native_width_scale=1.0):
    # Optional calibration of usable line width for the actual target/font.
    # Physical columns, font size and padding remain Art's chosen values.
    if not isfinite(native_width_scale) or native_width_scale < 1:
        raise ValueError("native width scale must be finite and at least 1")
    if not columns or not rows or not comparison.strip() or not caption.strip():
        raise ValueError("table needs data, visible scope caption and comparison intent")
    if any(c.role not in {"text", "number", "period"} or c.weight <= 0 for c in columns):
        raise ValueError("invalid column semantics")
    if any(len(row) != len(columns) for row in rows):
        raise ValueError("ragged table")
    if any(i < 0 or i >= len(rows) for i in total_rows):
        raise ValueError("total row outside data")
    widths = [width * c.weight / sum(c.weight for c in columns) for c in columns]
    px, py = padding
    cells, heights = [], []
    for ri, row in enumerate([[c.label for c in columns]] + rows):
        line = []
        for ci, value in enumerate(row):
            pt = header_size if ri == 0 else size
            font = ImageFont.truetype(str(font_path), round(pt * 120 / 72))
            lines = wrap(value, font, (widths[ci] - 2 * px) * 120 / native_width_scale, **({"protected_terms":protected_terms} if protected_terms is not None else {}))
            # Numeric wrapping changes a comparison's reading; return to Art.
            if ri and columns[ci].role == "number" and len(lines) != 1:
                raise ValueError(f"numeric cell wraps: data row {ri-1}, column {ci}")
            line.append(dict(text=str(value), lines=lines, size_pt=pt,
                             align="right" if columns[ci].role == "number" else "left",
                             bold=ri == 0 or ri - 1 in total_rows or (ri - 1, ci) in emphasis,
                             header=ri == 0, total=ri - 1 in total_rows,
                             emphasis=(ri - 1, ci) in emphasis))
        cells.append(line)
        heights.append(max(len(c["lines"]) * c["size_pt"] / 72 * 1.18 for c in line) + 2 * py)
    if max_height is not None and sum(heights) > max_height:
        raise ValueError(f"table needs {sum(heights):.2f}in, available {max_height:.2f}in; revise Art")
    return dict(columns=[vars(c) for c in columns], cells=cells, widths=widths,
                heights=heights, padding=list(padding), width=width,
                height=sum(heights), caption=caption, comparison=comparison,
                native_width_scale=native_width_scale)


def write(slide, spec, x, y, *, family="STKaiti", ink="172B3A",
          accent="185D72", line="CAD5DA", name="semantic-table"):
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    shape = slide.shapes.add_table(len(spec["cells"]), len(spec["widths"]),
                                   Inches(x), Inches(y), Inches(spec["width"]), Inches(spec["height"]))
    shape.name = name
    table = shape.table
    # Disable theme banding/first-column defaults; make every cell explicit.
    for attr in ("first_row", "first_col", "last_row", "last_col", "horz_banding", "vert_banding"):
        setattr(table, attr, False)
    for ci, width in enumerate(spec["widths"]):
        table.columns[ci].width = Inches(width)
    for ri, row in enumerate(spec["cells"]):
        table.rows[ri].height = Inches(spec["heights"][ri])
        for ci, item in enumerate(row):
            cell = table.cell(ri, ci)
            px, py = spec["padding"]
            cell.margin_left = cell.margin_right = Inches(px)
            cell.margin_top = cell.margin_bottom = Inches(py)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("E9F0F2" if item["header"] or item["total"] else "FFFFFF")
            tf = cell.text_frame
            tf.clear(); tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.vertical_anchor = MSO_ANCHOR.TOP
            # Do not call a generic text-box setter that resets cell padding.
            for li, value in enumerate(item["lines"]):
                p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.RIGHT if item["align"] == "right" else PP_ALIGN.LEFT
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = Pt(item["size_pt"] * 1.18)
                run = p.add_run(); run.text = value
                run.font.name = family; run.font.size = Pt(item["size_pt"]); run.font.bold = item["bold"]
                run.font.color.rgb = RGBColor.from_string(accent if item["emphasis"] or item["total"] else ink)
                for tag in ("a:ea", "a:cs"):
                    font = OxmlElement(tag); font.set("typeface", family)
                    run._r.get_or_add_rPr().append(font)
            tc = cell._tc.get_or_add_tcPr()
            for side in ("L", "R", "T", "B"):
                ln = OxmlElement("a:ln" + side); ln.set("w", "6350")
                fill = OxmlElement("a:solidFill"); col = OxmlElement("a:srgbClr")
                col.set("val", line if side == "B" else "FFFFFF"); fill.append(col); ln.append(fill); tc.append(ln)
    return shape


def preview(draw, spec, x, y, *, scale=120, font_path, ink="172B3A", accent="185D72"):
    """Uses the same widths, line breaks, padding and row heights as native output."""
    cy = y * scale
    for ri, row in enumerate(spec["cells"]):
        cx = x * scale; height = spec["heights"][ri] * scale
        for ci, item in enumerate(row):
            width = spec["widths"][ci] * scale
            draw.rectangle((cx, cy, cx + width, cy + height), fill="#E9F0F2" if item["header"] or item["total"] else "white")
            font = ImageFont.truetype(str(font_path), round(item["size_pt"] * scale / 72))
            for li, value in enumerate(item["lines"]):
                tx = cx + spec["padding"][0] * scale
                if item["align"] == "right": tx = cx + width - spec["padding"][0] * scale - font.getlength(value)
                ty = cy + spec["padding"][1] * scale + li * item["size_pt"] / 72 * 1.18 * scale
                draw.text((tx, ty), value, font=font, fill="#" + (accent if item["emphasis"] or item["total"] else ink), anchor="lt", stroke_width=0)
            draw.line((cx, cy + height, cx + width, cy + height), fill="#CAD5DA", width=1)
            cx += width
        cy += height


def column_demands(columns, rows, *, font_path, size=23, header_size=18, padding=.17, width_scale=1.06):
    """Minimum single-line content widths; Art still chooses available width.

    Not every text cell must stay on one line. This is an inspectable diagnostic,
    not an automatic full-slide-width template.
    """
    if any(len(row)!=len(columns) for row in rows):raise ValueError('ragged table')
    demands=[]
    for ci,col in enumerate(columns):
        labels=[(col.label,header_size)]+[(str(row[ci]),size) for row in rows]
        demands.append(max(ImageFont.truetype(str(font_path),round(pt*120/72)).getlength(value)*width_scale/120+2*padding for value,pt in labels))
    return demands
