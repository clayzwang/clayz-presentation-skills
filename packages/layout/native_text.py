"""Editable native text: logical paragraphs, native reflow and role defaults.

No width measurement is converted into hard/soft line breaks here. Explicit
newlines belong to the caller's semantic paragraphs; inherited content is not
rewritten. This module authors properties, not fit or GUI acceptance evidence.
"""
from __future__ import annotations

from math import isfinite
from typing import Any


def _positive(value: Any, name: str, *, zero: bool = False) -> float:
    if isinstance(value, bool) or not isinstance(value, (float, int)):
        raise ValueError(f"{name} must be a finite number")
    value = float(value)
    if not isfinite(value) or (value < 0 if zero else value <= 0):
        raise ValueError(f"invalid {name}")
    return value


def prepare_native_text(target: Any, *, role: str = "body", word_wrap: bool = True,
                        line_spacing: float = 1.0, space_before_pt: float = 0,
                        space_after_pt: float | None = None,
                        horizontal: str | None = None, vertical: str | None = None,
                        margins_in: dict[str, float] | None = None,
                        fixed_line_spacing_pt: float | None = None,
                        exception_reason: str | None = None) -> Any:
    """Prepare an existing Shape, Cell or TextFrame without replacing runs.

    Body defaults to top/left and 6pt paragraph-after; label/header/button/node
    defaults to middle/center and 0pt after. Margins default to .12in sides and
    .06in top/bottom. Pass actual table padding explicitly to preserve it.
    A float line_spacing writes native percentage spacing, not fixed points.
    Explicit fixed spacing or disabled wrap requires a named local exception.
    No automatic font shrink or text/shape resizing is performed.
    """
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Inches, Pt

    if role not in {"body", "label", "header", "button", "node"}:
        raise ValueError("unknown text role")
    if not isinstance(word_wrap, bool):
        raise ValueError("word_wrap must be bool")
    multiple = _positive(line_spacing, "line_spacing")
    before = _positive(space_before_pt, "space_before_pt", zero=True)
    after = _positive(6 if role == "body" else 0, "space_after_pt", zero=True) if space_after_pt is None else _positive(space_after_pt, "space_after_pt", zero=True)
    fixed = None if fixed_line_spacing_pt is None else _positive(fixed_line_spacing_pt, "fixed_line_spacing_pt")
    if (not word_wrap or fixed is not None) and (not isinstance(exception_reason, str) or not exception_reason.strip()):
        raise ValueError("fixed/no-wrap exception needs a nonempty reason")
    align = horizontal or ("left" if role == "body" else "center")
    anchor = vertical or ("top" if role == "body" else "middle")
    if align not in {"left", "center", "right", "justify"} or anchor not in {"top", "middle", "bottom"}:
        raise ValueError("invalid text alignment")
    margins = {"left": .12, "right": .12, "top": .06, "bottom": .06}
    if margins_in is not None:
        if set(margins_in) != set(margins):
            raise ValueError("margins_in needs left/right/top/bottom")
        margins = {k: _positive(v, "margin_" + k, zero=True) for k, v in margins_in.items()}
    tf = target.text_frame if hasattr(target, "text_frame") else target
    if not hasattr(tf, "_txBody"):
        raise TypeError("Pass a native Shape, Cell or TextFrame")
    padding_target = target if hasattr(target, "_tc") else tf
    tf.word_wrap = word_wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for side, value in margins.items():
        setattr(padding_target, "margin_" + side, Inches(value))
    # A table cell stores its vertical anchor on tcPr, not only bodyPr.
    anchor_target = target if hasattr(target, "_tc") else tf
    anchor_target.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    for paragraph in tf.paragraphs:
        paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}[align]
        paragraph.line_spacing = Pt(fixed) if fixed is not None else multiple
        paragraph.space_before = Pt(before)
        paragraph.space_after = Pt(after)
    return tf


def write_native_text(target: Any, text: str, *, family: str = "STKaiti",
                      size_pt: float = 20, color: str = "172B3A",
                      bold: bool = False, **flow: Any) -> Any:
    """Write exact logical text and explicit native font fields, then prepare.

    Only caller-authored \n/\v are written. No measured visual lines are added.
    For mixed emphasis, write original paragraph runs yourself and call
    prepare_native_text afterwards; it preserves those runs and their styling.
    """
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    if not isinstance(text, str):
        raise TypeError("text must be a string")
    size = _positive(size_pt, "size_pt")
    tf = target.text_frame if hasattr(target, "text_frame") else target
    tf.text = text
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.name = family
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
            rpr = run._r.get_or_add_rPr()
            for tag in ("latin", "ea", "cs"):
                node = rpr.find(qn("a:" + tag))
                if node is None:
                    node = OxmlElement("a:" + tag); rpr.append(node)
                node.set("typeface", family)
    return prepare_native_text(target, **flow)


def add_native_textbox(slide: Any, text: str, x: float, y: float, w: float, h: float,
                       *, name: str | None = None, **options: Any) -> Any:
    """Create one native text box; coordinates are inches."""
    from pptx.util import Inches
    obj = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name is not None: obj.name = name
    write_native_text(obj, text, **options)
    return obj


def add_text_shape(slide: Any, text: str, x: float, y: float, w: float, h: float,
                   *, shape_type: Any = None, name: str | None = None,
                   fill: str | None = None, line: str | None = None,
                   role: str = "label", **options: Any) -> Any:
    """Create a single native AutoShape with its own text, not an overlay.

    Coordinates are inches. One object owns fill, geometry, label and margins.
    Separate labels/annotations remain available when their placement needs it.
    """
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    obj = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if shape_type is None else shape_type,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    for effect in obj._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):
        effect.set('idx', '0')
    if name is not None: obj.name = name
    if fill is None: obj.fill.background()
    else: obj.fill.solid(); obj.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    if line is None: obj.line.fill.background()
    else: obj.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
    write_native_text(obj, text, role=role, **options)
    return obj
