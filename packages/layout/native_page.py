"""Neutral native page initialization; composition remains explicit."""
from __future__ import annotations
from typing import Any


def add_native_page(presentation: Any, *, layout: Any, background: str | None = None) -> Any:
    """Add only the caller-selected layout and optional background.

    No title, rule, navigation mark, body band or reserved coordinates are
    authored. A blank layout produces an empty canvas. A user-master layout
    retains its inherited placeholders/background; None preserves background.
    Title/content placement is a separate, explicit authoring decision.
    """
    rgb = None
    if background is not None:
        from pptx.dml.color import RGBColor
        if not isinstance(background, str) or len(background.lstrip('#')) != 6:
            raise ValueError('Pass an explicit six-digit RGB background')
        rgb = RGBColor.from_string(background.lstrip('#').upper())
    slide = presentation.slides.add_slide(layout)
    if rgb is not None:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb
    return slide
