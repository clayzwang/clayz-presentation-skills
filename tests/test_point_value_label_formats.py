import io
import unittest
import xml.etree.ElementTree as ET
import zipfile

from lxml import etree

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.util import Inches, Pt

from packages.layout.native_chart_text import set_point_value_label_formats


class PointValueLabelFormatTest(unittest.TestCase):
    def fixture(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ['Positive', 'Negative', 'Zero']
        data.add_series('Values', [1.25, -.5, 0])
        chart = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1),
                                      Inches(1), Inches(6), Inches(4), data).chart
        chart.plots[0].has_data_labels = True
        for index, point in enumerate(chart.series[0].points):
            label = point.data_label
            label.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
            label.font.name = 'STKaiti'
            label.font.size = Pt(17)
            label.font.color.rgb = RGBColor.from_string('18323D')
        return presentation, chart

    def request(self, index, value, code='0.00;-0.00;0.00'):
        return dict(series_index=0, point_index=index,
                    expected_value=value, number_format=code)

    def unchanged_siblings(self, chart):
        return [[etree.tostring(element, with_tail=False) for element in label
                 if element.tag.rsplit('}', 1)[-1] != 'numFmt']
                for label in chart.series[0]._element.xpath('./c:dLbls/c:dLbl')]

    def test_selected_numeric_formats_preserve_data_style_and_other_labels(self):
        presentation, chart = self.fixture()
        before = self.unchanged_siblings(chart)
        workbook = chart.part.chart_workbook.xlsx_part.blob
        result = set_point_value_label_formats(chart, [self.request(0, 1.25),
                                                       self.request(1, -.5)])
        self.assertEqual(chart.series[0].values, (1.25, -.5, 0))
        self.assertEqual(chart.part.chart_workbook.xlsx_part.blob, workbook)
        self.assertEqual(self.unchanged_siblings(chart), before)
        labels = chart.series[0]._element.xpath('./c:dLbls/c:dLbl')
        for label in labels[:2]:
            self.assertFalse(label.xpath('./c:tx'))
            fmt = label.xpath('./c:numFmt')[0]
            self.assertEqual(dict(fmt.attrib),
                             {'formatCode': '0.00;-0.00;0.00', 'sourceLinked': '0'})
            tags = [element.tag.rsplit('}', 1)[-1] for element in label]
            self.assertLess(tags.index('numFmt'), tags.index('txPr'))
        self.assertFalse(labels[2].xpath('./c:numFmt'))
        self.assertEqual(len(result['selected_labels']), 2)

    def test_replace_data_and_native_reopen_keep_value_generated_labels(self):
        presentation, chart = self.fixture()
        set_point_value_label_formats(chart, [self.request(0, 1.25),
                                              self.request(1, -.5),
                                              self.request(2, 0)])
        labels_before = [label.xml for label in
                         chart.series[0]._element.xpath('./c:dLbls/c:dLbl')]
        new_data = CategoryChartData()
        new_data.categories = ['Positive', 'Negative', 'Zero']
        new_data.add_series('Values', [2.5, -.75, 0])
        chart.replace_data(new_data)
        buffer = io.BytesIO()
        presentation.save(buffer)
        reopened = Presentation(io.BytesIO(buffer.getvalue()))
        actual = reopened.slides[0].shapes[0].chart
        self.assertEqual(actual.series[0].values, (2.5, -.75, 0))
        self.assertEqual([label.xml for label in
                          actual.series[0]._element.xpath('./c:dLbls/c:dLbl')],
                         labels_before)
        self.assertFalse(actual._chartSpace.xpath('.//c:dLbl/c:tx'))
        with zipfile.ZipFile(io.BytesIO(actual.part.chart_workbook.xlsx_part.blob)) as workbook:
            sheet = ET.fromstring(workbook.read('xl/worksheets/sheet1.xml'))
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        values = {cell.attrib['r']: float(cell.find('s:v', ns).text)
                  for cell in sheet.findall('.//s:c', ns)
                  if cell.attrib['r'] in {'B2', 'B3', 'B4'}}
        self.assertEqual(values, {'B2': 2.5, 'B3': -.75, 'B4': 0})

    def test_late_invalid_or_custom_text_requests_are_atomic(self):
        presentation, chart = self.fixture()
        cases = [self.request(1, -.6), self.request(1, -.5, ''),
                 dict(self.request(1, -.5), series_index=True),
                 self.request(0, 1.25)]
        for invalid in cases:
            with self.subTest(invalid=invalid):
                before = chart._chartSpace.xml
                with self.assertRaises((TypeError, ValueError)):
                    set_point_value_label_formats(chart,
                                                  [self.request(0, 1.25), invalid])
                self.assertEqual(chart._chartSpace.xml, before)
        chart.series[0].points[1].data_label.text_frame.text = 'Deliberate custom text'
        before = chart._chartSpace.xml
        with self.assertRaisesRegex(ValueError, 'Custom label text'):
            set_point_value_label_formats(chart, [self.request(0, 1.25),
                                                  self.request(1, -.5)])
        self.assertEqual(chart._chartSpace.xml, before)


if __name__ == '__main__':
    unittest.main()
