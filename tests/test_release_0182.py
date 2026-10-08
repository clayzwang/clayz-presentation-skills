"""Actual native tables and separate execution/integrity/quality outcomes."""
import copy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'packages/validators'))
sys.path.insert(0, str(ROOT))
from tests.test_release_0180 import clean_handoff, relock, write_native
from tests.test_story_visual_handoff import synthetic_handoff
from compare_package_to_pptx import compare, main as compare_main
from validate_art_direction_plan import validate_plan
from validate_build_deviation_log import validate as validate_deviation, sha256
from validate_output_qa import validate_qa, CHECK_KEYS
from verification_result import Issue, result
from pptx import Presentation
from pptx.util import Inches


class VerificationRegressionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        _, self.package, self.plan = clean_handoff(self.folder)
        self.plan['contract_version'] = '2.2'
        self.elements = self.plan['visual_baseline']['slides'][0]['elements']
        for element in self.elements:
            element.update(native_name='ART::' + element['element_id'], render_separately=True)
        relock(self.package, self.plan)

    def write_json(self, name, value):
        path = self.folder / name
        path.write_text(json.dumps(value), encoding='utf-8')
        return path

    def table_deck(self):
        element = copy.deepcopy(self.elements[0])
        element.update(element_id='DATA', native_name='ART::DATA', copy_ids=[],
                       kind='table', native_type='table',
                       table_spec={'headers': ['Region', 'Value'], 'rows': [['North', '12']]})
        self.elements.append(element)
        self.plan['slides'][0]['medium_execution_contract']['minimum_object_counts']['native-table'] = 1
        relock(self.package, self.plan)
        self.assertEqual([], validate_plan(self.package, self.plan))
        path = self.folder / 'native.pptx'
        write_native(self.package, self.plan, path)
        deck = Presentation(path)
        table = next(shape.table for shape in deck.slides[0].shapes if shape.name == 'ART::DATA')
        for row, values in enumerate([['Region', 'Value'], ['North', '12']]):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        deck.save(path)
        return path

    def test_approved_art_table_cells_are_not_extra_text(self):
        path = self.table_deck()
        self.assertEqual([], compare(self.package, self.plan, path))

    def test_changed_cell_remains_quality_failure_even_in_diagnostic_mode(self):
        path = self.table_deck()
        deck = Presentation(path)
        next(shape.table for shape in deck.slides[0].shapes if shape.name == 'ART::DATA').cell(1, 1).text = '99'
        deck.save(path)
        errors = compare(self.package, self.plan, path, allow_extra_text=True)
        diagnostic = result('comparison', errors, quality_checked=True)
        self.assertEqual('valid', diagnostic['record_status'])
        self.assertEqual('fail', diagnostic['quality_status'])
        self.assertTrue(any('native table cell differs' in error for error in diagnostic['quality_findings']))

    def test_table_content_does_not_whitelist_unbound_textbox(self):
        path = self.table_deck()
        deck = Presentation(path)
        deck.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = 'North'
        deck.save(path)
        self.assertTrue(any('extra visible text' in error for error in compare(self.package, self.plan, path)))

    def test_missing_table_cell_is_detected(self):
        path = self.table_deck()
        deck = Presentation(path)
        table = next(shape.table for shape in deck.slides[0].shapes if shape.name == 'ART::DATA')
        table._tbl.remove(table._tbl.tr_lst[1])
        deck.save(path)
        self.assertTrue(any('native table cell differs' in error for error in compare(self.package, self.plan, path)))

    def test_bad_pptx_is_integrity_failure_without_quality_verdict(self):
        path = self.folder / 'bad.pptx'
        path.write_bytes(b'not a zip file')
        diagnostic = result('comparison', compare(self.package, self.plan, path), quality_checked=True)
        self.assertEqual('invalid', diagnostic['record_status'])
        self.assertEqual('unverified', diagnostic['quality_status'])

    def test_cli_preserves_tool_exception_and_missing_quality_verdict(self):
        package = self.write_json('package.json', self.package)
        plan = self.write_json('plan.json', self.plan)
        output = self.folder / 'result.json'
        with patch('sys.argv', ['compare', str(package), str(plan), 'unused.pptx', '--result-json', str(output)]):
            with patch('compare_package_to_pptx.compare', side_effect=AttributeError('synthetic validator defect')):
                self.assertEqual(2, compare_main())
        diagnostic = json.loads(output.read_text())
        self.assertEqual('error', diagnostic['execution_status'])
        self.assertEqual('unverified', diagnostic['quality_status'])
        self.assertIn('AttributeError', diagnostic['tool_errors'][0])

    def test_real_comparator_cli_writes_scoped_result(self):
        path = self.table_deck()
        package = self.write_json('package.json', self.package)
        plan = self.write_json('plan.json', self.plan)
        output = self.folder / 'result.json'
        command = [sys.executable, str(ROOT / 'packages/validators/compare_package_to_pptx.py'),
                   str(package), str(plan), str(path), '--result-json', str(output)]
        run = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(0, run.returncode, run.stdout + run.stderr)
        diagnostic = json.loads(output.read_text())
        self.assertEqual(('completed', 'valid', 'pass'),
                         tuple(diagnostic[key] for key in ('execution_status', 'record_status', 'quality_status')))
        run = subprocess.run(command + ['--allow-extra-text'], capture_output=True, text=True)
        self.assertEqual(0, run.returncode, run.stdout + run.stderr)
        self.assertEqual('unverified', json.loads(output.read_text())['quality_status'])

    def test_deviation_log_accepts_clean_copy_without_logic_pages(self):
        package = self.write_json('package.json', self.package)
        plan = self.write_json('plan.json', self.plan)
        log = {
            'contract_version': '1.1', 'artifact_role': 'subordinate-operational-evidence',
            'package_id': self.package['package_id'], 'package_version': self.package['version'],
            'art_direction_plan_contract_version': '2.2',
            'source_bindings': {'package_sha256': sha256(package), 'art_direction_plan_sha256': sha256(plan), 'pptx_sha256': 'a' * 64},
            'environment_precedence': 'written-pptx-and-render-over-in-memory-state',
            'scoring_policy': 'evidence-not-score', 'deviations': [], 'challenges': [], 'final_status': 'known-risk',
            'cycles': [{'cycle_id': 'FINAL', 'phase': 'final-reopen', 'trigger': 'final-verification',
                        'repair_of': None, 'scope': {'slide_ids': ['S01'], 'target_ids': []}, 'actions': [],
                        'observation': {'machine_evidence': ['Synthetic native inventory'],
                                        'render_evidence': ['Synthetic deferred render record'],
                                        'model_interpretation': 'Synthetic log structure only; not real render acceptance.',
                                        'affected_slides': ['S01']},
                        'decision': {'outcome': 'continue-known-risk', 'owner_layer': 'output-qa',
                                     'reason': 'Synthetic structure fixture with deferred render evidence.'}}],
        }
        errors = validate_deviation(self.package, self.plan, log, package, plan, None)
        self.assertEqual([], errors)
        self.assertEqual('unverified', result('deviation', errors)['quality_status'])
        log['cycles'][0]['scope']['slide_ids'] = ['UNKNOWN']
        self.assertTrue(any('unknown package slide' in error for error in validate_deviation(self.package, self.plan, log, package, plan, None)))

    def qa_errors(self, state, remaining=None):
        _, package, plan, qa = synthetic_handoff(self.folder)
        qa['slides'] = [{'slide_id': package['logic_layer']['slides'][0]['slide_id'],
                         'checks': {key: 'pass' for key in CHECK_KEYS},
                         'not_applicable_reasons': {}, 'issues_remaining': remaining or []}]
        qa['slides'][0]['checks']['visual_hierarchy'] = state
        # Deliberately partial evidence isolates status interpretation in the real
        # validator; unrelated missing-evidence errors must remain present.
        return validate_qa(package, plan, qa)

    def test_honest_qa_failure_is_quality_not_unknown_status(self):
        errors = self.qa_errors('fail', ['Synthetic content overlap'])
        diagnostic = result('qa', errors, quality_checked=True)
        self.assertTrue(any('visual_hierarchy' in error for error in diagnostic['quality_findings']))
        self.assertTrue(any('issues_remaining' in error for error in diagnostic['quality_findings']))
        self.assertFalse(any('visual_hierarchy' in error or 'issues_remaining' in error for error in diagnostic['record_errors']))

    def test_deferred_qa_has_no_quality_pass_or_failure(self):
        diagnostic = result('qa', self.qa_errors('deferred'), quality_checked=True)
        self.assertTrue(any('visual_hierarchy' in error for error in diagnostic['coverage_gaps']))
        self.assertEqual('unverified', diagnostic['quality_status'])

    def test_unknown_qa_status_is_still_a_record_error(self):
        diagnostic = result('qa', self.qa_errors('looks-good'), quality_checked=True)
        self.assertTrue(any('visual_hierarchy' in error for error in diagnostic['record_errors']))

    def test_record_validation_cannot_imply_quality_pass(self):
        self.assertEqual('unverified', result('record-only', [])['quality_status'])
        diagnostic = result('mixed', [Issue('Actual numeric mismatch', 'quality'), Issue('Renderer unavailable', 'tool')], quality_checked=True)
        self.assertEqual('error', diagnostic['execution_status'])
        self.assertEqual('fail', diagnostic['quality_status'])
        self.assertEqual(['Actual numeric mismatch'], diagnostic['quality_findings'])


if __name__ == '__main__':
    unittest.main()
