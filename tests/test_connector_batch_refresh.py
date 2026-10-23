import io
import unittest

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from packages.layout.native_connector import (
    add_rectangular_connector, refresh_rectangular_connector,
    refresh_rectangular_connectors,
)


class ConnectorBatchRefreshTest(unittest.TestCase):
    def fixture(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        nodes = []
        for x, y, label in [(1, 1, 'A'), (4, 2, 'B'), (7, 1, 'C')]:
            node = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                           Inches(2), Inches(1))
            node.text = label
            nodes.append(node)
        a, b, c = nodes
        lines = [add_rectangular_connector(slide, begin=(a, 'right'), end=(b, 'left'),
                                           tail={'type': 'triangle', 'w': 'med'}),
                 add_rectangular_connector(slide, begin=(b, 'right'), end=(c, 'left'),
                                           tail={'type': 'triangle', 'len': 'lg'})]
        annotation = slide.shapes.add_textbox(Inches(3), Inches(1), Inches(1), Inches(.5))
        annotation.text = 'Independent label'
        return presentation, slide, nodes, lines, annotation

    def line_style(self, line):
        return etree.tostring(line._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln'))

    def registered(self, line):
        return tuple(line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:stCxn/@id')
                     + line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:stCxn/@idx')
                     + line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn/@id')
                     + line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn/@idx'))

    def assert_current_paths(self, nodes, lines):
        a, b, c = nodes
        self.assertEqual((lines[0].begin_x, lines[0].begin_y), (a.left+a.width, a.top+a.height//2))
        self.assertEqual((lines[0].end_x, lines[0].end_y), (b.left, b.top+b.height//2))
        self.assertEqual((lines[1].begin_x, lines[1].begin_y), (b.left+b.width, b.top+b.height//2))
        self.assertEqual((lines[1].end_x, lines[1].end_y), (c.left, c.top+c.height//2))

    def test_shared_node_move_refreshes_selected_paths_and_preserves_other_objects(self):
        presentation, slide, nodes, lines, annotation = self.fixture()
        styles = [self.line_style(line) for line in lines]
        registrations = [self.registered(line) for line in lines]
        text_bodies = [node._element.txBody.xml for node in nodes]
        annotation_xml = annotation._element.xml
        count = len(slide.shapes)
        old_first_end = (lines[0].end_x, lines[0].end_y)
        nodes[1].left = Inches(8)
        nodes[1].top += Inches(.5)
        nodes[1].width += Inches(.5)
        self.assertEqual((lines[0].end_x, lines[0].end_y), old_first_end)
        self.assertEqual(refresh_rectangular_connectors(lines), tuple(lines))
        self.assert_current_paths(nodes, lines)
        self.assertEqual([self.line_style(line) for line in lines], styles)
        self.assertEqual([self.registered(line) for line in lines], registrations)
        self.assertEqual([node._element.txBody.xml for node in nodes], text_bodies)
        self.assertEqual(annotation._element.xml, annotation_xml)
        self.assertEqual(len(slide.shapes), count)

    def test_late_bad_registration_rejects_before_any_selected_path_changes(self):
        for attribute, value in [('id', '99999'), ('idx', '99')]:
            with self.subTest(attribute=attribute):
                presentation, slide, nodes, lines, annotation = self.fixture()
                nodes[1].left += Inches(1)
                lines[1]._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn')[0].set(attribute, value)
                before = [line._element.xml for line in lines]
                with self.assertRaises(ValueError):
                    refresh_rectangular_connectors(lines)
                self.assertEqual([line._element.xml for line in lines], before)

    def test_duplicate_cross_slide_and_nonexplicit_batches_do_not_partially_update(self):
        presentation, slide, nodes, lines, annotation = self.fixture()
        nodes[1].left += Inches(1)
        other = presentation.slides.add_slide(presentation.slide_layouts[6])
        a = other.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
        b = other.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), 0, Inches(1), Inches(1))
        external = add_rectangular_connector(other, begin=(a, 'right'), end=(b, 'left'))
        before = [line._element.xml for line in lines]
        for batch in [[lines[0], lines[0]], [lines[0], external], [], iter(lines), [lines[0], nodes[1]]]:
            with self.subTest(batch_type=type(batch).__name__):
                with self.assertRaises(ValueError):
                    refresh_rectangular_connectors(batch)
                self.assertEqual([line._element.xml for line in lines], before)

    def test_ambiguous_native_target_ids_reject_before_any_path_update(self):
        presentation, slide, nodes, lines, annotation = self.fixture()
        nodes[1].left += Inches(1)
        nodes[2]._element.nvSpPr.cNvPr.set('id', str(nodes[0].shape_id))
        before = [line._element.xml for line in lines]
        with self.assertRaisesRegex(ValueError, 'Ambiguous duplicate native shape IDs'):
            refresh_rectangular_connectors(lines)
        self.assertEqual([line._element.xml for line in lines], before)

    def test_single_refresh_and_batch_native_reopen_preserve_ids_sites_and_endings(self):
        presentation, slide, nodes, lines, annotation = self.fixture()
        nodes[1].left -= Inches(2)
        self.assertIs(refresh_rectangular_connector(lines[0]), lines[0])
        refresh_rectangular_connectors(lines)
        self.assert_current_paths(nodes, lines)
        registrations = [self.registered(line) for line in lines]
        styles = [self.line_style(line) for line in lines]
        memory = io.BytesIO()
        presentation.save(memory)
        current = Presentation(io.BytesIO(memory.getvalue())).slides[0]
        shapes = list(current.shapes)
        new_nodes, new_lines = shapes[:3], shapes[3:5]
        self.assert_current_paths(new_nodes, new_lines)
        self.assertEqual([self.registered(line) for line in new_lines], registrations)
        self.assertEqual([self.line_style(line) for line in new_lines], styles)
        self.assertEqual(shapes[5].text, 'Independent label')


if __name__ == '__main__':
    unittest.main()
