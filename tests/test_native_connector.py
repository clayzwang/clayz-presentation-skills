import unittest

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from packages.layout.native_connector import attach_rectangular_connector


class NativeConnectorTest(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.a = self.rect(self.slide, 1, 1)
        self.b = self.rect(self.slide, 5, 3)
        self.line = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(0), Inches(0), Inches(2), Inches(2))

    def rect(self, slide, x, y):
        return slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(2), Inches(1))

    def rejected_unchanged(self, **kwargs):
        before = self.line._element.xml
        with self.assertRaises(ValueError):
            attach_rectangular_connector(self.line, **kwargs)
        self.assertEqual(self.line._element.xml, before)

    def test_both_ids_sites_and_coordinates_match_selected_sides(self):
        n = len(self.slide.shapes)
        self.assertIs(attach_rectangular_connector(self.line, begin=(self.a, 'right'), end=(self.b, 'left')), self.line)
        self.assertEqual(self.line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:stCxn/@id'), [str(self.a.shape_id)])
        self.assertEqual(self.line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:stCxn/@idx'), ['3'])
        self.assertEqual(self.line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn/@id'), [str(self.b.shape_id)])
        self.assertEqual(self.line._element.xpath('./p:nvCxnSpPr/p:cNvCxnSpPr/a:endCxn/@idx'), ['1'])
        self.assertEqual((self.line.begin_x, self.line.begin_y), (Inches(3), Inches(1.5)))
        self.assertEqual((self.line.end_x, self.line.end_y), (Inches(5), Inches(3.5)))
        self.assertEqual(len(self.slide.shapes), n)

    def test_single_end_preserves_other_end_and_no_implicit_selection(self):
        old_end = (self.line.end_x, self.line.end_y)
        attach_rectangular_connector(self.line, begin=(self.a, 'bottom'))
        self.assertEqual((self.line.end_x, self.line.end_y), old_end)
        self.assertFalse(self.line._element.xpath('.//a:endCxn'))
        self.rejected_unchanged()

    def test_all_sites_and_reverse_geometry(self):
        expected = {'top': (2, 1), 'left': (1, 1.5), 'bottom': (2, 2), 'right': (3, 1.5)}
        for site, (x, y) in expected.items():
            attach_rectangular_connector(self.line, begin=(self.b, 'left'), end=(self.a, site))
            self.assertEqual((self.line.end_x, self.line.end_y), (Inches(x), Inches(y)))

    def test_invalid_second_selection_does_not_modify_valid_first(self):
        for site in ['unknown', 4, -1, True]:
            self.rejected_unchanged(begin=(self.a, 'right'), end=(self.b, site))

    def test_wrong_target_kinds_and_zero_extent_leave_existing_binding(self):
        attach_rectangular_connector(self.line, begin=(self.a, 'bottom'))
        ellipse = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, Inches(1), Inches(1))
        textbox = self.slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        empty = self.rect(self.slide, 8, 1); empty.width = 0
        for bad in [ellipse, textbox, self.line, empty]:
            self.rejected_unchanged(begin=(self.a, 'right'), end=(bad, 'left'))

    def test_cross_slide_same_numeric_id_is_rejected_before_any_change(self):
        other = self.prs.slides.add_slide(self.prs.slide_layouts[6]); target = self.rect(other, 1, 1)
        self.assertEqual(target.shape_id, self.a.shape_id)
        self.rejected_unchanged(begin=(self.a, 'right'), end=(target, 'left'))

    def test_group_container_even_same_slide_part_is_rejected(self):
        group = self.slide.shapes.add_group_shape(); target = self.rect(group, 1, 1)
        self.assertIs(target.part, self.line.part)
        self.rejected_unchanged(begin=(self.a, 'right'), end=(target, 'left'))
        grouped_line = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(1), Inches(1))
        before = grouped_line._element.xml
        with self.assertRaises(ValueError):
            attach_rectangular_connector(grouped_line, begin=(target, 'right'))
        self.assertEqual(grouped_line._element.xml, before)

    def test_rotation_and_both_target_flips_are_rejected(self):
        self.b.rotation = 20
        self.rejected_unchanged(begin=(self.a, 'right'), end=(self.b, 'left'))
        self.b.rotation = 0
        for flag in ['flipH', 'flipV']:
            xfrm = self.b._element.spPr.xfrm; setattr(xfrm, flag, True)
            self.rejected_unchanged(begin=(self.a, 'right'), end=(self.b, 'left'))
            setattr(xfrm, flag, False)

    def test_explicit_refresh_after_programmatic_move_no_auto_layout_claim(self):
        attach_rectangular_connector(self.line, end=(self.b, 'left'))
        old = self.line.end_x; self.b.left += Inches(1)
        self.assertEqual(self.line.end_x, old)
        attach_rectangular_connector(self.line, end=(self.b, 'left'))
        self.assertEqual(self.line.end_x, self.b.left)

    def test_not_a_connector_and_curved_or_rotated_connector_rejected(self):
        with self.assertRaises(ValueError):
            attach_rectangular_connector(self.a, begin=(self.b, 'left'))
        curved = self.slide.shapes.add_connector(MSO_CONNECTOR.CURVE, 0, 0, Inches(2), Inches(1))
        with self.assertRaises(ValueError):
            attach_rectangular_connector(curved, begin=(self.a, 'top'))
        self.line.rotation = 15
        self.rejected_unchanged(begin=(self.a, 'right'))


if __name__ == '__main__':
    unittest.main()
