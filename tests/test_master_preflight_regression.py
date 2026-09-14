# SPDX-License-Identifier: Apache-2.0
"""Regression coverage for calibrated preflight of user supplied masters.

The fixtures in this module are deliberately local and synthetic.  The CLI
test still exercises the real configure-task, run-challenge, and preflight
scripts against the checked out 0.15.0-shaped bundle.
"""

from __future__ import annotations

import copy
import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import ExitStack
from datetime import datetime, timezone
from pathlib import Path
from pathlib import PurePosixPath
from unittest import mock


ROOT = Path(__file__).resolve().parents[1]
DEPENDENCY_ROOT = ROOT.parents[1] / "library-iteration-deps"
MASTER_CAPABILITIES = (
    "master-preservation",
    "layout-inheritance",
    "east-asian-font-name",
)
HOST_ATTESTATION_CONTRACT = "io.clayz.presentation.host-capability-attestation/1.0"
HOST_INVENTORY_CONTRACT = "io.clayz.presentation.host-tool-inventory/1.0"

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
VALIDATORS = ROOT / "packages" / "validators"
if str(VALIDATORS) not in sys.path:
    sys.path.insert(0, str(VALIDATORS))

from packages.runtime import preflight  # noqa: E402
from packages.validators.stage_work_records import create_record  # noqa: E402
from validate_supervision_report import validate_environment_observation  # noqa: E402
from scripts import runtime_preflight as runtime_preflight_cli  # noqa: E402


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _canonical(value: object) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8")


def _write_json(path: Path, value: object) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(_canonical(value))
    return path


def _test_environment() -> dict[str, str]:
    environment = dict(os.environ)
    environment.update({"PYTHONUTF8": "1", "PYTHONDONTWRITEBYTECODE": "1"})
    existing = environment.get("PYTHONPATH")
    paths = [str(DEPENDENCY_ROOT)]
    if existing:
        paths.append(existing)
    environment["PYTHONPATH"] = os.pathsep.join(paths)
    return environment


def _write_minimal_zip_master(path: Path) -> None:
    """Keep the CLI fixture readable even when python-pptx is unavailable."""

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'
        '</Types>'
    ).encode("utf-8")
    presentation = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'
    ).encode("utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("ppt/presentation.xml", presentation)


def _write_synthetic_master(path: Path) -> str:
    """Create a valid local template and return its byte hash.

    When the bundled runtime has python-pptx, the fixture contains a real
    slide with Chinese text and a requested CJK font.  The fallback remains a
    valid ZIP package so tests that simulate a missing author can still bind a
    concrete readable file.
    """

    path.parent.mkdir(parents=True, exist_ok=True)
    if importlib.util.find_spec("pptx") is None:
        _write_minimal_zip_master(path)
    else:
        from pptx import Presentation  # type: ignore[import-not-found]
        from pptx.util import Inches  # type: ignore[import-not-found]

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(8.0), Inches(1.0))
        paragraph = textbox.text_frame.paragraphs[0]
        run = paragraph.add_run()
        run.text = "合成母版"
        run.font.name = "华文楷体"
        body = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(8.0), Inches(0.8))
        body.text = "合成模板字体与版式目标"
        presentation.save(path)

    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise AssertionError("synthetic master contains a corrupt ZIP member")
        if "ppt/presentation.xml" not in archive.namelist():
            raise AssertionError("synthetic master is missing ppt/presentation.xml")
    if importlib.util.find_spec("pptx") is not None:
        from pptx import Presentation  # type: ignore[import-not-found]

        opened = Presentation(path)
        if len(opened.slides) < 1:
            raise AssertionError("synthetic master has no readable slide")
    return _sha256(path)


def _master_config(master_path: Path, *, extra: tuple[str, ...] = ()) -> dict[str, object]:
    config = json.loads((ROOT / "config" / "default.json").read_text(encoding="utf-8"))
    config = copy.deepcopy(config)
    theme = config["theme"]
    theme["profile"] = "synthetic-master"
    theme["source"] = "user-master"
    theme["master_path"] = str(master_path.resolve())
    theme["typography"]["primary_fonts"] = ["华文楷体"]
    theme["typography"]["cjk_fonts"] = ["华文楷体"]
    required = list(config["renderer"]["required_capabilities"])
    for capability in (*MASTER_CAPABILITIES, *extra):
        if capability not in required:
            required.append(capability)
    config["renderer"]["required_capabilities"] = required
    return config


def _component_gate(config: dict[str, object]) -> dict[str, object]:
    version = str(config["identity"]["version"])
    return {
        "artifact": "component-version-report.json",
        "sha256": "a" * 64,
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "status": "latest",
        "local_release_version": version,
        "latest_release_version": version,
        "manifest_sha256": "b" * 64,
        "all_components_current": True,
    }


def _bound_run(root: Path) -> dict[str, object]:
    root.mkdir(parents=True, exist_ok=True)
    request = root / "task-request.txt"
    request.write_bytes(b"synthetic master preflight request\n")
    challenge = preflight.issue_run_challenge(request.read_bytes(), task_root=root)
    challenge_path = root / "run-challenge.json"
    challenge_raw = (json.dumps(challenge, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    challenge_path.write_bytes(challenge_raw)
    challenge_sha256 = hashlib.sha256(challenge_raw).hexdigest()

    issuance = {
        "contract": preflight.RUN_CHALLENGE_ISSUANCE_CONTRACT,
        "challenge_sha256": challenge_sha256,
        "run_id": challenge["run_id"],
        "task_request_sha256": challenge["task_request_sha256"],
        "nonce": challenge["nonce"],
        "task_root_sha256": challenge["task_root_sha256"],
        "issued_at": challenge["issued_at"],
        "expires_at": challenge["expires_at"],
    }
    issuance_path = root / Path(*PurePosixPath(str(challenge["issuance_record"])).parts)
    issuance_raw = (json.dumps(issuance, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    issuance_path.parent.mkdir(parents=True, exist_ok=True)
    issuance_path.write_bytes(issuance_raw)
    issuance_receipt = {
        **issuance,
        "receipt_path": issuance_path.resolve().as_posix(),
        "receipt_sha256": hashlib.sha256(issuance_raw).hexdigest(),
    }

    consumption = {
        "contract": preflight.RUN_CHALLENGE_CONSUMPTION_CONTRACT,
        "challenge_sha256": challenge_sha256,
        "run_id": challenge["run_id"],
        "task_request_sha256": challenge["task_request_sha256"],
        "nonce": challenge["nonce"],
        "task_root_sha256": challenge["task_root_sha256"],
        "consumed_at": challenge["issued_at"],
    }
    consumption_path = root / ".clayz-run-challenges" / "consumed" / f"{challenge_sha256}.json"
    consumption_raw = (json.dumps(consumption, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    consumption_path.parent.mkdir(parents=True, exist_ok=True)
    consumption_path.write_bytes(consumption_raw)
    consumption_receipt = {
        **consumption,
        "receipt_path": consumption_path.resolve().as_posix(),
        "receipt_sha256": hashlib.sha256(consumption_raw).hexdigest(),
    }
    return {
        "root": root,
        "request": request,
        "challenge": challenge,
        "challenge_sha256": challenge_sha256,
        "issuance": issuance_receipt,
        "consumption": consumption_receipt,
    }


def _bound_host(root: Path, capabilities: list[str]) -> dict[str, object]:
    bundle = _bound_run(root)
    challenge = bundle["challenge"]
    challenge_sha256 = str(bundle["challenge_sha256"])
    inventory = {
        "contract": HOST_INVENTORY_CONTRACT,
        "run_id": challenge["run_id"],
        "task_request_sha256": challenge["task_request_sha256"],
        "nonce": challenge["nonce"],
        "challenge_sha256": challenge_sha256,
        "capabilities": capabilities,
    }
    inventory_path = root / "host-tool-inventory.json"
    inventory_raw = (json.dumps(inventory, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    inventory_path.write_bytes(inventory_raw)
    host = {
        "contract": HOST_ATTESTATION_CONTRACT,
        "host": "synthetic-host",
        "available": True,
        "capabilities": capabilities,
        "run_id": challenge["run_id"],
        "task_request_sha256": challenge["task_request_sha256"],
        "nonce": challenge["nonce"],
        "challenge_sha256": challenge_sha256,
        "source": "host-inspection",
        "observed_at": challenge["issued_at"],
        "evidence_receipts": [{
            "artifact": inventory_path.name,
            "sha256": hashlib.sha256(inventory_raw).hexdigest(),
            "contract": HOST_INVENTORY_CONTRACT,
        }],
    }
    host_path = root / "host-capabilities.json"
    _write_json(host_path, host)
    context = runtime_preflight_cli._validate_host_evidence(host, host_path, challenge, challenge_sha256)
    if context is None:
        raise AssertionError("synthetic host evidence was not validated")
    return {**bundle, "host": host, "host_context": context}


def _build_report(
    config: dict[str, object],
    *,
    python_pptx: bool,
    renderer: bool,
    powerpoint: bool = False,
    binding: dict[str, object] | None = None,
    host_bundle: dict[str, object] | None = None,
) -> dict[str, object]:
    """Run the real report builder while controlling only host observations."""

    modules = {"pptx": python_pptx, "PIL": True, "yaml": True}

    def module_available(name: str) -> bool:
        return modules.get(name, False)

    def command(*names: str) -> str | None:
        if renderer and any(name in {"soffice", "libreoffice"} for name in names):
            return "C:/synthetic/soffice.exe"
        return None

    if host_bundle is not None:
        binding = host_bundle
    kwargs: dict[str, object] = {
        "model_profile": "A",
        "component_version_gate": _component_gate(config),
    }
    if binding is not None:
        kwargs.update({
            "run_challenge": binding["challenge"],
            "run_challenge_sha256": binding["challenge_sha256"],
            "run_challenge_issuance": binding["issuance"],
            "run_challenge_consumption": binding["consumption"],
        })
    if host_bundle is not None:
        kwargs.update({
            "host_capabilities": host_bundle["host"],
            "host_attestation_context": host_bundle["host_context"],
        })

    with ExitStack() as stack:
        stack.enter_context(mock.patch.object(preflight, "_module", side_effect=module_available))
        stack.enter_context(mock.patch.object(
            preflight,
            "_artifact_tool",
            return_value={"available": False, "node": None, "node_modules": None, "bin_dir": None, "package": None},
        ))
        stack.enter_context(mock.patch.object(
            preflight,
            "_windows_powerpoint",
            return_value={"available": powerpoint, "evidence": "synthetic" if powerpoint else None},
        ))
        stack.enter_context(mock.patch.object(preflight, "_command", side_effect=command))
        return preflight.build_preflight_report(config, **kwargs)


def _capability_mirror(report: dict[str, object]) -> dict[str, list[str]]:
    selected = report["selected_route"]
    configured = sorted(report["required_capabilities"])
    missing = sorted(selected.get("missing_capabilities", []))
    declared_unverified = (
        sorted(set(configured) - set(missing))
        if selected.get("assurance_level") == "host-declared-unverified"
        else []
    )
    return {
        "configured": configured,
        "satisfied": [] if declared_unverified else sorted(set(configured) - set(missing)),
        "declared_unverified": declared_unverified,
        "missing": missing,
    }


def _environment_observation(report: dict[str, object], preflight_path: Path) -> dict[str, object]:
    binding = report["run_binding"]
    selected = report["selected_route"]
    targets = []
    for item in report["target_application_checks"]:
        final_status = "deferred" if item["availability"] == "unavailable" else "not-selected"
        targets.append({
            "application": item["application"],
            "capability": item["capability"],
            "availability": item["availability"],
            "final_status": final_status,
            "authoring_gate": False,
            "evidence_refs": [f"runtime-preflight.json#target_application_checks.{item['application']}"],
        })
    status = "ready" if selected["available"] else "provisional" if selected["attemptable"] else "blocked"
    preflight_record = {
        "artifact": str(preflight_path.resolve()),
        "scan_id": report["scan_id"],
        "sha256": _sha256(preflight_path),
        "run_id": binding["run_id"],
        "task_request_sha256": binding["task_request_sha256"],
        "config_sha256": report["config_binding"]["sha256"],
        "nonce": binding["nonce"],
        "challenge_sha256": binding["challenge_sha256"],
        "task_root_sha256": binding["task_root_sha256"],
        "issued_at": binding["issued_at"],
        "expires_at": binding["expires_at"],
        "issuance_receipt_sha256": binding["issuance_receipt_sha256"],
        "consumption_receipt_sha256": binding["consumption_receipt_sha256"],
    }
    return {
        "preflight": preflight_record,
        "route": {
            "route_id": selected["route_id"],
            "authoring_backend": selected["authoring_backend"],
            "render_backend": selected["render_backend"],
            "status": status,
        },
        "required_capabilities": _capability_mirror(report),
        "target_applications": targets,
        "compatibility_scope": "none",
        "attribution_summary": "Synthetic preflight has authoring evidence but deferred quality and target acceptance.",
    }


class MasterPreflightRegressionTests(unittest.TestCase):
    def _temporary_root(self, prefix: str) -> tempfile.TemporaryDirectory[str]:
        fixture_root = ROOT / ".tmp"
        fixture_root.mkdir(parents=True, exist_ok=True)
        return tempfile.TemporaryDirectory(prefix=f"{prefix}-", dir=fixture_root)

    def _run_cli(self, script: Path, *args: object, cwd: Path, expected: int = 0) -> subprocess.CompletedProcess[str]:
        result = subprocess.run(
            [sys.executable, "-B", str(script), *map(str, args)],
            cwd=cwd,
            env=_test_environment(),
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=60,
        )
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        if expected == 0:
            self.assertNotIn("Traceback", result.stderr)
        return result

    def _write_component_report(self, path: Path, config: dict[str, object]) -> None:
        version = str(config["identity"]["version"])
        _write_json(
            path,
            {
                "contract": "io.clayz.presentation.component-version-report/1.0",
                "artifact": path.name,
                "sha256": "c" * 64,
                "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
                "status": "latest",
                "local_release_version": version,
                "latest_release_version": version,
                "latest_release": {"version": version},
                "manifest_sha256": "d" * 64,
                "all_components_current": True,
                "error_codes": [],
                "components": [{"component_id": "synthetic-runtime", "status": "current"}],
            },
        )

    def test_real_configure_challenge_preflight_binds_readable_master_and_keeps_quality_missing(self) -> None:
        with tempfile.TemporaryDirectory(prefix="luna-cli-master-", dir=ROOT.parents[1]) as directory:
            root = Path(directory)
            master = root / "synthetic-master.pptx"
            master_sha256 = _write_synthetic_master(master)
            personal = root / "personal-config.json"
            _write_json(
                personal,
                {
                    "config": {
                        "theme": {
                            "profile": "synthetic-master",
                            "source": "user-master",
                            "master_path": str(master.resolve()),
                            "typography": {"primary_fonts": ["华文楷体"], "cjk_fonts": ["华文楷体"]},
                        }
                    },
                    "preferences": {"fixture": "local-synthetic-master"},
                },
            )
            configured = root / "configured"
            self._run_cli(
                ROOT / "scripts" / "cloud_learning_cli.py",
                "configure-task",
                "--personal-config",
                personal,
                "--without-library",
                "--output-dir",
                configured,
                cwd=ROOT,
            )
            config_path = configured / "task-config.json"
            selection_path = configured / "task-selection.json"
            config = json.loads(config_path.read_text(encoding="utf-8"))
            selection = json.loads(selection_path.read_text(encoding="utf-8"))
            self.assertEqual(config["theme"]["source"], "user-master")
            self.assertEqual(Path(config["theme"]["master_path"]).resolve(), master.resolve())
            self.assertEqual(config["theme"]["typography"]["primary_fonts"], ["华文楷体"])
            default_config = json.loads((ROOT / "config" / "default.json").read_text(encoding="utf-8"))
            self.assertTrue(
                set(default_config["renderer"]["required_capabilities"]).issubset(
                    set(config["renderer"]["required_capabilities"])
                )
            )
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(config["renderer"]["required_capabilities"])))
            self.assertEqual(selection["config_sha256"], _sha256(config_path))
            self.assertEqual(master_sha256, _sha256(master))
            with zipfile.ZipFile(master) as archive:
                self.assertIsNone(archive.testzip())

            request = root / "request.txt"
            request.write_bytes(b"synthetic configure to preflight request\n")
            challenge = root / "run-challenge.json"
            self._run_cli(
                ROOT / "scripts" / "runtime_preflight.py",
                "--issue-challenge",
                "--task-request",
                request,
                "--output",
                challenge,
                cwd=root,
            )
            component = root / "component-version-report.json"
            self._write_component_report(component, config)
            output = root / "runtime-preflight.json"
            self._run_cli(
                ROOT / "scripts" / "runtime_preflight.py",
                "--challenge",
                challenge,
                "--task-request",
                request,
                "--task-selection",
                selection_path,
                "--config",
                config_path,
                "--model-profile",
                "A",
                "--component-version-report",
                component,
                "--output",
                output,
                cwd=root,
            )
            report = json.loads(output.read_text(encoding="utf-8"))
            required = set(report["required_capabilities"])
            selected = report["selected_route"]
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(required))
            self.assertEqual(report["config_binding"]["source"], "task-selected")
            self.assertEqual(report["config_binding"]["sha256"], _sha256(config_path))
            if importlib.util.find_spec("pptx") is not None:
                self.assertEqual(selected["authoring_backend"], "python-pptx", report)
                self.assertTrue(selected["attemptable"], report)
                self.assertNotEqual(selected["authoring_backend"], "spec-only")
                self.assertFalse(selected["available"], report)
                self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(selected["missing_capabilities"])), report)
                self.assertFalse(any("emit an internal spec only" in warning for warning in report["warnings"]))
                warning_text = " ".join(report["warnings"]).lower()
                self.assertTrue(any(word in warning_text for word in ("defer", "quality", "attempt", "evidence")), report)
            else:
                self.assertEqual(selected["authoring_backend"], "spec-only", report)
                self.assertFalse(selected["attemptable"], report)

    def test_python_pptx_with_renderer_is_attemptable_but_master_quality_is_still_missing(self) -> None:
        with self._temporary_root("luna-python-render") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            report = _build_report(_master_config(master), python_pptx=True, renderer=True)
            selected = report["selected_route"]
            self.assertEqual(selected["authoring_backend"], "python-pptx", report)
            self.assertEqual(selected["render_backend"], "libreoffice", report)
            self.assertTrue(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(selected["missing_capabilities"])), report)
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(report["required_capabilities"])))
            self.assertFalse(any("emit an internal spec only" in warning for warning in report["warnings"]))

    def test_python_pptx_without_renderer_defers_render_but_keeps_real_writer(self) -> None:
        with self._temporary_root("luna-python-no-render") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            report = _build_report(_master_config(master), python_pptx=True, renderer=False)
            selected = report["selected_route"]
            self.assertEqual(selected["authoring_backend"], "python-pptx", report)
            self.assertEqual(selected["render_backend"], "none", report)
            self.assertTrue(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertIn("render-preview", selected["missing_capabilities"], report)
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(selected["missing_capabilities"])), report)
            warning_text = " ".join(report["warnings"]).lower()
            self.assertIn("deferred", warning_text, report)
            self.assertFalse(any("emit an internal spec only" in warning for warning in report["warnings"]))

    def test_unknown_quality_feature_stays_missing_without_blocking_a_real_writer(self) -> None:
        unknown = "experimental-master-fidelity-v2"
        with self._temporary_root("luna-unknown-quality") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            report = _build_report(_master_config(master, extra=(unknown,)), python_pptx=True, renderer=True)
            selected = report["selected_route"]
            self.assertEqual(selected["authoring_backend"], "python-pptx", report)
            self.assertTrue(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertIn(unknown, report["required_capabilities"])
            self.assertIn(unknown, selected["missing_capabilities"], report)
            self.assertNotIn(unknown, {"editable-text", "editable-shapes", "pptx-inspection"})
            self.assertFalse(any("emit an internal spec only" in warning for warning in report["warnings"]))

    def test_renderer_without_python_author_does_not_create_a_writable_route(self) -> None:
        with self._temporary_root("luna-render-only") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            report = _build_report(_master_config(master), python_pptx=False, renderer=True)
            selected = report["selected_route"]
            self.assertEqual(selected["authoring_backend"], "spec-only", report)
            self.assertFalse(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertIn("editable-text", selected["missing_capabilities"], report)
            self.assertTrue(any(route["authoring_backend"] == "spec-only" for route in report["fallback_routes"]))
            self.assertFalse(any(
                route["authoring_backend"] != "spec-only" and route["attemptable"]
                for route in report["fallback_routes"]
            ))

    def test_no_author_and_no_renderer_remains_blocked_spec_only(self) -> None:
        with self._temporary_root("luna-no-author") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            report = _build_report(_master_config(master), python_pptx=False, renderer=False)
            selected = report["selected_route"]
            self.assertEqual(selected["route_id"], "spec-only+none", report)
            self.assertEqual(selected["authoring_backend"], "spec-only", report)
            self.assertEqual(selected["render_backend"], "none", report)
            self.assertFalse(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertIn("editable-text", selected["missing_capabilities"], report)
            self.assertTrue(any("spec only" in warning.lower() for warning in report["warnings"]))

    def test_complete_challenge_bound_host_is_attemptable_but_remains_unverified(self) -> None:
        with self._temporary_root("luna-host-unverified") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            config = _master_config(master)
            capabilities = list(config["renderer"]["required_capabilities"])
            host_bundle = _bound_host(root / "host", capabilities)
            report = _build_report(config, python_pptx=False, renderer=False, host_bundle=host_bundle)
            selected = report["selected_route"]
            self.assertEqual(selected["authoring_backend"], "native-presentation-tool", report)
            self.assertEqual(selected["render_backend"], "native-presentation-tool", report)
            self.assertTrue(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertEqual(selected["assurance_level"], "host-declared-unverified", report)
            observation = report["dependencies"]["host_tools"]["observation"]
            self.assertEqual(observation["assurance_level"], "host-declared-unverified")
            self.assertFalse(observation["route_eligible"])
            self.assertTrue(observation["evidence_receipts"])
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(report["required_capabilities"])))

    def test_host_renderer_and_host_author_capabilities_do_not_cross_route_boundaries(self) -> None:
        with self._temporary_root("luna-host-roles") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            config = _master_config(master)

            render_only = _bound_host(root / "render-only", ["render-preview", "pptx-inspection"])
            render_report = _build_report(config, python_pptx=False, renderer=False, host_bundle=render_only)
            render_selected = render_report["selected_route"]
            self.assertNotEqual(render_selected["authoring_backend"], "native-presentation-tool", render_report)
            self.assertFalse(render_selected["attemptable"], render_report)
            native_author_routes = [
                route for route in render_report["fallback_routes"]
                if route["authoring_backend"] == "native-presentation-tool"
            ]
            self.assertTrue(native_author_routes)
            self.assertTrue(all("editable-text" in route["missing_capabilities"] for route in native_author_routes))

            author_capabilities = [
                capability for capability in config["renderer"]["required_capabilities"]
                if capability != "render-preview"
            ]
            author_only = _bound_host(root / "author-only", author_capabilities)
            author_report = _build_report(config, python_pptx=False, renderer=False, host_bundle=author_only)
            author_selected = author_report["selected_route"]
            self.assertEqual(author_selected["authoring_backend"], "native-presentation-tool", author_report)
            self.assertEqual(author_selected["render_backend"], "none", author_report)
            self.assertTrue(author_selected["attemptable"], author_report)
            self.assertFalse(author_selected["available"], author_report)

    def test_preflight_missing_mirror_and_quality_limited_record_can_coexist(self) -> None:
        with self._temporary_root("luna-quality-limited") as directory:
            root = Path(directory)
            master = root / "master.pptx"
            _write_synthetic_master(master)
            config = _master_config(master)
            binding = _bound_run(root / "run")
            report = _build_report(config, python_pptx=True, renderer=False, binding=binding)
            selected = report["selected_route"]
            self.assertEqual(selected["authoring_backend"], "python-pptx", report)
            self.assertTrue(selected["attemptable"], report)
            self.assertFalse(selected["available"], report)
            self.assertTrue(set(MASTER_CAPABILITIES).issubset(set(selected["missing_capabilities"])), report)

            preflight_path = root / "runtime-preflight.json"
            _write_json(preflight_path, report)
            observation = {
                "run_id": report["run_binding"]["run_id"],
                "task_request_sha256": report["run_binding"]["task_request_sha256"],
                "supervised_at": report["run_binding"]["issued_at"],
                "artifact_paths": {"runtime_preflight": str(preflight_path.resolve())},
                "environment_observation": _environment_observation(report, preflight_path),
            }
            errors = validate_environment_observation(
                observation,
                runtime_preflight=report,
                runtime_preflight_sha256=_sha256(preflight_path),
                resolved_config=config,
                resolved_config_sha256=report["config_binding"]["sha256"],
            )
            self.assertEqual(errors, [], errors)
            self.assertEqual(
                observation["environment_observation"]["required_capabilities"]["missing"],
                sorted(selected["missing_capabilities"]),
            )

            acceptance = root / "acceptance.json"
            acceptance.write_text("quality-only continuation\n", encoding="utf-8")
            record = create_record(
                "logic",
                {
                    "summary": "Master fidelity evidence remains pending after preflight.",
                    "decisions": ["Continue with a quality-limited authoring attempt."],
                    "checks": [{
                        "name": "master fidelity evidence",
                        "status": "fail",
                        "blocking": False,
                        "evidence_roles": ["preflight"],
                    }],
                    "open_issues": ["Native master fidelity verification is pending until output audit."],
                },
                {"run_binding": report["run_binding"]},
                {"preflight": preflight_path},
                acceptance_contract=acceptance,
            )
            self.assertEqual(record["readiness"], "quality-limited")
            self.assertFalse(record["checks"][0]["blocking"])


if __name__ == "__main__":
    unittest.main()
