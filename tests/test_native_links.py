import io,unittest,zipfile
from pptx import Presentation
from pptx.util import Inches,Pt
from packages.layout.native_links import link_table

class NativeLinksTest(unittest.TestCase):
 def setUp(self):
  self.prs=Presentation();s=self.prs.slides.add_slide(self.prs.slide_layouts[6])
  self.table=s.shapes.add_table(1,1,Inches(1),Inches(1),Inches(5),Inches(1)).table
  r=self.table.cell(0,0).text_frame.paragraphs[0].add_run();r.text='S7、S17\n美光';r.font.name='STKaiti';r.font.size=Pt(23)
 def test_preserves_visible_text_and_type(self):
  old=self.table.cell(0,0).text
  self.assertEqual(link_table(self.table,{'S7':'https://example.org/7','S17':'https://example.org/17','美光':'https://example.org/micron'}),3)
  self.assertEqual(self.table.cell(0,0).text,old)
  for p in self.table.cell(0,0).text_frame.paragraphs:
   for r in p.runs:self.assertEqual(r.font.name,'STKaiti');self.assertEqual(r.font.size,Pt(23))
 def test_identifier_is_not_prefix_match(self):
  self.assertEqual(link_table(self.table,{'S7':'https://example.org/7'}),1)
  urls=[r.hyperlink.address for p in self.table.cell(0,0).text_frame.paragraphs for r in p.runs if r.hyperlink.address]
  self.assertEqual(urls,['https://example.org/7'])
 def test_native_package_carries_real_url(self):
  url='https://example.org/source?a=1&b=2';link_table(self.table,{'S17':url})
  out=io.BytesIO();self.prs.save(out)
  with zipfile.ZipFile(out) as z:self.assertIn('https://example.org/source?a=1&amp;b=2',z.read('ppt/slides/_rels/slide1.xml.rels').decode())

if __name__=='__main__':unittest.main()
