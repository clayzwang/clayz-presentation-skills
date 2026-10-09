"""Real package lookup and native PPTX changes; teaching/reader quality is not simulated proof."""
import copy
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'packages' / 'validators'))
sys.path.insert(0, str(ROOT / 'scripts'))
from packages.validators.art_learning import load_pack, lookup, validate_cognition, verify_cognition_sources, digest
from packages.validators.art_content import presentation_pages, differences, validate_art_content, validate_change_audit
from tests.test_release_0190 import planning_handoff, relock, write_native, planning_reference, validate_plan, compare
from packages.validators import reader_review as rr
from packages.validators.work_report import _art_direction
from scripts.build_runtime_packs import build_learning, include_light
PACK = ROOT / 'learning-packs/art-design-foundations-v1'

def cognition():
    example = json.loads((PACK / 'worked-example.json').read_text())
    steps = []
    for entry in example['steps']:
        step = {k: v for k, v in entry.items() if k != 'knowledge_codes'}
        step['scope'] = 'deck: hypothetical business-mode comparison'
        step['knowledge_refs'] = [{'code': entry['knowledge_codes'][0],
                                  'receipt': lookup(PACK, entry['knowledge_codes'][0]),
                                  'how_applied': entry['actions'][0]}]
        steps.append(step)
    return {'contract': 'io.clayz.presentation.art-cognition/1.0', 'steps': steps}

class ArtLearningTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory(); self.addCleanup(tmp.cleanup); self.root = Path(tmp.name)

    def test_all_codes_have_principles_and_bounded_lookup(self):
        _, nodes = load_pack(PACK)
        self.assertEqual(66, len(nodes))
        self.assertTrue(all(n['principle'].strip() for n in nodes.values()))
        self.assertEqual(4, len(lookup(PACK, 'A06')['nodes']))
        self.assertEqual(1, len(lookup(PACK, 'P.PER.HIERARCHY')['nodes']))
        with self.assertRaises(ValueError): lookup(PACK, 'P.FAKE')

    def test_hash_tampering_and_empty_category_principle_rejected(self):
        target = self.root / 'pack'; shutil.copytree(PACK, target)
        path = target / 'nodes.json'; value = json.loads(path.read_text()); value['nodes'][0]['principle'] = ''
        path.write_text(json.dumps(value))
        with self.assertRaisesRegex(ValueError, 'hash'): load_pack(target)
        import hashlib
        manifest = json.loads((target / 'manifest.json').read_text()); manifest['files']['nodes.json'] = hashlib.sha256(path.read_bytes()).hexdigest()
        (target / 'manifest.json').write_text(json.dumps(manifest))
        with self.assertRaisesRegex(ValueError, 'principle'): load_pack(target)

    def test_conclusions_actions_and_actual_code_use_are_required(self):
        record = cognition()
        self.assertEqual([], verify_cognition_sources(record, [PACK]))
        self.assertTrue(verify_cognition_sources(record, []))
        broken = copy.deepcopy(record); broken['steps'].pop()
        self.assertTrue(validate_cognition(broken))
        broken = copy.deepcopy(record); broken['steps'][0]['actions'] = []
        self.assertTrue(validate_cognition(broken))
        receipt = record['steps'][0]['knowledge_refs'][0]['receipt']
        receipt['nodes'][0]['principle'] = 'Invented quote'; receipt['selection_sha256'] = digest(receipt['nodes'])
        self.assertTrue(verify_cognition_sources(record, [PACK]))

    def test_external_pack_archive_and_light_exclusion(self):
        archive = build_learning(self.root, '0.20.2')
        with zipfile.ZipFile(archive) as z:
            z.extractall(self.root / 'unpacked')
        self.assertEqual(66, len(load_pack(self.root / 'unpacked/art-design-foundations-v1')[1]))
        self.assertFalse(include_light(PACK / 'nodes.json', 'cloud'))
        self.assertFalse(include_light(PACK / 'nodes.json', 'local'))

    def handoff(self):
        origin, package, plan = planning_handoff(self.root)
        pages = copy.deepcopy(package['copy_layer']['slides'])
        pages[0]['copy_units'][0]['text'] = 'A revised, bounded pilot judgment'
        extension = {'contract': 'io.clayz.presentation.art-content/1.0', 'copy_package_sha256': digest(package), 'slides': pages,
                     'changes': [{**row, 'reason': 'Clarify the bounded inference.', 'evidence': 'Synthetic example: original pilot evidence; no real business claim.',
                                  'impact': 'Keep the limitation while changing expression.'} for row in differences(package, pages)]}
        plan.update(art_content=extension, art_cognition=cognition())
        record = copy.deepcopy(plan['page_planning']['content'])
        record.update(art_content=extension, art_cognition=plan['art_cognition'])
        path = self.root / 'revised-planning.json'; path.write_text(json.dumps(record))
        plan['page_planning'] = planning_reference(path); relock(package, plan)
        return package, plan

    def test_original_copy_preserved_actual_art_text_reaches_native_output_and_report(self):
        package, plan = self.handoff(); original = copy.deepcopy(package)
        self.assertEqual([], validate_plan(package, plan))
        output = self.root / 'actual.pptx'
        view = {**package, 'copy_layer': {'slides': presentation_pages(package, plan)}}
        write_native(view, plan, output)
        self.assertEqual([], compare(package, plan, output))
        self.assertEqual(original, package)
        write_native(package, plan, output)
        self.assertTrue(compare(package, plan, output))
        sections = _art_direction({'plan': plan})['plan_sections']
        self.assertEqual(plan['art_cognition'], sections['art_cognition'])
        self.assertEqual(plan['art_content'], sections['art_content'])

    def test_final_reader_freezes_pixels_then_audits_original_copy_and_art_changes(self):
        from tests.test_reader_review import ReaderReviewTests
        helper = ReaderReviewTests(); helper.setUp(); self.addCleanup(helper.doCleanups)
        package = rr.read(helper.package); package['status'] = 'copy-approved'
        helper.package.write_text(json.dumps(package))
        pages = copy.deepcopy(package['copy_layer']['slides'])
        pages[0]['copy_units'][0]['text'] = 'A has a stated deadline; B has timing principles'
        changes = [{**r, 'reason': 'Name the supported difference.', 'evidence': 'Body wording.', 'impact': 'Title becomes specific.'}
                   for r in differences(package, pages)]
        plan = {'art_content': {'contract': 'io.clayz.presentation.art-content/1.0',
                'copy_package_sha256': digest(package), 'slides': pages, 'changes': changes}}
        plan_path = helper.put('plan.json', plan)
        packet_path = helper.root / 'art-final-packet.json'
        packet = rr.prepare_packet(phase='final', package=helper.package, brief=helper.brief, plan=plan_path,
            directory=helper.root/'blind-input', output=packet_path, pptx=helper.pptx, renders=helper.renders)
        payload = rr.validate_packet(packet)
        self.assertNotIn('art_content', json.dumps(payload))
        self.assertNotIn('Name the supported difference', json.dumps(payload))
        first = helper.root/'first.json'
        rr.record_first(packet=packet_path, context=helper.context(packet_path, 'art-final'),
                        response=helper.put('response.json', helper.response()), output=first)
        comparison = {'baseline': rr.ref(helper.package), 'checks': {'accuracy': 'Synthetic content check.',
            'completeness': 'Qualifier remains.', 'reasoning': 'Claim remains bounded.'},
            'verdict': 'pass', 'explanation': 'Synthetic review tests transport only.',
            'art_changes': [{'change_id': c['change_id'], 'assessment': 'justified', 'explanation': 'Names what body supports.',
                             'evidence': 'Original and actual title compared with unchanged body.'} for c in changes]}
        path = helper.root/'review.json'
        value = rr.record_review(first_read=first, dispositions=helper.put('dispositions.json', []),
                    evidence=[helper.package, plan_path], comparison=helper.put('comparison.json', comparison), output=path)
        rr.require_pass(rr.validate_review(value))
        broken = copy.deepcopy(value); broken['comparison']['art_changes'] = []
        with self.assertRaisesRegex(ValueError, 'every Art difference'): rr.validate_review(broken)
        auditor = helper.audit({'final': path}); auditor['source_records'].append({'kind': 'plan', **rr.ref(plan_path)})
        self.assertIn('final', rr.audit_reviews(auditor))

    def test_omitted_difference_stale_baseline_and_snapshot_mismatch_rejected(self):
        package, plan = self.handoff()
        broken = copy.deepcopy(plan); broken['art_content']['changes'] = []
        self.assertTrue(validate_art_content(package, broken))
        broken = copy.deepcopy(plan); broken['art_content']['copy_package_sha256'] = '0'*64
        self.assertTrue(validate_art_content(package, broken))
        plan['art_cognition']['steps'][0]['conclusion'] = 'Different later conclusion'
        relock(package, plan)
        self.assertTrue(validate_plan(package, plan))

    def test_split_page_maps_actual_final_reader_order_without_leaking_rationale(self):
        package, plan = self.handoff()
        pages = plan['art_content']['slides']
        new = copy.deepcopy(pages[0]); new['slide_id'] = 'S02'
        for u in new['copy_units']: u['copy_id'] += '-2'
        for row in new['content_relationships']['body_relations']:
            row['copy_id'] += '-2'; row['heading_ids'] = [i+'-2' for i in row['heading_ids']]
        pages.append(new)
        plan['art_content']['changes'] = [{**row, 'reason': 'Split for readability.', 'evidence': 'Synthetic extra page.', 'impact': 'Adds a page.'} for row in differences(package, pages)]
        self.assertEqual([], validate_art_content(package, plan))
        design = copy.deepcopy(plan['slides'][0]); design['slide_id'] = 'S02'
        design['reading_sequence'] = [i+'-2' for i in design['reading_sequence']]
        for mapping in design['copy_unit_map']: mapping['copy_id'] += '-2'
        plan['slides'].append(design)
        preview = copy.deepcopy(plan['visual_baseline']['slides'][0]); preview['slide_id'] = 'S02'
        for element in preview['elements']: element['copy_ids'] = [i+'-2' for i in element['copy_ids']]
        plan['visual_baseline']['slides'].append(preview)
        record = copy.deepcopy(plan['page_planning']['content'])
        record['slides'].append({**record['slides'][0], 'slide_id': 'S02'})
        record['art_content'] = plan['art_content']
        path = self.root / 'two-page-planning.json'; path.write_text(json.dumps(record))
        plan['page_planning'] = planning_reference(path); relock(package, plan)
        self.assertEqual([], validate_plan(package, plan))
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        for page, preview in zip(pages, plan['visual_baseline']['slides']):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            units = {u['copy_id']: u['text'] for u in page['copy_units']}
            for element in preview['elements']:
                x,y,w,h = element['box']
                shape = slide.shapes.add_textbox(Inches(x*10), Inches(y*7.5), Inches(w*10), Inches(h*7.5))
                shape.name = element['native_name']; shape.text = '\n'.join(units[c] for c in element['copy_ids'])
        output = self.root / 'two-pages.pptx'; deck.save(output)
        self.assertEqual([], compare(package, plan, output))
        self.assertEqual(['S01','S02'], [p['slide_id'] for p in rr.projection(package, 'final', plan)])
        self.assertNotIn('art_content', str(rr.projection(package, 'final', plan)))
        changes = plan['art_content']['changes']
        obs = [{'change_id': c['change_id'], 'assessment': 'justified', 'explanation': 'Synthetic audit transport exercise.', 'evidence': 'Original baseline and actual rendered wording.'} for c in changes]
        self.assertEqual([], validate_change_audit(package, plan, obs, 'pass'))
        self.assertTrue(validate_change_audit(package, plan, [], 'pass'))
        obs[0]['assessment'] = 'unsupported'
        self.assertTrue(validate_change_audit(package, plan, obs, 'pass'))
        self.assertEqual([], validate_change_audit(package, plan, obs, 'return-art'))

if __name__ == '__main__': unittest.main()
