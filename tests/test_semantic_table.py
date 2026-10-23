import tempfile
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from PIL import ImageFont
from packages.layout.semantic_table import Column, data_cell_by_label, plan, write
from packages.validators.visual_evidence import validate

from tests.font_support import test_font_path
FONT = test_font_path()


class TableRegression(unittest.TestCase):
    def spec(self, **kwargs):
        return plan([Column('项目', weight=2), Column('金额', 'number')],
                    [['经营现金流', '182.935'], ['资本开支', '115.948'], ['差额', '66.987']],
                    width=6, font_path=FONT, caption='FY2026；十亿美元',
                    comparison='比较现金产生、现金投资与剩余', total_rows=[2], **kwargs)

    def test_native_roundtrip_preserves_padding_and_numeric_alignment(self):
        spec = self.spec(); prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        write(slide, spec, 1, 1)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'table.pptx'; prs.save(path)
            tbl = Presentation(path).slides[0].shapes[0].table
            self.assertAlmostEqual(tbl.cell(1, 1).margin_left / 914400, .17)
            self.assertEqual(tbl.cell(1, 1).text_frame.paragraphs[0].alignment, PP_ALIGN.RIGHT)
            self.assertTrue(tbl.cell(3, 1).text_frame.paragraphs[0].runs[0].font.bold)
            self.assertFalse(tbl.first_row)
            self.assertGreater(tbl.columns[0].width, tbl.columns[1].width)

    def test_capacity_cannot_pass_by_silent_font_shrink(self):
        with self.assertRaisesRegex(ValueError, 'revise Art'): self.spec(max_height=.5)

    def test_native_width_reserve_reflows_before_capacity_check(self):
        # R4 native reopening split Google Cloud despite a measured single line.
        text = 'AWS、微软Azure、Google Cloud'
        font = ImageFont.truetype(FONT, round(23 * 120 / 72))
        args = dict(columns=[Column('企业样本')], rows=[[text]],
                    width=font.getlength(text)/120*1.04+.34,
                    font_path=FONT, size=23, caption='关键角色样本',
                    comparison='交付与企业对应', protected_terms=['Google Cloud'])
        original = plan(**args)
        calibrated = plan(**args, native_width_scale=1.08)
        self.assertEqual(original['cells'][1][0]['lines'], [text])
        self.assertGreater(calibrated['height'], original['height'])
        self.assertEqual(''.join(calibrated['cells'][1][0]['lines']).replace(' ', ''), text.replace(' ', ''))
        self.assertEqual(calibrated['cells'][1][0]['size_pt'], 23)
        self.assertEqual(calibrated['widths'], original['widths'])
        self.assertEqual(calibrated['padding'], original['padding'])
        with self.assertRaisesRegex(ValueError, 'revise Art'):
            plan(**args, native_width_scale=1.08, max_height=original['height']+.01)

    def test_missing_visual_review_stays_unreviewed(self):
        self.assertTrue(validate({'pages': {'S01': {'content': 'numbers correct'}}}, ['S01']))


class NamedCellSelection(unittest.TestCase):
    def setUp(self):
        self.columns = [Column('工作'), Column('上期', 'number'), Column('本期', 'number')]
        self.rows = [['客户服务', '8', '10'], ['资本经营', '5', '7']]

    def select(self, columns=None, rows=None, **kwargs):
        labels = dict(row_label_column='工作', row_label='资本经营', column_label='本期')
        labels.update(kwargs)
        return data_cell_by_label(self.columns if columns is None else columns,
                                  self.rows if rows is None else rows, **labels)

    def test_reordered_data_keeps_explicit_target_in_native_output(self):
        # Select the smaller value deliberately: importance is the author's.
        self.assertEqual(self.select(), (1, 2))
        columns = [self.columns[i] for i in (2, 0, 1)]
        rows = [[row[i] for i in (2, 0, 1)] for row in reversed(self.rows)]
        cell = self.select(columns, rows)
        self.assertEqual(cell, (0, 0))  # data row, not native header-inclusive row
        spec = plan(columns, rows, width=6, font_path=FONT, emphasis=[cell],
                    caption='合成测试数据', comparison='作者选定资本经营本期值')
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        write(slide, spec, 1, 1)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'selected-cell.pptx'; prs.save(path)
            table = Presentation(path).slides[0].shapes[0].table
            selected = table.cell(1, 0).text_frame.paragraphs[0].runs[0]
            other = table.cell(2, 0).text_frame.paragraphs[0].runs[0]
            self.assertEqual(selected.text, '7')
            self.assertTrue(selected.font.bold)
            self.assertEqual(str(selected.font.color.rgb), '185D72')
            self.assertEqual(other.text, '10')
            self.assertFalse(other.font.bold)
            self.assertEqual(str(other.font.color.rgb), '172B3A')

    def test_missing_or_ambiguous_selected_column_is_rejected(self):
        for labels in [('工作', '上期', '上期'), ('工作', '本期', '本期'),
                       ('工作', '工作', '本期')]:
            with self.subTest(labels=labels), self.assertRaisesRegex(ValueError, 'column label'):
                self.select([Column(label) for label in labels])

    def test_missing_or_ambiguous_selected_row_is_rejected(self):
        for rows in ([self.rows[0]], [self.rows[1], self.rows[1]]):
            with self.subTest(rows=rows), self.assertRaisesRegex(ValueError, 'row label'):
                self.select(rows=rows)
        # Unrelated repeated row labels do not force a table-wide uniqueness rule.
        self.assertEqual(self.select(rows=[self.rows[0], self.rows[0], self.rows[1]]), (2, 2))

    def test_exact_identity_does_not_normalize_or_match_substrings(self):
        for label in ('资本', ' 资本经营', '资本经营 '):
            with self.subTest(label=label), self.assertRaisesRegex(ValueError, 'row label'):
                self.select(row_label=label)

    def test_empty_ragged_or_unidentified_selection_is_rejected(self):
        for rows in ([], [['资本经营', '7']]):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                self.select(rows=rows)
        for value in ('', ' ', 1, None):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, 'nonempty string'):
                self.select(row_label=value)


if __name__ == '__main__': unittest.main()
