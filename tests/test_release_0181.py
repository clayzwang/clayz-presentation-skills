"""Art-owned editing boundaries checked against actual native PPTX objects."""
import copy
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'packages/validators'))
sys.path.insert(0, str(ROOT))
from tests.test_release_0180 import clean_handoff, relock, write_native
from validate_art_direction_plan import validate_plan
from compare_package_to_pptx import compare
from pptx import Presentation


class ArtEditingBoundaryTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.folder = Path(temporary.name)
        _, self.package, self.plan = clean_handoff(self.folder)
        self.plan['contract_version'] = '2.2'
        for element in self.elements:
            element.update(native_name='ART::'+element['element_id'], render_separately=True)
        relock(self.package, self.plan)

    @property
    def elements(self):
        return self.plan['visual_baseline']['slides'][0]['elements']

    def separate_heading_and_body(self):
        heading = next(e for e in self.elements if e['element_id']=='E-H')
        body = copy.deepcopy(heading)
        heading['copy_ids'] = ['H']
        body.update(element_id='E-B', native_name='ART::E-B', copy_ids=['B'])
        self.elements.append(body)
        for mapping in self.plan['slides'][0]['copy_unit_map']:
            if mapping['copy_id'] in {'H','B'}:
                mapping['native_location'].pop('text_range')
            if mapping['copy_id']=='B':
                mapping.update(render_target_id='E-B', native_location={'shape_name':'ART::E-B'})
        self.plan['slides'][0]['medium_execution_contract']['minimum_object_counts']={'shape':5}
        relock(self.package, self.plan)

    def write(self):
        path = self.folder/'native.pptx'
        write_native(self.package, self.plan, path)
        return path

    def test_explicit_shared_text_remains_editable(self):
        # Several Copy units form one Art-declared physical object.
        self.assertEqual([], validate_plan(self.package, self.plan))
        self.assertEqual([], compare(self.package, self.plan, self.write()))

    def test_art_separates_objects_without_copy_rendering_commands(self):
        self.separate_heading_and_body()
        self.assertEqual([], validate_plan(self.package, self.plan))
        path = self.write()
        self.assertEqual([], compare(self.package, self.plan, path))
        self.assertEqual(5, len(Presentation(path).slides[0].shapes))
        self.assertFalse(any('render_separately' in u for u in self.package['copy_layer']['slides'][0]['copy_units']))

    def test_output_cannot_merge_art_declared_objects(self):
        self.separate_heading_and_body()
        path = self.write()
        deck = Presentation(path)
        heading = next(s for s in deck.slides[0].shapes if s.name=='ART::E-H')
        body = next(s for s in deck.slides[0].shapes if s.name=='ART::E-B')
        heading.text += '\n'+body.text
        body._element.getparent().remove(body._element)
        deck.save(path)
        self.assertTrue(any('requires one separate native object' in e for e in compare(self.package, self.plan, path)))

    def test_art_cannot_alias_two_declared_objects_to_one_native_name(self):
        self.separate_heading_and_body()
        body = next(e for e in self.elements if e['element_id']=='E-B')
        body['native_name']='ART::E-H'
        units={u['copy_id']:u['text'] for u in self.package['copy_layer']['slides'][0]['copy_units']}
        for mapping in self.plan['slides'][0]['copy_unit_map']:
            if mapping['copy_id']=='H':mapping['native_location']['text_range']=[0,len(units['H'])]
            elif mapping['copy_id']=='B':
                start=len(units['H'])+1
                mapping['native_location']={'shape_name':'ART::E-H','text_range':[start,start+len(units['B'])]}
        relock(self.package, self.plan)
        self.assertTrue(any('unique native_name' in e for e in validate_plan(self.package, self.plan)))

    def test_art_must_declare_object_identity_and_separation(self):
        for field in ('native_name','render_separately'):
            with self.subTest(field=field):
                plan = copy.deepcopy(self.plan)
                plan['visual_baseline']['slides'][0]['elements'][0].pop(field)
                relock(self.package, plan)
                self.assertTrue(validate_plan(self.package, plan))

    def test_native_group_keeps_two_editable_child_objects(self):
        self.separate_heading_and_body()
        for element in self.elements:
            if element['element_id'] in {'E-H','E-B'}:
                element['native_group_path']=['ART::editing-group']
        relock(self.package, self.plan)
        path = self.write()
        deck = Presentation(path)
        children = [s for s in deck.slides[0].shapes if s.name in {'ART::E-H','ART::E-B'}]
        group = deck.slides[0].shapes.add_group_shape(children)
        group.name='ART::editing-group'
        deck.save(path)
        self.assertEqual([], compare(self.package, self.plan, path))
        reopened = Presentation(path)
        group = next(s for s in reopened.slides[0].shapes if s.name=='ART::editing-group')
        self.assertEqual({'ART::E-H','ART::E-B'}, {s.name for s in group.shapes})
        self.assertTrue(all(s.has_text_frame for s in group.shapes))
        # Ungrouping changes the approved interaction even when all text survives.
        slide = reopened.slides[0]
        for child in list(group.shapes):
            slide.shapes._spTree.append(child._element)
        group._element.getparent().remove(group._element)
        reopened.save(path)
        self.assertTrue(any('native grouping differs' in e for e in compare(self.package, self.plan, path)))

    def test_output_cannot_add_unapproved_native_grouping(self):
        self.separate_heading_and_body()
        path = self.write()
        deck = Presentation(path)
        group = deck.slides[0].shapes.add_group_shape([s for s in deck.slides[0].shapes if s.name in {'ART::E-H','ART::E-B'}])
        group.name='Unapproved convenience group'
        deck.save(path)
        self.assertTrue(any('native grouping differs' in e for e in compare(self.package, self.plan, path)))

    def test_flattened_picture_cannot_replace_native_text(self):
        from PIL import Image
        path = self.write()
        image = self.folder/'flattened.png'
        Image.new('RGB',(120,40),'white').save(image)
        deck = Presentation(path)
        target = next(s for s in deck.slides[0].shapes if s.name=='ART::E-H')
        picture = deck.slides[0].shapes.add_picture(str(image),target.left,target.top,target.width,target.height)
        picture.name=target.name
        target._element.getparent().remove(target._element)
        deck.save(path)
        self.assertTrue(any('native type changed' in e for e in compare(self.package, self.plan, path)))

    def test_historical_art_21_replay_remains_supported(self):
        self.plan['contract_version']='2.1'
        for element in self.elements:
            element.pop('native_name');element.pop('render_separately')
        relock(self.package, self.plan)
        self.assertEqual([], compare(self.package, self.plan, self.write()))


if __name__=='__main__':
    unittest.main()
