"""Real native-object regressions for the v0.18.0 content/design boundary."""
import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'packages/validators'))
sys.path.insert(0, str(ROOT))
from tests.test_story_visual_handoff import synthetic_handoff, reference
from tests.test_research_content_handoff import convert_to_research_handoff
from story_handoff import digest, spec_digest, render_document, build_stage_documents, validate_embedded_documents
from validate_logic_package import validate_package as validate_logic
from validate_ppt_package import validate_package as validate_copy
from validate_art_direction_plan import validate_plan
from compare_package_to_pptx import compare
from clean_content import RETIRED_COPY_KEYS, RETIRED_QA_CHECKS
from validate_supervision_report import target_counts
from scripts.build_runtime_packs import build_light
from scripts.verify_release_bundles import verify_light
from packages.index_runtime.capability import CORE_BY_STAGE, mandatory_core


def clean_handoff(folder):
    _, package, old_plan, _ = synthetic_handoff(folder)
    convert_to_research_handoff(package, old_plan)
    package['contract_version'] = '3.3'
    package['logic_layer'] = None
    units = [
        dict(copy_id='T', role='title', text='The pilot supports a bounded next step'),
        dict(copy_id='H', role='heading', heading_level=3, text='Observed result'),
        dict(copy_id='B', role='body', text='The synthetic result retains its original uncertainty.'),
        dict(copy_id='H2', role='heading', heading_level=3, text='A different emphasis'),
        dict(copy_id='A', role='annotation', text='Synthetic evidence only.'),
    ]
    layer = package['copy_layer']
    package['copy_layer'] = {key: layer[key] for key in ('logic_version','pagination_owner','research_sha256','semantic_preservation_review')}
    page = dict(slide_id='S01', narrative_role='body', copy_units=units)
    package['copy_layer']['slides'] = [page]
    package['copy_provenance'] = {u['copy_id']: ['B1'] for u in units}
    origin = copy.deepcopy(package)
    origin.update(status='logic-approved', copy_layer=None)
    origin.pop('copy_provenance')
    origin.pop('logic_artifact', None)
    path = folder/'logic-33.json'
    path.write_text(json.dumps(origin), encoding='utf-8')
    package['logic_artifact'] = reference(path)
    elements, mappings = [], []
    for i, unit in enumerate(units):
        if unit['copy_id'] == 'B':
            continue
        target = 'E-'+unit['copy_id']
        cids = ['H','B'] if unit['copy_id'] == 'H' else [unit['copy_id']]
        elements.append(dict(element_id=target, copy_ids=cids, kind='text', purpose='Synthetic editorial object',
            box=[.05,.05+i*.15,.85,.12], native_type='text', group_id='page', alignment='left',
            typography=dict(font_family='Arial',size_pt=18+i*2), locked_properties=['text'], allowed_adjustments={'box_delta_max':.01}))
        for cid in cids:
            location = dict(shape_name='ART::'+target)
            if cid == 'H':
                location['text_range']=[0,len(units[1]['text'])]
            elif cid == 'B':
                start=len(units[1]['text'])+1
                location['text_range']=[start,start+len(units[2]['text'])]
            mappings.append(dict(copy_id=cid, render_target_id=target, target_type='shape', native_location=location,
                                 style_token='emphasis-'+cid, reading_order=6-units.index(next(u for u in units if u['copy_id']==cid))))
    design=dict(slide_id='S01', reading_sequence=['A','H2','H','B','T'], copy_unit_map=mappings,
                medium_execution_contract=dict(structure_type='free-editorial-composition',minimum_object_counts={'shape':4}))
    plan={key:copy.deepcopy(old_plan[key]) for key in ('package_id','package_version','acceptance_contract','resource_inventory_lock','communication_contract','index_evidence','art_direction')}
    plan.update(contract_version='2.1',package_contract_version='3.3',status='art-direction-approved',slides=[design],
        reference_research=dict(learning_package_available=False,source_strategy='autonomous',references=[],
                               notes='Synthetic test has no browsing source; original layout is permitted.'))
    preview=copy.deepcopy(old_plan['visual_baseline']['slides'][0])
    preview.update(slide_id='S01',elements=elements)
    plan['visual_baseline']=dict(status='locked',locked_at='2026-10-04T01:00:00+00:00',
        copy_package_sha256=digest(package),spec_sha256=spec_digest(plan),slides=[preview])
    return origin,package,plan


def relock(package, plan):
    plan['visual_baseline'].update(copy_package_sha256=digest(package),spec_sha256=spec_digest(plan))


def write_native(package, plan, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation=Presentation()
    slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    units={u['copy_id']:u for u in package['copy_layer']['slides'][0]['copy_units']}
    for element in plan['visual_baseline']['slides'][0]['elements']:
        x,y,w,h=element['box']
        if element['native_type']=='table':
            shape=slide.shapes.add_table(2,2,Inches(x*10),Inches(y*7.5),Inches(w*10),Inches(h*7.5))
            frame=shape.table.cell(0,0).text_frame
        else:
            shape=slide.shapes.add_textbox(Inches(x*10),Inches(y*7.5),Inches(w*10),Inches(h*7.5))
            frame=shape.text_frame
        shape.name='ART::'+element['element_id']
        frame.text='\n'.join(units[cid]['text'] for cid in element['copy_ids'])
        for i,p in enumerate(frame.paragraphs):
            p.font.name='Arial';p.font.size=Pt(element['typography']['size_pt']+i*2)
    presentation.save(path)


class CleanContentTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name)
        self.origin,self.package,self.plan=clean_handoff(self.folder)

    def test_clean_text_reaches_art_without_old_hierarchy(self):
        self.assertEqual([],validate_logic(self.origin))
        self.assertEqual([],validate_copy(self.package))
        self.assertEqual([],validate_plan(self.package,self.plan))
        text=render_document('copy',self.package)
        self.assertIn('heading, level=3',text)
        self.assertNotIn('parent=',text)
        self.assertEqual({'shape':4,'table-cell':0,'chart-label':0},target_counts(self.plan['slides'][0]))

    def test_body_only_and_annotations_only_pages_are_valid(self):
        for role in ('body','annotation','subtitle','heading'):
            with self.subTest(role=role):
                package=copy.deepcopy(self.package)
                unit=dict(copy_id='ONLY',role=role,text='A complete statement of the same research.')
                if role=='heading':unit['heading_level']=7
                package['copy_layer']['slides'][0]['copy_units']=[unit]
                package['copy_provenance']={'ONLY':['B1']}
                self.assertEqual([],validate_copy(package))

    def test_retired_visual_fields_cannot_leak_into_new_copy(self):
        for key in RETIRED_COPY_KEYS:
            with self.subTest(key=key):
                package=copy.deepcopy(self.package)
                package['copy_layer']['slides'][0][key]=True
                self.assertTrue(any('retired Copy presentation field' in e for e in validate_copy(package)))

    def test_real_pptx_merges_heading_body_and_changes_visual_order(self):
        path=self.folder/'merged.pptx';write_native(self.package,self.plan,path)
        self.assertEqual([],compare(self.package,self.plan,path))
        from pptx import Presentation
        slide=Presentation(path).slides[0]
        self.assertEqual(4,len(slide.shapes))
        merged=next(s for s in slide.shapes if s.name=='ART::E-H')
        self.assertEqual(2,len(merged.text_frame.paragraphs))
        self.assertNotEqual(merged.text_frame.paragraphs[0].font.size,merged.text_frame.paragraphs[1].font.size)

    def test_missing_changed_or_extra_text_still_fails(self):
        from pptx import Presentation
        for change in ('missing','changed','extra','case'):
            with self.subTest(change=change):
                path=self.folder/(change+'.pptx');write_native(self.package,self.plan,path)
                deck=Presentation(path);shape=next(s for s in deck.slides[0].shapes if s.name=='ART::E-H')
                if change=='missing':shape.text='Observed result'
                elif change=='changed':shape.text=shape.text.replace('uncertainty','certainty')
                elif change=='case':shape.text=shape.text.lower()
                else:shape.text+=' Added unsupported claim.'
                deck.save(path)
                self.assertTrue(compare(self.package,self.plan,path))

    def test_shared_ranges_must_not_overlap(self):
        mapping=next(m for m in self.plan['slides'][0]['copy_unit_map'] if m['copy_id']=='B')
        mapping['native_location']['text_range']=[0,10]
        relock(self.package,self.plan)
        self.assertTrue(any('disjoint text_range' in e for e in validate_plan(self.package,self.plan)))

    def test_two_units_can_share_one_native_table_cell(self):
        element=next(e for e in self.plan['visual_baseline']['slides'][0]['elements'] if e['element_id']=='E-H')
        element.update(native_type='table',kind='table')
        design=self.plan['slides'][0]
        for mapping in design['copy_unit_map']:
            if mapping['copy_id'] in {'H','B'}:
                mapping['target_type']='table-cell'
                mapping['native_location'].update(row=0,column=0)
        design['medium_execution_contract']['minimum_object_counts']={'shape':3,'native-table':1}
        relock(self.package,self.plan)
        path=self.folder/'combined-cell.pptx';write_native(self.package,self.plan,path)
        self.assertEqual([],compare(self.package,self.plan,path))

    def test_one_table_can_bind_text_in_different_cells(self):
        from pptx import Presentation
        element=next(e for e in self.plan['visual_baseline']['slides'][0]['elements'] if e['element_id']=='E-H')
        element.update(native_type='table',kind='table')
        for mapping in self.plan['slides'][0]['copy_unit_map']:
            if mapping['copy_id'] in {'H','B'}:
                mapping['target_type']='table-cell'
                mapping['native_location']={'shape_name':'ART::E-H','row':0,'column':0 if mapping['copy_id']=='H' else 1}
        self.plan['slides'][0]['medium_execution_contract']['minimum_object_counts']={'shape':3,'native-table':1}
        relock(self.package,self.plan)
        path=self.folder/'separate-cells.pptx';write_native(self.package,self.plan,path)
        deck=Presentation(path);table=next(s.table for s in deck.slides[0].shapes if s.name=='ART::E-H')
        units={u['copy_id']:u['text'] for u in self.package['copy_layer']['slides'][0]['copy_units']}
        table.cell(0,0).text=units['H'];table.cell(0,1).text=units['B'];deck.save(path)
        self.assertEqual([],compare(self.package,self.plan,path))

    def test_native_ranges_preserve_spaces_and_soft_line_breaks(self):
        from pptx import Presentation
        path=self.folder/'native-breaks.pptx';write_native(self.package,self.plan,path)
        deck=Presentation(path);shape=next(s for s in deck.slides[0].shapes if s.name=='ART::E-H')
        heading='  Observed\vresult  '
        body=next(u['text'] for u in self.package['copy_layer']['slides'][0]['copy_units'] if u['copy_id']=='B')
        shape.text_frame.text=heading+'\n'+body;deck.save(path)
        start=len(heading)+1
        for mapping in self.plan['slides'][0]['copy_unit_map']:
            if mapping['copy_id']=='H':mapping['native_location']['text_range']=[0,len(heading)]
            elif mapping['copy_id']=='B':mapping['native_location']['text_range']=[start,start+len(body)]
        relock(self.package,self.plan)
        self.assertEqual([],compare(self.package,self.plan,path))

    def test_malformed_native_bindings_report_errors(self):
        for field,value in (('reading_sequence',[{}]),('copy_unit_map',[{'copy_id':[]}]),
                            ('copy_unit_map',[{'copy_id':'T','render_target_id':[],'native_location':[]}])):
            with self.subTest(field=field,value=value):
                plan=copy.deepcopy(self.plan);plan['slides'][0][field]=value;relock(self.package,plan)
                self.assertTrue(validate_plan(self.package,plan))

    def test_native_chart_title_has_an_exact_editable_location(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        title=next(u['text'] for u in self.package['copy_layer']['slides'][0]['copy_units'] if u['copy_id']=='T')
        data_id='SYNTHETIC-DATA-1'
        self.origin['research']['data']=[dict(data_id=data_id,metric_name='Observed',display_value='1',raw_value=1,
            unit='count',period='Synthetic period',definition_ref='Synthetic definition',
            source_ids=[self.origin['research']['sources'][0]['source_id']],evidence_status='source-fact')]
        self.origin['research']['findings'][0]['data_ids']=[data_id]
        self.package['research']=copy.deepcopy(self.origin['research'])
        self.package['copy_layer']['research_sha256']=digest(self.package['research'])
        origin_path=Path(self.package['logic_artifact']['path'])
        origin_path.write_text(json.dumps(self.origin),encoding='utf-8')
        self.package['logic_artifact']=reference(origin_path)
        element=next(e for e in self.plan['visual_baseline']['slides'][0]['elements'] if e['element_id']=='E-T')
        element.update(native_type='chart',kind='chart',chart_type='bar',data_ids=[data_id])
        mapping=next(m for m in self.plan['slides'][0]['copy_unit_map'] if m['copy_id']=='T')
        mapping.update(target_type='chart-label',native_location={'shape_name':'ART::E-T','label_kind':'title'})
        self.plan['slides'][0]['medium_execution_contract']['minimum_object_counts']={'shape':3,'native-chart':1}
        relock(self.package,self.plan)
        # Start with the same approved text and replace its editable title shape
        # with a real native chart title. No image/OCR can establish this check.
        temporary=copy.deepcopy(self.plan)
        temporary['visual_baseline']['slides'][0]['elements'][0]['native_type']='text'
        path=self.folder/'chart-title.pptx';write_native(self.package,temporary,path)
        deck=Presentation(path);slide=deck.slides[0]
        shape=next(s for s in slide.shapes if s.name=='ART::E-T')
        node=shape._element;node.getparent().remove(node)
        data=CategoryChartData();data.categories=['Synthetic'];data.add_series('Observed',[1])
        shape=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.5),Inches(.5),Inches(8),Inches(2),data)
        shape.name='ART::E-T';shape.chart.has_title=True;shape.chart.chart_title.text_frame.text=title
        deck.save(path)
        self.assertEqual([],compare(self.package,self.plan,path))

    def test_copy_meaning_and_source_coverage_remain_bound(self):
        self.package['research']['summary']='An unsupported new conclusion'
        self.package['copy_layer']['research_sha256']=digest(self.package['research'])
        self.assertTrue(any('immutable Logic research' in e for e in validate_copy(self.package)))
        self.package['research']=copy.deepcopy(self.origin['research'])
        self.package['copy_layer']['research_sha256']=digest(self.package['research'])
        self.package['copy_provenance']={cid:[] for cid in self.package['copy_provenance']}
        self.assertTrue(any('required research finding' in e for e in validate_copy(self.package)))

    def test_learning_priority_and_no_match_original_design(self):
        r=self.plan['reference_research'];r.update(learning_package_available=True,source_strategy='learning-first',notes='Relevant learning content was inspected; no layout matched, so this design is original.')
        relock(self.package,self.plan)
        self.assertEqual([],validate_plan(self.package,self.plan))
        r['source_strategy']='web-first';relock(self.package,self.plan)
        self.assertTrue(any('first reference source' in e for e in validate_plan(self.package,self.plan)))

    def test_portable_documents_keep_clean_handoff(self):
        report={'stage_documents':build_stage_documents(self.package,self.plan)}
        self.assertEqual([],validate_embedded_documents(report))


class LightBoundaryTests(unittest.TestCase):
    def test_both_light_targets_contain_no_art_layout_or_design_index(self):
        with tempfile.TemporaryDirectory() as td:
            for target in ('cloud','local'):
                with self.subTest(target=target):
                    path=build_light(Path(td),(ROOT/'VERSION').read_text().strip(),target)
                    self.assertEqual([],verify_light(path,target))
                    with zipfile.ZipFile(path) as archive:
                        records=[json.loads(line) for line in archive.read('clayz-presentation-skills/catalog/records.jsonl').decode().splitlines()]
                        self.assertTrue(records)
                        self.assertFalse(any('art-direction' in r['classification']['stages'] for r in records))
                        self.assertIn('clayz-presentation-skills/packages/validators/clean_content.py',archive.namelist())
                        for stage in CORE_BY_STAGE:
                            for required in mandatory_core(stage):
                                self.assertIn('clayz-presentation-skills/'+required['ref'],archive.namelist())


if __name__=='__main__':unittest.main()
