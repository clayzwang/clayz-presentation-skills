import unittest
from copy import deepcopy
from io import BytesIO
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from packages.layout.native_chart_text import set_point_label_wrapping


def fixture():
    prs=Presentation();sl=prs.slides.add_slide(prs.slide_layouts[6])
    data=CategoryChartData();data.categories=['first','second'];data.add_series('A',[13.526,3]);data.add_series('B',[6.888,2])
    chart=sl.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED,Inches(1),Inches(1),Inches(8),Inches(3),data).chart
    for si,text in [(0,'13.526'),(1,'6.888')]:
        label=chart.series[si].points[0].data_label;label.has_text_frame=True;label.text_frame.text=text
        label.text_frame.paragraphs[0].runs[0].font.size=Pt(17);label.position=XL_LABEL_POSITION.INSIDE_END
    return prs,chart


def request(si=0,pi=0,text='13.526',wrap=False):
    return dict(series_index=si,point_index=pi,expected_text=text,word_wrap=wrap)


class PointWrappingTests(unittest.TestCase):
    def test_selected_only_and_reopen_preserves_style_data(self):
        prs,ch=fixture();blob=ch.part.related_part(ch._chartSpace.xpath('./c:externalData/@r:id')[0]).blob
        label=ch.series[0]._element.get_dLbl(0);label.append(OxmlElement('c:spPr'));before=deepcopy(ch._chartSpace)
        receipt=set_point_label_wrapping(ch,[request()]);self.assertIs(receipt['selected_labels'][0]['authored_word_wrap'],False)
        before.xpath('.//c:ser')[0].xpath('./c:dLbls/c:dLbl/c:tx/c:rich/a:bodyPr')[0].set('wrap','none')
        self.assertEqual(ch._chartSpace.xml,before.xml)
        self.assertEqual(blob,ch.part.related_part(ch._chartSpace.xpath('./c:externalData/@r:id')[0]).blob)
        stream=BytesIO();prs.save(stream);stream.seek(0);new=Presentation(stream).slides[0].shapes[0].chart
        self.assertEqual(new._chartSpace.xml,ch._chartSpace.xml)
    def test_true_and_none_are_distinct_choices(self):
        _,ch=fixture();set_point_label_wrapping(ch,[request(wrap=True)]);self.assertEqual(ch._chartSpace.xpath('.//c:ser')[0].xpath('.//a:bodyPr/@wrap'),['square'])
        set_point_label_wrapping(ch,[request(wrap=None)]);self.assertEqual(ch._chartSpace.xpath('.//c:ser')[0].xpath('.//a:bodyPr/@wrap'),[])
    def test_invalid_second_request_changes_nothing(self):
        _,ch=fixture();before=ch._chartSpace.xml
        with self.assertRaises(ValueError):set_point_label_wrapping(ch,[request(),request(1,text='wrong')])
        self.assertEqual(ch._chartSpace.xml,before)
    def test_automatic_label_not_converted(self):
        _,ch=fixture();before=ch._chartSpace.xml
        with self.assertRaises(ValueError):set_point_label_wrapping(ch,[request(pi=1,text='3')])
        self.assertEqual(ch._chartSpace.xml,before)
    def test_duplicate_and_invalid_selectors_preflight(self):
        for invalid in [request(si=True),request(pi=-1),request(si=20),request(pi=20),request(wrap='none'),request(text='')]:
            _,ch=fixture();before=ch._chartSpace.xml
            with self.assertRaises(ValueError):set_point_label_wrapping(ch,[request(),invalid])
            self.assertEqual(ch._chartSpace.xml,before)
        _,ch=fixture();before=ch._chartSpace.xml
        with self.assertRaises(ValueError):set_point_label_wrapping(ch,[request(),request()])
        self.assertEqual(ch._chartSpace.xml,before)
    def test_hard_breaks_and_multiple_selected_labels_preserved(self):
        _,ch=fixture();ch.series[0].points[0].data_label.text_frame.text='13.\n526'
        receipt=set_point_label_wrapping(ch,[request(text='13.\n526'),request(1,text='6.888',wrap=True)])
        self.assertEqual(receipt['selected_labels'][0]['text'],'13.\n526');self.assertTrue(receipt['selected_labels'][1]['authored_word_wrap'])
    def test_empty_selection_no_mutation(self):
        _,ch=fixture();before=ch._chartSpace.xml;self.assertEqual(set_point_label_wrapping(ch,[])['selected_labels'],[]);self.assertEqual(ch._chartSpace.xml,before)
    def test_nonchart_and_missing_fields_rejected(self):
        with self.assertRaises(TypeError):set_point_label_wrapping(object(),[])
        _,ch=fixture();before=ch._chartSpace.xml
        with self.assertRaises(ValueError):set_point_label_wrapping(ch,[{'word_wrap':False}])
        self.assertEqual(ch._chartSpace.xml,before)

if __name__=='__main__':unittest.main()
