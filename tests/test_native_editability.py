"""Native editing defaults and maintained node links, without GUI claims."""
import unittest
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from packages.layout.native_text import add_native_textbox, add_text_shape, write_native_text, prepare_native_text
from packages.layout.native_connector import add_rectangular_connector, refresh_rectangular_connector
from packages.layout.semantic_table import Column, plan, write
from packages.adapters.python_pptx.render import _add_object, _imports
from tests.font_support import test_font_path


class NativeEditabilityTest(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def test_logical_text_native_wrap_percentage_spacing_and_paragraphs(self):
        value = 'A continuous sentence remains editable and is not pre-broken for this width.'
        obj = add_native_textbox(self.slide, value, 1, 1, 2, 2, line_spacing=1.15)
        self.assertEqual(obj.text, value)
        self.assertEqual(len(obj.text_frame.paragraphs), 1)
        self.assertTrue(obj.text_frame.word_wrap)
        self.assertAlmostEqual(obj.text_frame.paragraphs[0].line_spacing, 1.15)
        self.assertIn('spcPct', obj._element.xml)
        self.assertNotIn('<a:br', obj._element.xml)
        write_native_text(obj, value+'\nA real second paragraph.', line_spacing=1.0)
        self.assertEqual(len(obj.text_frame.paragraphs), 2)
        self.assertEqual(obj.text_frame.vertical_anchor, MSO_ANCHOR.TOP)
        self.assertEqual(obj.text_frame.paragraphs[0].space_after.pt, 6)
        write_native_text(obj, 'Shorter body after deleting the second paragraph.', line_spacing=1.15)
        self.assertEqual(len(obj.text_frame.paragraphs), 1)
        self.assertAlmostEqual(obj.text_frame.paragraphs[0].line_spacing, 1.15)

    def test_label_defaults_single_shape_and_adapter_own_text(self):
        node = add_text_shape(self.slide, 'A node label', 1, 1, 2, 1)
        self.assertEqual(len(self.slide.shapes), 1)
        self.assertEqual(node.text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)
        self.assertEqual(node.text_frame.paragraphs[0].alignment, PP_ALIGN.CENTER)
        self.assertAlmostEqual(node.text_frame.margin_left.inches, .12)
        _add_object(self.slide, {'object_id':'OWN-TEXT','type':'shape','shape':'rect',
                    'text':'Adapter node', 'options':{'x':4,'y':1,'w':2,'h':1}}, Path('.'), _imports())
        self.assertEqual(len(self.slide.shapes), 2)
        self.assertEqual(self.slide.shapes[1].text, 'Adapter node')
        self.assertEqual(self.slide.shapes[1].text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)
        _add_object(self.slide, {'object_id':'TABLE','type':'table','rows':[['Header'],['Body']],
                    'options':{'x':1,'y':3,'w':4,'h':2,'margin':.2}}, Path('.'), _imports())
        cell=self.slide.shapes[2].table.cell(0,0)
        self.assertEqual(cell.vertical_anchor, MSO_ANCHOR.MIDDLE)
        self.assertAlmostEqual(cell.margin_left.inches,.2)

    def test_table_estimated_lines_are_not_native_breaks_and_header_cell_anchor(self):
        spec = plan([Column('工作标签')], [['这项完整工作说明在当前较窄表格中预计需要多行显示']],
                    width=1.8, font_path=test_font_path(), caption='Synthetic sample', comparison='Work identity')
        self.assertGreater(len(spec['cells'][1][0]['lines']), 1)
        shape = write(self.slide, spec, 1, 1)
        body = shape.table.cell(1, 0); header = shape.table.cell(0, 0)
        self.assertEqual(len(body.text_frame.paragraphs), 1)
        self.assertNotIn('<a:br', body._tc.xml)
        self.assertTrue(body.text_frame.word_wrap)
        self.assertEqual(body.text_frame.paragraphs[0].line_spacing, 1.0)
        self.assertEqual(header.vertical_anchor, MSO_ANCHOR.MIDDLE)
        self.assertEqual(header.text_frame.paragraphs[0].alignment, PP_ALIGN.CENTER)
        self.assertAlmostEqual(body.margin_left.inches, .17)

    def test_programmatic_move_requires_refresh_and_native_end_marker(self):
        a = add_text_shape(self.slide, 'A', 1, 1, 2, 1)
        b = add_text_shape(self.slide, 'B', 5, 1, 2, 1)
        line = add_rectangular_connector(self.slide, begin=(a,'right'), end=(b,'left'), tail={'type':'triangle'})
        self.assertEqual(len(self.slide.shapes), 3)
        self.assertIn('stCxn', line._element.xml)
        self.assertIn('endCxn', line._element.xml)
        self.assertIn('tailEnd', line._element.xml)
        old = (line.end_x, line.end_y)
        b.top += Inches(1)
        self.assertEqual((line.end_x,line.end_y), old)  # No automatic observer.
        refresh_rectangular_connector(line)
        self.assertEqual(line.end_x, b.left)
        self.assertEqual(line.end_y, b.top+b.height//2)
        self.assertNotEqual((line.end_x,line.end_y), old)
        self.assertEqual(line._element.xpath('.//a:endCxn')[0].get('id'), str(b.shape_id))

    def test_cross_slide_rejected_before_creating_a_line(self):
        a = add_text_shape(self.slide, 'A', 1, 1, 2, 1)
        other = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        b = add_text_shape(other, 'B', 5, 1, 2, 1)
        before = len(self.slide.shapes)
        with self.assertRaises(ValueError):
            add_rectangular_connector(self.slide, begin=(a,'right'), end=(b,'left'))
        self.assertEqual(len(self.slide.shapes), before)

    def test_local_fixed_or_no_wrap_requires_a_reason(self):
        obj = add_native_textbox(self.slide, 'Label', 1, 1, 2, 1)
        with self.assertRaisesRegex(ValueError, 'reason'):
            prepare_native_text(obj, word_wrap=False)
        prepare_native_text(obj, role='label', word_wrap=False,
                            fixed_line_spacing_pt=24, exception_reason='One-line technical field')
        self.assertEqual(obj.text_frame.paragraphs[0].line_spacing.pt, 24)

    def test_table_multiple_spacing_changes_candidate_height_and_max_height(self):
        args=dict(columns=[Column('工作')], rows=[['连续工作说明用于比较行高估算']],
                  width=2.4,font_path=test_font_path(),caption='合成内容',comparison='同文字与行距')
        single=plan(**args,line_spacing=1.0)
        double=plan(**args,line_spacing=2.0)
        self.assertGreater(double['height'],single['height'])
        available=(single['height']+double['height'])/2
        plan(**args,line_spacing=1.0,max_height=available)
        with self.assertRaisesRegex(ValueError,'revise Art'):
            plan(**args,line_spacing=2.0,max_height=available)


if __name__ == '__main__': unittest.main()
