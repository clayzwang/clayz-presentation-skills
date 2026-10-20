import copy
import unittest

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from packages.layout import set_connector_line_ends
from packages.layout.native_connector import attach_rectangular_connector


class NativeLineEndsTest(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.sl = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.a = self.sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
        self.b = self.sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(3), Inches(2), Inches(1))
        self.line = self.sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(3), Inches(2))
        self.line.line.width = Pt(2)
        self.line.line.color.rgb = RGBColor.from_string('197F84')
        self.line.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def markers(self, line=None):
        ln = (line or self.line)._element.spPr.ln
        return {name: dict(node.attrib) for name in ['headEnd', 'tailEnd']
                if ln is not None and (node := ln.find(qn('a:' + name))) is not None}

    def non_marker_xml(self):
        root = copy.deepcopy(self.line._element)
        for node in root.xpath('.//a:headEnd | .//a:tailEnd'):
            node.getparent().remove(node)
        return etree.tostring(root)

    def reject_unchanged(self, **kwargs):
        before = self.line._element.xml
        with self.assertRaises(ValueError):
            set_connector_line_ends(self.line, **kwargs)
        self.assertEqual(self.line._element.xml, before)

    def test_two_native_ends_no_new_object_and_all_other_xml_preserved(self):
        attach_rectangular_connector(self.line, begin=(self.a, 'right'), end=(self.b, 'left'))
        before, n = self.non_marker_xml(), len(self.sl.shapes)
        self.assertIs(set_connector_line_ends(self.line, head={'type': 'oval', 'w': 'sm'},
                      tail={'type': 'triangle', 'w': 'lg', 'len': 'med'}), self.line)
        self.assertEqual(self.markers(), {'headEnd': {'type': 'oval', 'w': 'sm'},
                                        'tailEnd': {'type': 'triangle', 'w': 'lg', 'len': 'med'}})
        self.assertEqual(self.non_marker_xml(), before)
        self.assertEqual(len(self.sl.shapes), n)

    def test_omitted_other_end_and_omitted_attributes_are_preserved(self):
        set_connector_line_ends(self.line, head={'type': 'oval', 'len': 'lg'},
                                tail={'type': 'triangle', 'w': 'sm', 'len': 'med'})
        head = etree.tostring(self.line._element.spPr.ln.find(qn('a:headEnd')))
        set_connector_line_ends(self.line, tail={'type': 'arrow'})
        self.assertEqual(etree.tostring(self.line._element.spPr.ln.find(qn('a:headEnd'))), head)
        self.assertEqual(self.markers()['tailEnd'], {'type': 'arrow', 'w': 'sm', 'len': 'med'})

    def test_explicit_none_and_inherit_are_distinct(self):
        set_connector_line_ends(self.line, tail={'type': 'none'})
        self.assertEqual(self.markers()['tailEnd'], {'type': 'none'})
        set_connector_line_ends(self.line, tail='inherit')
        self.assertNotIn('tailEnd', self.markers())
        self.assertIsNotNone(self.line._element.spPr.ln)
        self.assertEqual(self.line.line.width, Pt(2))

    def test_inherit_without_existing_line_properties_does_not_create_any(self):
        plain = self.sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(1), 0)
        before = plain._element.xml
        set_connector_line_ends(plain, head='inherit', tail='inherit')
        self.assertEqual(plain._element.xml, before)

    def test_invalid_second_end_leaves_all_first_end_and_style_bytes_unchanged(self):
        set_connector_line_ends(self.line, head={'type': 'oval'})
        for bad in [{'type': 'unrecognised'}, {'type': 'arrow', 'w': 'huge'},
                    {'type': 'arrow', 'len': None}, {'type': []}, {}, False,
                    {'type': 'arrow', 'size': 'sm'}, 'none']:
            self.reject_unchanged(head={'type': 'triangle'}, tail=bad)

    def test_invalid_second_end_does_not_create_line_properties(self):
        plain = self.sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(1), 0)
        before = plain._element.xml
        with self.assertRaises(ValueError):
            set_connector_line_ends(plain, head={'type': 'triangle'}, tail={'type': 'arrow', 'w': 1})
        self.assertEqual(plain._element.xml, before)

    def test_new_ends_respect_line_schema_order_and_unrelated_extension(self):
        ln = self.line._element.spPr.ln
        join = OxmlElement('a:round'); ln.append(join)
        ext = OxmlElement('a:extLst'); ln.append(ext)
        set_connector_line_ends(self.line, tail={'type': 'diamond'})
        set_connector_line_ends(self.line, head={'type': 'oval'})
        tags = [etree.QName(x).localname for x in ln]
        self.assertEqual(tags[-4:], ['round', 'headEnd', 'tailEnd', 'extLst'])
        self.assertIs(ln[-1], ext)

    def test_reverse_and_both_flips_preserve_geometry_and_end_marker_selection(self):
        for start, end in [((8, 5), (2, 1)), ((8, 1), (2, 5)), ((2, 5), (8, 1))]:
            line = self.sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                *[Inches(v) for v in (*start, *end)])
            before = etree.tostring(line._element.spPr.xfrm)
            set_connector_line_ends(line, head={'type': 'oval'}, tail={'type': 'triangle'})
            self.assertEqual(etree.tostring(line._element.spPr.xfrm), before)
            self.assertEqual(self.markers(line)['tailEnd']['type'], 'triangle')

    def test_explicit_rebind_preserves_native_marker_and_other_end(self):
        attach_rectangular_connector(self.line, begin=(self.a, 'right'), end=(self.b, 'left'))
        set_connector_line_ends(self.line, head={'type': 'oval'}, tail={'type': 'triangle'})
        old_begin = (self.line.begin_x, self.line.begin_y)
        marker_xml = etree.tostring(self.line._element.spPr.ln)
        old_end = (self.line.end_x, self.line.end_y)
        self.b.left += Inches(1)
        self.assertEqual((self.line.end_x, self.line.end_y), old_end)
        attach_rectangular_connector(self.line, end=(self.b, 'left'))
        self.assertEqual(self.line.end_x, self.b.left)
        self.assertEqual((self.line.begin_x, self.line.begin_y), old_begin)
        self.assertEqual(etree.tostring(self.line._element.spPr.ln), marker_xml)

    def test_existing_duplicate_marker_is_rejected_without_partial_mutation(self):
        ln = self.line._element.spPr.ln
        ln.append(OxmlElement('a:tailEnd')); ln.append(OxmlElement('a:tailEnd'))
        self.reject_unchanged(head={'type': 'triangle'}, tail='inherit')

    def test_scope_no_selection_zero_length_curved_rotated_group_and_nonconnector(self):
        self.reject_unchanged()
        for bad in [self.a,
                    self.sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0),
                    self.sl.shapes.add_connector(MSO_CONNECTOR.CURVE, 0, 0, Inches(2), Inches(2))]:
            before = bad._element.xml
            with self.assertRaises(ValueError):
                set_connector_line_ends(bad, tail={'type': 'triangle'})
            self.assertEqual(bad._element.xml, before)
        self.line.rotation = 20
        self.reject_unchanged(tail={'type': 'triangle'})
        group = self.sl.shapes.add_group_shape()
        grouped = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(1), 0)
        before = grouped._element.xml
        with self.assertRaises(ValueError):
            set_connector_line_ends(grouped, tail={'type': 'triangle'})
        self.assertEqual(grouped._element.xml, before)

    def test_idempotent_call_does_not_duplicate_markers_or_touch_separate_head(self):
        external = self.sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, 0, 0, Inches(.1), Inches(.1))
        before = external._element.xml
        for _ in range(2):
            set_connector_line_ends(self.line, tail={'type': 'triangle', 'w': 'med', 'len': 'med'})
        self.assertEqual(len(self.line._element.spPr.ln.findall(qn('a:tailEnd'))), 1)
        self.assertEqual(external._element.xml, before)


if __name__ == '__main__':
    unittest.main()
