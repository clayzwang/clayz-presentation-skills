"""Synthetic, inspectable PPTX and JSON fixtures for work-report tests.

The fixtures intentionally contain no private or production material.  They
are valid OOXML packages so the work-report acceptance tests can compare
machine-derived page/object/notes/media observations with the files bound by
the real CLI.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
REL_CHART = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
REL_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
REL_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
REL_TARGET = "Type"

_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8")


def write_json(path: Path, value: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(json_bytes(value))
    return path


def file_ref(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    return {"path": str(path.resolve()), "sha256": sha256_bytes(raw), "bytes": len(raw)}


def _set_run_font(run: Any, face: str = "Aptos") -> None:
    run.font.name = face


def _write_notes(slide: Any, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def build_report_pptx(path: Path, *, subject: str, rich: bool) -> Path:
    """Create a valid small presentation with deliberately different shapes.

    ``rich=True`` produces one slide with one native chart, one table, notes,
    and a media relationship on the slide master.  ``rich=False`` produces a
    single text-only slide with notes and no chart/table/master media.
    """

    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(blank)
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.7))
    title.name = f"COPY::{subject}-TITLE"
    run = title.text_frame.paragraphs[0].add_run()
    run.text = f"{subject.title()} research brief"
    run.font.size = Pt(26)
    _set_run_font(run)
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.25), Inches(12), Inches(0.8))
    body.name = f"COPY::{subject}-BODY"
    body_run = body.text_frame.paragraphs[0].add_run()
    body_run.text = "Synthetic evidence is bound to the work report."
    body_run.font.size = Pt(18)
    _set_run_font(body_run)
    _write_notes(slide, f"{subject} notes: research question, assumptions, and decision context.")

    if rich:
        table_shape = slide.shapes.add_table(3, 3, Inches(0.6), Inches(2.3), Inches(5.3), Inches(2.1))
        table_shape.name = f"TABLE::{subject}-EVIDENCE"
        table = table_shape.table
        values = (("Period", "Signal", "Value"), ("Q1", "alpha", "12"), ("Q2", "beta", "17"))
        for row_index, row in enumerate(values):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value

        chart_data = ChartData()
        chart_data.categories = ["Q1", "Q2", "Q3"]
        chart_data.add_series("Observed", (12, 17, 21))
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(6.25),
            Inches(2.05),
            Inches(6.2),
            Inches(3.3),
            chart_data,
        )
        chart.chart.has_legend = False
        chart.name = f"CHART::{subject}-TREND"

    presentation.save(path)
    if rich:
        _add_master_media(path)
    return path


def _add_master_media(path: Path) -> None:
    """Attach a real PNG to the first slide master and its relationship graph."""

    ET.register_namespace("p", P_NS)
    ET.register_namespace("a", A_NS)
    ET.register_namespace("r", R_NS)
    ET.register_namespace("pr", PKG_REL_NS)
    source = path.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source), "r") as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        names = set(archive.namelist())
        master_name = "ppt/slideMasters/slideMaster1.xml"
        rels_name = "ppt/slideMasters/_rels/slideMaster1.xml.rels"
        if master_name not in names or rels_name not in names:
            raise AssertionError("python-pptx fixture is missing the first slide master")

        master_root = ET.fromstring(archive.read(master_name))
        rels_root = ET.fromstring(archive.read(rels_name))
        ids = []
        for relationship in rels_root:
            value = relationship.get("Id", "")
            match = re.fullmatch(r"rId(\d+)", value)
            if match:
                ids.append(int(match.group(1)))
        relation_id = f"rId{max(ids, default=0) + 1}"
        ET.SubElement(
            rels_root,
            f"{{{PKG_REL_NS}}}Relationship",
            {"Id": relation_id, REL_TARGET: REL_IMAGE, "Target": "../media/master-brand.png"},
        )

        sp_tree = master_root.find(f".//{{{P_NS}}}spTree")
        if sp_tree is None:
            raise AssertionError("first slide master has no shape tree")
        picture = ET.Element(f"{{{P_NS}}}pic")
        nv_pic = ET.SubElement(picture, f"{{{P_NS}}}nvPicPr")
        ET.SubElement(nv_pic, f"{{{P_NS}}}cNvPr", {"id": "9001", "name": "Master media"})
        ET.SubElement(nv_pic, f"{{{P_NS}}}cNvPicPr")
        ET.SubElement(nv_pic, f"{{{P_NS}}}nvPr")
        blip_fill = ET.SubElement(picture, f"{{{P_NS}}}blipFill")
        ET.SubElement(blip_fill, f"{{{A_NS}}}blip", {f"{{{R_NS}}}embed": relation_id})
        stretch = ET.SubElement(blip_fill, f"{{{A_NS}}}stretch")
        ET.SubElement(stretch, f"{{{A_NS}}}fillRect")
        shape_properties = ET.SubElement(picture, f"{{{P_NS}}}spPr")
        transform = ET.SubElement(shape_properties, f"{{{A_NS}}}xfrm")
        ET.SubElement(transform, f"{{{A_NS}}}off", {"x": "120000", "y": "120000"})
        ET.SubElement(transform, f"{{{A_NS}}}ext", {"cx": "240000", "cy": "240000"})
        geometry = ET.SubElement(shape_properties, f"{{{A_NS}}}prstGeom", {"prst": "rect"})
        ET.SubElement(geometry, f"{{{A_NS}}}avLst")
        insert_at = len(sp_tree)
        for index, child in enumerate(sp_tree):
            if child.tag == f"{{{P_NS}}}extLst":
                insert_at = index
                break
        sp_tree.insert(insert_at, picture)

        content_types_name = "[Content_Types].xml"
        content_root = ET.fromstring(archive.read(content_types_name))
        if not any(item.get("Extension") == "png" for item in content_root):
            ET.SubElement(content_root, "{http://schemas.openxmlformats.org/package/2006/content-types}Default", {"Extension": "png", "ContentType": "image/png"})

        for name in archive.namelist():
            if name == master_name:
                payload = ET.tostring(master_root, encoding="utf-8", xml_declaration=True)
            elif name == rels_name:
                payload = ET.tostring(rels_root, encoding="utf-8", xml_declaration=True)
            elif name == content_types_name:
                payload = ET.tostring(content_root, encoding="utf-8", xml_declaration=True)
            else:
                payload = archive.read(name)
            target.writestr(name, payload)
        target.writestr("ppt/media/master-brand.png", _ONE_PIXEL_PNG)
    path.write_bytes(output.getvalue())


def _relationship_rows(raw: bytes) -> list[dict[str, str]]:
    root = ET.fromstring(raw)
    return [{"id": item.get("Id", ""), "type": item.get("Type", ""), "target": item.get("Target", "")} for item in root]


def pptx_object_stats(path: Path) -> dict[str, Any]:
    """Read page/text/notes/chart/table/media stats directly from OOXML."""

    slides: list[dict[str, Any]] = []
    slide_pattern = re.compile(r"ppt/slides/slide(\d+)\.xml$")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        slide_names = sorted((name for name in names if slide_pattern.fullmatch(name)), key=lambda name: int(slide_pattern.fullmatch(name).group(1)))
        for slide_name in slide_names:
            number = int(slide_pattern.fullmatch(slide_name).group(1))
            root = ET.fromstring(archive.read(slide_name))
            rels_name = f"ppt/slides/_rels/slide{number}.xml.rels"
            relationships = _relationship_rows(archive.read(rels_name)) if rels_name in names else []
            chart_count = sum(item["type"] == REL_CHART for item in relationships)
            notes_targets = [item["target"] for item in relationships if item["type"] == REL_NOTES]
            note_chars = 0
            note_text_runs = 0
            for target in notes_targets:
                notes_name = str(Path("ppt/slides") / target).replace("\\", "/")
                if notes_name.startswith("ppt/slides/../"):
                    notes_name = notes_name.replace("ppt/slides/../", "ppt/")
                if notes_name in names:
                    notes_root = ET.fromstring(archive.read(notes_name))
                    values = [item.text or "" for item in notes_root.findall(f".//{{{A_NS}}}t")]
                    note_text_runs += len(values)
                    note_chars += sum(len(value) for value in values)
            image_count = sum(item["type"] == REL_IMAGE for item in relationships)
            text_values = [item.text or "" for item in root.findall(f".//{{{A_NS}}}t")]
            table_count = len(root.findall(f".//{{{A_NS}}}tbl"))
            slides.append({
                "page": number,
                "text_runs": len(text_values),
                "text_chars": sum(len(value) for value in text_values),
                "notes_text_runs": note_text_runs,
                "notes_chars": note_chars,
                "tables": table_count,
                "charts": chart_count,
                "media": image_count,
            })

        master_media = 0
        for rels_name in names:
            if rels_name.startswith("ppt/slideMasters/_rels/") and rels_name.endswith(".rels"):
                master_media += sum(item["type"] == REL_IMAGE for item in _relationship_rows(archive.read(rels_name)))
        package_media = sorted(name for name in names if name.startswith("ppt/media/") and not name.endswith("/"))

    return {
        "page_count": len(slides),
        "slide_order": slide_names,
        "slides": slides,
        "totals": {
            "text_runs": sum(item["text_runs"] for item in slides),
            "text_chars": sum(item["text_chars"] for item in slides),
            "notes_text_runs": sum(item["notes_text_runs"] for item in slides),
            "notes_chars": sum(item["notes_chars"] for item in slides),
            "tables": sum(item["tables"] for item in slides),
            "charts": sum(item["charts"] for item in slides),
            "media": sum(item["media"] for item in slides),
            "master_media": master_media,
            "package_media": len(package_media),
        },
    }


__all__ = [
    "build_report_pptx",
    "file_ref",
    "json_bytes",
    "pptx_object_stats",
    "sha256_file",
    "write_json",
]
