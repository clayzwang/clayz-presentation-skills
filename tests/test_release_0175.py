"""Regression cases from 0.17.4 production, plus portable evidence integrity."""
import base64
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'packages/validators'))
sys.path.insert(0, str(ROOT / 'scripts'))
from report_evidence import compact_report, expand_report, load_report, render_compact_markdown
from stamp_pptx_metadata import stamp, remove_stamp
from audit_ppt_legibility import audit
from audit_pptx_size import inspect
from compare_package_to_pptx import compare, text_locations
from task_runtime import run_check


class Production175Tests(unittest.TestCase):
    def test_art_judgment_does_not_generate_size_commitments(self):
        from task_commitments import enrich_task_acceptance
        config = json.loads((ROOT / 'config/default.json').read_text())
        enriched, _ = enrich_task_acceptance({}, config)
        self.assertFalse(any(item['requirement_id'].startswith('COMMIT-SIZE-') for item in enriched['requirements']))

    def test_font_manifest_requires_original_authorized_bytes(self):
        from font_bundle import verified_fonts
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); directory = root / 'assets/fonts'; directory.mkdir(parents=True)
            manifest = directory / 'manifest.json'
            manifest.write_text('{"fonts": []}')
            with self.assertRaisesRegex(ValueError, 'pending'):
                verified_fonts(root)
            font = directory / 'synthetic-test.ttf'; font.write_bytes(b'test-only-fixture')
            license_file = directory / 'LICENSE.txt'; license_file.write_text('Synthetic test fixture only')
            entry = {'family':'STKaiti', 'path':'assets/fonts/synthetic-test.ttf',
                     'bytes':font.stat().st_size, 'sha256':hashlib.sha256(font.read_bytes()).hexdigest(),
                     'source':'synthetic regression fixture', 'redistribution_authorized':True,
                     'license_path':'assets/fonts/LICENSE.txt', 'license_sha256':hashlib.sha256(license_file.read_bytes()).hexdigest()}
            manifest.write_text(json.dumps({'fonts':[entry]}))
            self.assertEqual(len(verified_fonts(root)), 1)
            entry['redistribution_authorized'] = False
            manifest.write_text(json.dumps({'fonts':[entry]}))
            with self.assertRaisesRegex(ValueError, 'authorization'):
                verified_fonts(root)
            entry['redistribution_authorized'] = True
            manifest.write_text(json.dumps({'fonts':[entry]}))
            font.write_bytes(b'changed')
            with self.assertRaisesRegex(ValueError, 'differ'):
                verified_fonts(root)

    def test_stamp_preserves_content_and_uses_default_opc_namespaces(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); src=root/'before.pptx'; dst=root/'after.pptx'
            Presentation().save(src)
            stamp(src,dst,{'ClayzVersion':'0.17.5'})
            with zipfile.ZipFile(src) as before, zipfile.ZipFile(dst) as after:
                for name in before.namelist():
                    if name not in {'_rels/.rels','[Content_Types].xml','docProps/custom.xml'}:
                        self.assertEqual(before.read(name),after.read(name))
                for name,tag in [('_rels/.rels','Relationships'),('[Content_Types].xml','Types'),('docProps/custom.xml','Properties')]:
                    self.assertIn(('<'+tag+' xmlns=').encode(),after.read(name))
            Presentation(dst)
            restored=root/'removed.pptx'; remove_stamp(dst,restored,{'ClayzVersion'})
            with zipfile.ZipFile(restored) as z:
                self.assertNotIn('docProps/custom.xml',z.namelist())
                self.assertIn(b'<Relationships xmlns=',z.read('_rels/.rels'))

    def test_font_sizes_are_observations_not_failures(self):
        from pptx import Presentation
        from pptx.util import Inches,Pt
        with tempfile.TemporaryDirectory() as td:
            p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);t=s.shapes.add_textbox(0,0,Inches(2),Inches(1))
            t.text='Small technical label';t.text_frame.paragraphs[0].runs[0].font.size=Pt(7.5)
            path=Path(td)/'deck.pptx';p.save(path);result=audit(path,None)
            self.assertEqual(result['errors'],[])
            self.assertIn(7.5,result['slides'][0]['explicit_font_sizes_pt'])
            self.assertIsNone(result['slides'][0]['audience_text_minimum_required_pt'])

    def test_duplicate_multiline_cells_are_matched_by_location(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as td:
            p=Presentation();s=p.slides.add_slide(p.slide_layouts[6]);t=s.shapes.add_table(2,2,0,0, Inches(6), Inches(3));t.name='metrics'
            texts=['2025','profit\n10','2025','profit\n20']
            for i,text in enumerate(texts):t.table.cell(i//2,i%2).text=text
            path=Path(td)/'deck.pptx';p.save(path)
            units=[{'copy_id':str(i),'text':v} for i,v in enumerate(texts)]
            package={'logic_layer':{'slides':[{}]},'copy_layer':{'slides':[{'copy_units':units}]}}
            maps=[{'copy_id':str(i),'verification_method':'paragraph-exact'} for i in range(4)]
            plan={'slides':[{'copy_unit_map':maps,'medium_execution_contract':{'minimum_object_counts':{}}}]}
            with mock.patch('compare_package_to_pptx.validate_plan',return_value=[]):
                self.assertEqual(compare(package,plan,path),[])
                for i,m in enumerate(maps):m.update(verification_method='table-cell',native_location={'shape_name':'metrics','row':i//2,'column':i%2})
                self.assertEqual(compare(package,plan,path),[])
                maps[3]['native_location']['row']=0
                self.assertTrue(compare(package,plan,path))
                maps[3]['native_location']['row']=1
                t.table.cell(1,1).text='profit\n200';p.save(path)
                self.assertTrue(compare(package,plan,path))

    def test_twenty_mb_pptx_requires_optimization(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'large.pptx';Presentation().save(p)
            with zipfile.ZipFile(p,'a',compression=zipfile.ZIP_STORED) as z:z.writestr('ppt/media/oversize.bin',b'x'*20_000_000)
            result=inspect(p,'high-fidelity')
            self.assertIn('PPTX_REQUIRES_SIZE_OPTIMIZATION',[x['code'] for x in result['blockers']])
            self.assertEqual(result['largest_parts'][0]['path'],'ppt/media/oversize.bin')

    def test_evidence_is_lossless_deduplicated_and_tamper_evident(self):
        text='完整来源\n'*4000
        original={'work_report':{'source_snapshots':[{'raw_text':text},{'raw_text':text}]},
                  'stage_documents':{'previews':[{'base64':base64.b64encode(b'\x89PNG'+b'a'*100).decode()}]}}
        with tempfile.TemporaryDirectory() as td:
            root=Path(td); compact=compact_report(original,root)
            self.assertEqual(expand_report(compact,root),original)
            self.assertEqual(len(compact['evidence_storage']['files']),2)
            ref=compact['work_report']['source_snapshots'][0]['raw_text']['$evidence']
            target=root/ref['path'];target.write_bytes(b'changed')
            with self.assertRaisesRegex(ValueError,'hash'):expand_report(compact,root)

    def test_evidence_path_escape_is_rejected(self):
        with tempfile.TemporaryDirectory() as td:
            compact=compact_report({'raw_text':'x'*5000},td)
            ref=compact['raw_text']['$evidence'];ref['path']='../outside'
            with self.assertRaisesRegex(ValueError,'escapes'):expand_report(compact,td)

    def test_repeated_failure_stops_and_changed_input_can_retry(self):
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);i=root/'input';i.write_text('one');o=root/'output'
            command=[sys.executable,'-c','raise SystemExit(3)']
            for _ in range(2):
                with self.assertRaises(RuntimeError):run_check(root,'probe',command,[i],[o])
            with self.assertRaisesRegex(RuntimeError,'failed twice'):run_check(root,'probe',command,[i],[o])
            receipt=json.loads((root/'.clayz-checks/probe.json').read_text())
            self.assertIn('started_at',receipt);self.assertEqual(receipt['consecutive_failures'],2)
            i.write_text('two')
            with self.assertRaises(RuntimeError):run_check(root,'probe',command,[i],[o])
            self.assertEqual(json.loads((root/'.clayz-checks/probe.json').read_text())['consecutive_failures'],1)

    def test_external_evidence_real_publisher_and_verify(self):
        from tests.test_work_report_adversarial import _ReportRun
        from tests.test_calibrated_audit_adversarial import RealCliReleaseTests
        with mock.patch.object(RealCliReleaseTests,'report_storage','external-evidence',create=True):
            fixture=_ReportRun(subject='alpha',rich=True)
        try:
            stages=fixture.build_records_and_audit();bundle=fixture.publish(stages)
            report=bundle/'ppt-supervision-report.json';compact=json.loads(report.read_text())
            self.assertIn('evidence_storage',compact)
            self.assertIn('work_report',load_report(report))
            self.assertEqual((bundle/'work-report.md').read_text(),render_compact_markdown(compact))
            fixture.verify(bundle)
            ref=compact['evidence_storage']['files'][0];(bundle/ref['path']).write_bytes(b'corrupt')
            fixture.verify(bundle,expected=1)
        finally:fixture.close()

if __name__=='__main__':unittest.main()
