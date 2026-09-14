#!/usr/bin/env python3
"""Black-box acceptance tests for the report3.6 work-report delivery.

The tests deliberately compose the existing calibrated CLI fixture instead of
subclassing its ``TestCase``.  That keeps the real record/calibration/audit/
assembly/publish path while avoiding discovery and execution of the older
test methods.  All topic material is synthetic alpha/beta evidence.
"""

from __future__ import annotations

import io
import hashlib
import json
import subprocess
import tempfile
import unittest
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

import zipfile

import tests.test_calibrated_audit_adversarial as calibrated_helpers
from tests.work_report_fixtures import (
    build_report_pptx,
    file_ref,
    json_bytes,
    pptx_object_stats,
    sha256_file,
    write_json,
)


def _canonical_sha256(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _contains(value: Any, needle: str) -> bool:
    """Search structured output without imposing a prose layout on the core."""

    if isinstance(value, str):
        return needle in value
    if isinstance(value, dict):
        return any(_contains(key, needle) or _contains(item, needle) for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return any(_contains(item, needle) for item in value)
    return False


def _find_values(value: Any, keys: Iterable[str]) -> list[Any]:
    wanted = set(keys)
    found: list[Any] = []
    if isinstance(value, dict):
        for key, item in value.items():
            if key in wanted:
                found.append(item)
            found.extend(_find_values(item, wanted))
    elif isinstance(value, (list, tuple)):
        for item in value:
            found.extend(_find_values(item, wanted))
    return found


def _first_value(value: Any, keys: Iterable[str], default: Any = None) -> Any:
    values = _find_values(value, keys)
    return values[0] if values else default


def _stage_row(work_report: dict[str, Any], stage: str) -> dict[str, Any]:
    rows = work_report.get("stages")
    if not isinstance(rows, list):
        raise AssertionError("work_report.stages must be an array")
    row = next((item for item in rows if isinstance(item, dict) and item.get("stage") == stage), None)
    if not isinstance(row, dict):
        raise AssertionError(f"work_report has no {stage!r} stage row")
    return row


def _rewrite_presentation_package(path: Path, *, reverse_order: bool = False, remove_slide_relationship: bool = False) -> None:
    """Apply a narrow OOXML mutation for the independent inspector case."""

    presentation_name = "ppt/presentation.xml"
    relationships_name = "ppt/_rels/presentation.xml.rels"
    presentation_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    relationship_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ET.register_namespace("p", presentation_ns)
    ET.register_namespace("pr", relationship_ns)
    source = path.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source), "r") as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        presentation = ET.fromstring(archive.read(presentation_name))
        slide_ids = presentation.find(f"{{{presentation_ns}}}sldIdLst")
        if slide_ids is None:
            raise AssertionError("fixture has no presentation slide id list")
        if reverse_order:
            slide_ids[:] = list(reversed(list(slide_ids)))
        relationships = ET.fromstring(archive.read(relationships_name))
        if remove_slide_relationship:
            slide_relationship = next(
                (
                    item
                    for item in relationships
                    if item.get("Type", "").endswith("/slide")
                ),
                None,
            )
            if slide_relationship is None:
                raise AssertionError("fixture has no slide relationship to remove")
            relationships.remove(slide_relationship)
        for name in archive.namelist():
            if name == presentation_name:
                payload = ET.tostring(presentation, encoding="utf-8", xml_declaration=True)
            elif name == relationships_name:
                payload = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
            else:
                payload = archive.read(name)
            target.writestr(name, payload)
    path.write_bytes(output.getvalue())


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise AssertionError(f"expected JSON object: {path}")
    return value


class _ReportRun:
    """A disposable real CLI run assembled with the existing calibrated helper."""

    def __init__(self, *, subject: str = "alpha", rich: bool = True, missing_stage: str | None = None) -> None:
        self.subject = subject
        self.rich = rich
        self.missing_stage = missing_stage
        self.case = calibrated_helpers.RealCliReleaseTests("runTest")
        self.case.setUp()
        self.work = self.case.work

        # The existing fixture creates a synthetic task during setUp.  Replace
        # it before issuing a fresh challenge so alpha and beta have distinct
        # task hashes while retaining the helper's complete calibrated inputs.
        self.case.task.write_bytes(
            json_bytes(
                {
                    "request": f"Decide whether the synthetic {subject} pilot should proceed.",
                    "topic": subject,
                    "revision": 1,
                }
            )
        )
        self.case.challenge = self.case._write_challenge()
        self.case.preflight_path = self.case._write_preflight()
        build_report_pptx(self.case.pptx, subject=subject, rich=rich)

        self.note_paths: dict[str, Path] = {}
        for stage in ("logic", "copy", "art-direction", "output", "supervisor"):
            if stage == missing_stage:
                continue
            self.note_paths[stage] = self._write_note(stage)

    def close(self) -> None:
        self.case.tearDown()

    def _write_note(self, stage: str) -> Path:
        marker = f"{self.subject.upper()}-{stage.upper()}-REPORT-MARKER"
        common = {
            "case": self.subject,
            "stage": stage,
            "report_marker": marker,
            "task_question": f"Should the {self.subject} pilot proceed under the observed evidence?",
            "research": {
                "sources": [f"synthetic-source-{self.subject}-01"],
                "evidence": f"Observed synthetic {self.subject} signal in the bounded fixture.",
                "counterevidence": "The sample is small and may not generalize.",
            },
            "assumptions": ["The fixture is synthetic and the pilot remains reversible."],
            "excluded_options": ["Immediate full rollout was excluded because the sample is small."],
            "storyline": "Observed signal -> bounded pilot -> explicit limitation.",
            "design_tradeoff": "Keep the decision visible while retaining evidence detail in the report.",
            "calibration_response": f"{stage} absorbed the preceding synthetic calibration where applicable.",
            "actual_output": f"The {stage} artifact was written and hash-bound by the CLI.",
            "limitations": ["Native Office rendering is not claimed by the synthetic pixel evidence."],
        }
        if stage in {"copy", "output"}:
            path = self.work / f"{stage}-work-notes.md"
            lines = [
                f"# {self.subject.title()} {stage} work notes",
                "",
                f"report_marker: {marker}",
                f"task_question: Should the {self.subject} pilot proceed under the observed evidence?",
                "research_source: synthetic-source-" + self.subject + "-01",
                "counterevidence: The sample is small and may not generalize.",
                "assumption: The fixture is synthetic and the pilot remains reversible.",
                "excluded_option: Immediate full rollout was excluded because the sample is small.",
                "storyline: Observed signal -> bounded pilot -> explicit limitation.",
                "design_tradeoff: Keep the decision visible while retaining evidence detail in the report.",
                "limitation: Native Office rendering is not claimed by the synthetic pixel evidence.",
                "",
            ]
            path.write_text("\n".join(lines), encoding="utf-8", newline="\n")
            return path
        return write_json(self.work / f"{stage}-work-notes.json", common)

    def _draft(self, stage: str, roles: list[str]) -> Path:
        draft_name = f"{stage}-stage-record-draft.json" if stage == "supervisor" else f"{stage}-draft.json"
        return write_json(
            self.work / draft_name,
            {
                "summary": f"Recorded observable synthetic {stage} work for {self.subject}.",
                "decisions": [f"Keep the {self.subject} {stage} decision evidence-bound."],
                "checks": [{"name": "synthetic evidence inspected", "status": "pass", "evidence_roles": roles}],
                "open_issues": [],
            },
        )

    def _record_stage(
        self,
        stage: str,
        *,
        artifacts: dict[str, Path],
        previous: Path | None = None,
        calibration: Path | None = None,
    ) -> Path:
        roles = list(artifacts)
        if stage in self.note_paths:
            artifacts["work-notes"] = self.note_paths[stage]
            roles.append("work-notes")
        draft = self._draft(stage, roles)
        output = self.work / f"{stage}-record.json"
        args: list[object] = [
            "record-stage",
            "--stage",
            stage,
            "--draft",
            draft,
            "--challenge",
            self.case.challenge,
        ]
        if previous is not None:
            args.extend(("--previous-record", previous))
        for role, path in artifacts.items():
            args.extend(("--artifact", f"{role}={path}"))
        if calibration is not None:
            args.extend(("--calibration", calibration, "--calibration-dispositions", self.case.dispositions_path))
        args.extend(("--acceptance-contract", self.case.acceptance_path, "--output", output))
        self.case._run_cli(*args)
        return output

    def build_records_and_audit(self) -> dict[str, Path]:
        stages: dict[str, Path] = {}
        stages["logic"] = self._record_stage("logic", artifacts={"package": self.case.package_path})
        stages["calibration_logic_copy"] = self.work / "calibration-logic-copy.json"
        self.case._run_cli(
            "record-calibration",
            "--step",
            "logic-to-copy",
            "--source-record",
            stages["logic"],
            "--acceptance-contract",
            self.case.acceptance_path,
            "--shared-rules",
            self.case.supervisor_rules_path,
            "--findings",
            self.case.empty_findings_path,
            "--challenge",
            self.case.challenge,
            "--output",
            stages["calibration_logic_copy"],
        )
        stages["copy"] = self._record_stage(
            "copy",
            artifacts={"package": self.case.package_path},
            previous=stages["logic"],
            calibration=stages["calibration_logic_copy"],
        )
        stages["calibration_copy_art_direction"] = self.work / "calibration-copy-art-direction.json"
        self.case._run_cli(
            "record-calibration",
            "--step",
            "copy-to-art-direction",
            "--source-record",
            stages["copy"],
            "--acceptance-contract",
            self.case.acceptance_path,
            "--shared-rules",
            self.case.supervisor_rules_path,
            "--findings",
            self.case.empty_findings_path,
            "--challenge",
            self.case.challenge,
            "--previous-calibration",
            stages["calibration_logic_copy"],
            "--output",
            stages["calibration_copy_art_direction"],
        )
        stages["art-direction"] = self._record_stage(
            "art-direction",
            artifacts={"plan": self.case.plan_path},
            previous=stages["copy"],
            calibration=stages["calibration_copy_art_direction"],
        )
        stages["calibration_art_direction_output"] = self.work / "calibration-art-direction-output.json"
        self.case._run_cli(
            "record-calibration",
            "--step",
            "art-direction-to-output",
            "--source-record",
            stages["art-direction"],
            "--acceptance-contract",
            self.case.acceptance_path,
            "--shared-rules",
            self.case.supervisor_rules_path,
            "--findings",
            self.case.empty_findings_path,
            "--challenge",
            self.case.challenge,
            "--previous-calibration",
            stages["calibration_copy_art_direction"],
            "--output",
            stages["calibration_art_direction_output"],
        )
        stages["output"] = self._record_stage(
            "output",
            artifacts={"qa": self.case.qa_path, "inventory": self.case.inventory_path, "pptx": self.case.pptx},
            previous=stages["art-direction"],
            calibration=stages["calibration_art_direction_output"],
        )
        auditor = self.case._write_audit_inputs(stages=stages)
        stages["auditor"] = auditor
        draft = self.case._build_report(stages=stages, auditor_path=auditor)
        stages["supervisor"] = self._record_stage(
            "supervisor",
            artifacts={"draft": draft, "pptx": self.case.pptx, "auditor": auditor},
            previous=stages["output"],
        )
        stages["supervisor_draft"] = draft
        stages["report"] = self.case._assemble_report(stages, draft, auditor)
        return stages

    def publish(self, stages: dict[str, Path]) -> Path:
        output = self.work / "published"
        self.case._run_cli(
            self.case.package_path,
            self.case.plan_path,
            self.case.qa_path,
            self.case.inventory_path,
            stages["report"],
            "--pptx",
            self.case.pptx,
            "--runtime-preflight",
            self.case.preflight_path,
            "--config",
            self.case.config_path,
            "--render-root",
            self.case.render_root,
            "--output-dir",
            output,
        )
        return output

    def verify(self, bundle: Path, *, expected: int = 0, extra: list[object] | None = None) -> subprocess.CompletedProcess[str]:
        args: list[object] = ["verify-handoff", "--bundle", bundle]
        if extra:
            args.extend(extra)
        return self.case._run_cli(*args, expected=expected)


class WorkReportAdversarialTests(unittest.TestCase):
    maxDiff = None

    def _run(self, **kwargs: Any) -> tuple[_ReportRun, dict[str, Path], Path]:
        fixture = _ReportRun(**kwargs)
        try:
            stages = fixture.build_records_and_audit()
            bundle = fixture.publish(stages)
            # Return ownership to the test while keeping the temporary root
            # alive.  The caller must close the fixture in a finally block.
            return fixture, stages, bundle
        except Exception:
            fixture.close()
            raise

    @staticmethod
    def _report(bundle: Path) -> dict[str, Any]:
        return _read_json(bundle / "ppt-supervision-report.json")

    @staticmethod
    def _markdown(bundle: Path) -> Path:
        path = bundle / "work-report.md"
        if not path.is_file():
            raise AssertionError(f"publisher did not emit exact derived Markdown path: {path}")
        return path

    def test_real_cli_baseline_contains_richer_sources_and_derived_markdown(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            report = self._report(bundle)
            work_report = report.get("work_report")
            self.assertIsInstance(work_report, dict)
            self.assertEqual(work_report.get("contract"), "io.clayz.presentation.work-report/1.0")
            self.assertEqual(work_report.get("version"), "3.6")
            for section in (
                "task",
                "substantive_content",
                "storyline",
                "evidence",
                "copy",
                "art_direction",
                "calibrations",
                "stages",
                "actual_pptx",
                "auditor",
                "release",
                "source_snapshots",
                "provenance",
            ):
                self.assertIn(section, work_report)
            for marker in (
                "ALPHA-LOGIC-REPORT-MARKER",
                "ALPHA-COPY-REPORT-MARKER",
                "ALPHA-ART-DIRECTION-REPORT-MARKER",
                "ALPHA-OUTPUT-REPORT-MARKER",
                "ALPHA-SUPERVISOR-REPORT-MARKER",
                "Immediate full rollout was excluded",
                "The sample is small and may not generalize",
            ):
                self.assertTrue(_contains(work_report, marker), marker)
            for note_path in fixture.note_paths.values():
                snapshots = [
                    item
                    for item in work_report.get("source_snapshots", [])
                    if isinstance(item, dict) and item.get("path") == str(note_path.resolve())
                ]
                self.assertEqual(len(snapshots), 1)
                self.assertEqual(snapshots[0].get("raw_text"), note_path.read_text(encoding="utf-8"))
            auditor = work_report.get("auditor", {})
            self.assertTrue(_contains(auditor, "AUD-F-SYNTHETIC-RENDER"))
            self.assertTrue(_contains(auditor, "Native Office rendering was not executed"))
            self.assertEqual(auditor.get("raw_result"), _read_json(stages["auditor"]))
            self.assertEqual(report.get("work_report_sha256"), _canonical_sha256(work_report))

            markdown = self._markdown(bundle)
            markdown_text = markdown.read_text(encoding="utf-8")
            self.assertIn("ALPHA-LOGIC-REPORT-MARKER", markdown_text)
            self.assertIn("Immediate full rollout was excluded", markdown_text)
            self.assertIn(report["work_report_sha256"], markdown_text)

            manifest = _read_json(bundle / "delivery-manifest.json")
            derived = manifest.get("derived_files")
            self.assertIsInstance(derived, list)
            md_entry = next((item for item in derived if isinstance(item, dict) and item.get("path") == "work-report.md"), None)
            self.assertIsNotNone(md_entry)
            self.assertEqual(md_entry["sha256"], sha256_file(markdown))
            self.assertEqual(md_entry["bytes"], markdown.stat().st_size)

            result = fixture.verify(bundle)
            verified = json.loads(result.stdout)
            self.assertEqual(verified["status"], "handoff-verified")
            self.assertEqual(verified["task_identity"]["task_request_sha256"], report["task_request_sha256"])
            self.assertEqual(verified["full_assembly"], {"status": "validated", "records": 5})
            markdown_record = next(item for item in verified["artifacts"] if item["role"] == "work-report-markdown")
            self.assertEqual(Path(markdown_record["path"]).resolve(), markdown.resolve())
        finally:
            fixture.close()

    def test_missing_stage_notes_are_explicitly_not_recorded(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True, missing_stage="copy")
        try:
            report = self._report(bundle)
            work_report = report["work_report"]
            stage_value = _stage_row(work_report, "copy")
            self.assertEqual(stage_value.get("stage"), "copy")
            notes = [
                item
                for item in work_report.get("copy", {}).get("notes", [])
                if isinstance(item, dict) and item.get("stage") == "copy"
            ]
            self.assertTrue(notes)
            self.assertTrue(any(item.get("status") == "not-recorded" for item in notes))
            self.assertFalse(_contains(notes, "ALPHA-COPY-REPORT-MARKER"))
            self.assertFalse(_contains(work_report, "complete work history"))
            self.assertIn(report.get("run_status"), {"incomplete-evidence", "issues-found", "complete-with-deferred-acceptance"})
            result = fixture.verify(bundle)
            self.assertEqual(json.loads(result.stdout)["status"], "handoff-verified")
        finally:
            fixture.close()

    def test_actual_ooxml_order_notes_master_and_object_stats_are_bound(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            report = self._report(bundle)
            actual = report["work_report"]["actual_pptx"]
            expected = pptx_object_stats(fixture.case.pptx)
            self.assertEqual(_first_value(actual, {"page_count", "slide_count"}), expected["page_count"])
            self.assertEqual(actual.get("slide_order"), expected["slide_order"])
            totals = actual.get("totals")
            self.assertIsInstance(totals, dict)
            for key, expected_key in (
                ("charts", "charts"),
                ("tables", "tables"),
            ):
                self.assertIn(key, totals)
                self.assertEqual(totals[key], expected["totals"][expected_key], key)
            pages = actual.get("slides") or actual.get("pages")
            self.assertIsInstance(pages, list)
            self.assertEqual(len(pages), len(expected["slides"]))
            for observed, expected_page in zip(pages, expected["slides"]):
                self.assertEqual(observed.get("slide_index", observed.get("page", observed.get("slide_number"))), expected_page["page"])
                inventory = observed.get("inventory", {})
                self.assertEqual(inventory.get("tables", inventory.get("table_count")), expected_page["tables"])
                self.assertEqual(inventory.get("charts", inventory.get("chart_count")), expected_page["charts"])
                self.assertEqual(observed.get("notes", {}).get("text", ""), "alpha notes: research question, assumptions, and decision context.")
                self.assertIn("Alpha", json.dumps(observed, ensure_ascii=False))
                self.assertIn("Period", observed.get("page_text", ""))
                self.assertIn("Q2", observed.get("page_text", ""))
                self.assertEqual(observed.get("title_candidate"), "Alpha research brief")
            masters = actual.get("masters")
            self.assertIsInstance(masters, dict)
            self.assertEqual(masters.get("count"), 1)
            self.assertEqual(masters["parts"][0]["inventory"]["pictures"], 1)
            scope = actual.get("scope", {})
            self.assertEqual(scope.get("master_inheritance_observed"), "not-observed")
            self.assertFalse(scope.get("master_objects_in_slide_counts"))
            self.assertIsNot(scope.get("master_inheritance_verified"), True)
            result = fixture.verify(bundle)
            verified = json.loads(result.stdout)
            self.assertEqual(verified["presentation"]["slide_count"], expected["page_count"])
            self.assertEqual(verified["presentation"]["statistics"]["charts"], expected["totals"]["charts"])
        finally:
            fixture.close()

    def test_inspector_follows_reversed_presentation_order_and_slide_relationships(self) -> None:
        from packages.validators.work_report import WorkReportError, inspect_pptx

        from pptx import Presentation
        from pptx.chart.data import ChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches, Pt

        with tempfile.TemporaryDirectory(prefix="work-report-inspector-") as directory:
            root = Path(directory)
            pptx_path = root / "reordered.pptx"
            # Reuse the synthetic first-page builder, then add a second page
            # with independently identifiable text/notes and native objects.
            build_report_pptx(pptx_path, subject="order", rich=False)
            presentation = Presentation(pptx_path)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.7))
            title.name = "ORDER::SECOND-TITLE"
            run = title.text_frame.paragraphs[0].add_run()
            run.text = "SECOND-ORDER-TITLE"
            run.font.size = Pt(24)
            body = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12), Inches(0.7))
            body.name = "ORDER::SECOND-BODY"
            body.text_frame.text = "SECOND-ORDER-BODY"
            body.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
            slide.notes_slide.notes_text_frame.text = "SECOND-ORDER-NOTE"

            table_shape = slide.shapes.add_table(2, 2, Inches(0.6), Inches(2.1), Inches(4.0), Inches(1.5))
            table_shape.name = "ORDER::SECOND-TABLE"
            table_shape.table.cell(0, 0).text = "ORDER-CELL"
            table_shape.table.cell(0, 1).text = "VALUE"
            table_shape.table.cell(1, 0).text = "A"
            table_shape.table.cell(1, 1).text = "7"
            chart_data = ChartData()
            chart_data.categories = ["A", "B"]
            chart_data.add_series("ORDER-SERIES", (7, 9))
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED,
                Inches(5.0),
                Inches(2.0),
                Inches(5.5),
                Inches(3.0),
                chart_data,
            )
            presentation.save(pptx_path)
            _rewrite_presentation_package(pptx_path, reverse_order=True)

            observed = inspect_pptx(pptx_path)
            self.assertEqual(observed["slide_order"], ["ppt/slides/slide2.xml", "ppt/slides/slide1.xml"])
            self.assertEqual(observed["slides"][0]["slide_index"], 1)
            self.assertEqual(observed["slides"][0]["part"], "ppt/slides/slide2.xml")
            self.assertIn("SECOND-ORDER-TITLE", observed["slides"][0]["page_text"])
            self.assertIn("ORDER-CELL", observed["slides"][0]["page_text"])
            self.assertEqual(observed["slides"][0]["notes"]["text"], "SECOND-ORDER-NOTE")
            self.assertEqual(observed["slides"][0]["inventory"]["tables"], 1)
            self.assertEqual(observed["slides"][0]["inventory"]["charts"], 1)
            self.assertEqual(observed["slides"][1]["part"], "ppt/slides/slide1.xml")
            self.assertEqual(observed["slides"][1]["notes"]["text"], "order notes: research question, assumptions, and decision context.")
            self.assertEqual(observed["totals"]["tables"], 1)
            self.assertEqual(observed["totals"]["charts"], 1)

            missing_relationship_path = root / "missing-relationship.pptx"
            missing_relationship_path.write_bytes(pptx_path.read_bytes())
            _rewrite_presentation_package(missing_relationship_path, remove_slide_relationship=True)
            with self.assertRaises(WorkReportError):
                inspect_pptx(missing_relationship_path)

    def test_pure_text_beta_has_different_actual_stats_without_art_shape_quota(self) -> None:
        fixture, stages, bundle = self._run(subject="beta", rich=False)
        try:
            report = self._report(bundle)
            actual = report["work_report"]["actual_pptx"]
            expected = pptx_object_stats(fixture.case.pptx)
            totals = actual.get("totals")
            self.assertEqual(totals["charts"], 0)
            self.assertEqual(totals["tables"], 0)
            self.assertEqual(actual["masters"]["parts"][0]["inventory"]["pictures"], 0)
            self.assertEqual(_first_value(actual, {"page_count", "slide_count"}), expected["page_count"])
            self.assertNotIn("11 native chart", json.dumps(report, ensure_ascii=False).casefold())
        finally:
            fixture.close()

    def test_legacy_seven_field_bank_summary_cannot_be_published(self) -> None:
        fixture = _ReportRun(subject="alpha", rich=True)
        try:
            stages = fixture.build_records_and_audit()
            legacy = {
                "run_id": _read_json(fixture.case.challenge)["run_id"],
                "task_request_sha256": _read_json(fixture.case.challenge)["task_request_sha256"],
                "status": "supervised",
                "title": "Bank performance summary",
                "summary": "Revenue, margin, liquidity, and recommendation.",
                "charts": 11,
                "tables": 1,
            }
            stages["report"].write_bytes(json_bytes(legacy))
            output = fixture.work / "legacy-rejected"
            result = fixture.case._run_cli(
                fixture.case.package_path,
                fixture.case.plan_path,
                fixture.case.qa_path,
                fixture.case.inventory_path,
                stages["report"],
                "--pptx",
                fixture.case.pptx,
                "--runtime-preflight",
                fixture.case.preflight_path,
                "--config",
                fixture.case.config_path,
                "--render-root",
                fixture.case.render_root,
                "--output-dir",
                output,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), r"work.?report|assembly|3\.6|auditor")
        finally:
            fixture.close()

    def test_same_named_alpha_beta_bundles_cannot_exchange_report_and_pptx(self) -> None:
        alpha, alpha_stages, alpha_bundle = self._run(subject="alpha", rich=True)
        beta, beta_stages, beta_bundle = self._run(subject="beta", rich=False)
        try:
            swapped = beta.work / "swapped-rejected"
            result = beta.case._run_cli(
                beta.case.package_path,
                beta.case.plan_path,
                beta.case.qa_path,
                beta.case.inventory_path,
                alpha_stages["report"],
                "--pptx",
                beta.case.pptx,
                "--runtime-preflight",
                beta.case.preflight_path,
                "--config",
                beta.case.config_path,
                "--render-root",
                beta.case.render_root,
                "--output-dir",
                swapped,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(swapped.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), "task|run|hash|pptx|bundle|assembly")

            # Same basenames are intentional; identity must come from bound
            # task/run/hash values rather than filename selection.
            self.assertEqual((alpha_bundle / "final.pptx").name, (beta_bundle / "final.pptx").name)
            self.assertEqual((alpha_bundle / "ppt-supervision-report.json").name, (beta_bundle / "ppt-supervision-report.json").name)
            self.assertNotEqual(_read_json(alpha_bundle / "ppt-supervision-report.json")["task_request_sha256"], _read_json(beta_bundle / "ppt-supervision-report.json")["task_request_sha256"])
            (beta_bundle / "ppt-supervision-report.json").write_bytes(
                (alpha_bundle / "ppt-supervision-report.json").read_bytes()
            )
            swapped_bundle_result = beta.verify(beta_bundle, expected=1)
            self.assertNotEqual(swapped_bundle_result.returncode, 0)
            self.assertRegex((swapped_bundle_result.stdout + swapped_bundle_result.stderr).casefold(), "hash|manifest|task|assembly|record")
        finally:
            alpha.close()
            beta.close()

    def test_stage_notes_mutation_is_rejected_even_after_self_rehash(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            note = fixture.note_paths["logic"]
            note.write_text(note.read_text(encoding="utf-8") + "\nFORGED AFTER ASSEMBLY\n", encoding="utf-8", newline="\n")
            record_path = stages["logic"]
            record = _read_json(record_path)
            record["artifacts"]["work-notes"] = file_ref(note)
            from packages.validators.stage_work_records import canonical_record_sha256

            record["record_sha256"] = canonical_record_sha256(record)
            record_path.write_bytes(json_bytes(record))
            report_path = stages["report"]
            report = _read_json(report_path)
            report["work_records"][0] = record
            report["assembly"]["record_set_sha256"] = _canonical_sha256(report["work_records"])
            report["assembly"]["report_content_sha256"] = _canonical_sha256({key: value for key, value in report.items() if key != "assembly"})
            report_path.write_bytes(json_bytes(report))
            output = fixture.work / "stale-notes-rejected"
            result = fixture.case._run_cli(
                fixture.case.package_path,
                fixture.case.plan_path,
                fixture.case.qa_path,
                fixture.case.inventory_path,
                report_path,
                "--pptx",
                fixture.case.pptx,
                "--runtime-preflight",
                fixture.case.preflight_path,
                "--config",
                fixture.case.config_path,
                "--render-root",
                fixture.case.render_root,
                "--output-dir",
                output,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), "note|stale|hash|source|work")
        finally:
            fixture.close()

    def test_deleting_source_snapshots_cannot_be_hidden_by_rehash(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            report_path = stages["report"]
            report = _read_json(report_path)
            report["work_report"]["source_snapshots"] = []
            report["work_report"]["stage_record_refs"] = []
            report["work_report"]["stage_artifact_bindings"] = []
            report["work_report_sha256"] = _canonical_sha256(report["work_report"])
            report["assembly"]["report_content_sha256"] = _canonical_sha256(
                {key: value for key, value in report.items() if key != "assembly"}
            )
            report_path.write_bytes(json_bytes(report))
            output = fixture.work / "source-snapshots-deletion-rejected"
            result = fixture.case._run_cli(
                fixture.case.package_path,
                fixture.case.plan_path,
                fixture.case.qa_path,
                fixture.case.inventory_path,
                report_path,
                "--pptx",
                fixture.case.pptx,
                "--runtime-preflight",
                fixture.case.preflight_path,
                "--config",
                fixture.case.config_path,
                "--render-root",
                fixture.case.render_root,
                "--output-dir",
                output,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), "source|snapshot|work.?report|record|assembly")
        finally:
            fixture.close()

    def test_handwritten_or_drifted_markdown_is_rejected_by_verify_handoff(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            markdown = self._markdown(bundle)
            markdown.write_text(markdown.read_text(encoding="utf-8") + "\nHANDWRITTEN DRIFT\n", encoding="utf-8", newline="\n")
            result = fixture.verify(bundle, expected=1)
            self.assertNotEqual(result.returncode, 0)
            self.assertRegex((result.stdout + result.stderr).casefold(), "markdown|derived|hash|drift")
        finally:
            fixture.close()

    def test_missing_stage_details_cannot_claim_complete_after_rehash(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            report_path = stages["report"]
            report = _read_json(report_path)
            work_report = report["work_report"]
            work_report["stages"] = [
                item
                for item in work_report.get("stages", [])
                if not (isinstance(item, dict) and item.get("stage") == "copy")
            ]
            work_report["status"] = "complete"
            report["work_report_sha256"] = _canonical_sha256(work_report)
            report["assembly"]["report_content_sha256"] = _canonical_sha256({key: value for key, value in report.items() if key != "assembly"})
            report_path.write_bytes(json_bytes(report))
            output = fixture.work / "incomplete-claim-rejected"
            result = fixture.case._run_cli(
                fixture.case.package_path,
                fixture.case.plan_path,
                fixture.case.qa_path,
                fixture.case.inventory_path,
                report_path,
                "--pptx",
                fixture.case.pptx,
                "--runtime-preflight",
                fixture.case.preflight_path,
                "--config",
                fixture.case.config_path,
                "--render-root",
                fixture.case.render_root,
                "--output-dir",
                output,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), "stage|complete|missing|work.?report|assembly")
        finally:
            fixture.close()

    def test_declared_chart_count_disagrees_with_actual_ooxml_is_rejected(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            report_path = stages["report"]
            report = _read_json(report_path)
            actual = report["work_report"]["actual_pptx"]
            totals = actual.get("totals") or actual.get("statistics")
            self.assertIsInstance(totals, dict)
            totals["charts"] = 11
            report["work_report_sha256"] = _canonical_sha256(report["work_report"])
            report["assembly"]["report_content_sha256"] = _canonical_sha256({key: value for key, value in report.items() if key != "assembly"})
            report_path.write_bytes(json_bytes(report))
            output = fixture.work / "chart-mismatch-rejected"
            result = fixture.case._run_cli(
                fixture.case.package_path,
                fixture.case.plan_path,
                fixture.case.qa_path,
                fixture.case.inventory_path,
                report_path,
                "--pptx",
                fixture.case.pptx,
                "--runtime-preflight",
                fixture.case.preflight_path,
                "--config",
                fixture.case.config_path,
                "--render-root",
                fixture.case.render_root,
                "--output-dir",
                output,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), "chart|stat|ooxml|pptx|actual")
        finally:
            fixture.close()

    def test_deleting_actual_page_body_cannot_be_hidden_by_rehash(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            report_path = stages["report"]
            report = _read_json(report_path)
            actual = report["work_report"]["actual_pptx"]
            pages = actual.get("slides")
            self.assertIsInstance(pages, list)
            self.assertTrue(pages)
            pages[0]["body"] = ""
            pages[0]["page_text"] = pages[0].get("title_candidate", "")
            report["work_report_sha256"] = _canonical_sha256(report["work_report"])
            report["assembly"]["report_content_sha256"] = _canonical_sha256(
                {key: value for key, value in report.items() if key != "assembly"}
            )
            report_path.write_bytes(json_bytes(report))
            output = fixture.work / "page-body-deletion-rejected"
            result = fixture.case._run_cli(
                fixture.case.package_path,
                fixture.case.plan_path,
                fixture.case.qa_path,
                fixture.case.inventory_path,
                report_path,
                "--pptx",
                fixture.case.pptx,
                "--runtime-preflight",
                fixture.case.preflight_path,
                "--config",
                fixture.case.config_path,
                "--render-root",
                fixture.case.render_root,
                "--output-dir",
                output,
                expected=1,
            )
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(output.exists())
            self.assertRegex((result.stdout + result.stderr).casefold(), "actual|ooxml|page|body|pptx|work.?report")
        finally:
            fixture.close()

    def test_final_path_mismatch_is_rejected_by_verify_handoff(self) -> None:
        fixture, stages, bundle = self._run(subject="alpha", rich=True)
        try:
            manifest_path = bundle / "delivery-manifest.json"
            manifest = _read_json(manifest_path)
            manifest["files"][0]["path"] = "wrong-final.pptx"
            manifest_path.write_bytes(json_bytes(manifest))
            result = fixture.verify(bundle, expected=1)
            self.assertNotEqual(result.returncode, 0)
            self.assertRegex((result.stdout + result.stderr).casefold(), "path|manifest|pptx|missing|hash")
        finally:
            fixture.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
