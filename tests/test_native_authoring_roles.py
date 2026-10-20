import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from packages.adapters.python_pptx import render as adapter
from packages.layout.native_chart_text import set_manual_legend_colors
from packages.layout.native_page import add_native_page


class NativeAuthoringRolesTest(unittest.TestCase):
    def test_selected_blank_page_has_no_implicit_content(self):
        prs = Presentation()
        slide = add_native_page(prs, layout=prs.slide_layouts[6], background='F8FAF9')
        self.assertEqual(len(slide.shapes), 0)
        self.assertEqual(str(slide.background.fill.fore_color.rgb), 'F8FAF9')
        # Layout selection must come from the caller, rather than a title role.
        with self.assertRaises(TypeError):
            add_native_page(prs)

    def test_caller_selected_layout_is_not_stripped_to_blank(self):
        prs = Presentation()
        chosen = prs.slide_layouts[0]
        expected = len([p for p in chosen.placeholders
                        if p.placeholder_format.type not in {13, 15, 16}])
        slide = add_native_page(prs, layout=chosen)
        self.assertEqual(slide.slide_layout, chosen)
        self.assertEqual(len(slide.shapes), expected)

    def test_explicit_heading_geometry_survives_render_and_reopen(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            manifest = {'contract': 'io.clayz.presentation.render-manifest/1.0',
                        'slides': [
                            {'slide_id': 'empty', 'objects': []},
                            {'slide_id': 'positioned', 'objects': [{
                                'object_id': 'heading', 'type': 'text', 'text': 'Explicit heading',
                                'options': {'x': 4.1, 'y': .18, 'w': 4.5, 'h': .65,
                                            'fontFace': 'STKaiti', 'fontSize': 25}}]}]}
            source = root / 'input.json'; source.write_text(json.dumps(manifest))
            target = root / 'actual.pptx'
            with patch.object(adapter, 'add_native_page', wraps=adapter.add_native_page) as entry:
                adapter.render(source, target)
                self.assertEqual(entry.call_count, 2)
            reopened = Presentation(target)
            self.assertEqual(len(reopened.slides[0].shapes), 0)
            self.assertEqual(len(reopened.slides[1].shapes), 1)
            text = reopened.slides[1].shapes[0]
            self.assertEqual(text.text, 'Explicit heading')
            self.assertEqual(text.left, Inches(4.1))
            self.assertEqual(text.top, Inches(.18))
            self.assertEqual(text.width, Inches(4.5))
            self.assertEqual(text.height, Inches(.65))

    def test_legend_mark_and_words_preserve_independent_choices(self):
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        mark = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(.2), Inches(.1))
        label = slide.shapes.add_textbox(Inches(1.4), Inches(.9), Inches(2.3), Inches(.5))
        label.text = 'Historical period'; run = label.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(19); run.font.bold = True
        before = (label.left, label.top, label.width, label.height)
        set_manual_legend_colors(mark, label.text_frame, mark_color='D8E0E2', text_color='667782')
        self.assertEqual(str(mark.fill.fore_color.rgb), 'D8E0E2')
        self.assertEqual(str(run.font.color.rgb), '667782')
        self.assertEqual((label.left, label.top, label.width, label.height), before)
        self.assertEqual(run.font.size.pt, 19)
        self.assertTrue(run.font.bold)
        # The helper does not pick text color or force inequality/contrast.
        with self.assertRaises(TypeError):
            set_manual_legend_colors(mark, label.text_frame, mark_color='D8E0E2')
        set_manual_legend_colors(mark, label.text_frame, mark_color='197F86', text_color='197F86')
        self.assertEqual(str(run.font.color.rgb), '197F86')


if __name__ == '__main__':
    unittest.main()
