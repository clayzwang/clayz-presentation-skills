"""Small, explicit fixtures for the calibrated-audit acceptance tests.

The presentation is a real OOXML package created with python-pptx.  The PNG
helper creates a real raster image with Pillow and marks it as synthetic in
the PNG metadata; it is intentionally not presented as a native PowerPoint
render.
"""

from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def json_bytes(value: Any) -> bytes:
    import json

    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n").encode("utf-8")


def write_json(path: Path, value: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(json_bytes(value))
    return path


def file_ref(path: Path, *, role: str | None = None) -> dict[str, Any]:
    data = path.read_bytes()
    result: dict[str, Any] = {
        "path": str(path.resolve()),
        "sha256": sha256_bytes(data),
        "bytes": len(data),
    }
    if role:
        result["role"] = role
    return result


def build_minimal_pptx(
    path: Path,
    *,
    run_font: str | None = "Aptos",
    east_asian_font: str | None = "Aptos",
    title_text: str = "校准验收真实 PPTX",
    body_text: str = "The synthetic fixture is editable and byte-bound.",
) -> Path:
    """Create a valid one-slide OOXML deck with a Chinese text run.

    python-pptx does not consistently write East Asian run fonts, so the
    optional values are applied to the generated XML after saving.  That
    keeps the fixture realistic while making the adversarial XML explicit.
    """

    from pptx import Presentation
    from pptx.util import Inches, Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.6), Inches(12.0), Inches(1.2))
    box.name = "COPY::C-S01-01"
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = title_text
    run.font.size = Pt(28)
    if run_font:
        run.font.name = run_font
    body = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(12.0), Inches(1.0))
    body.name = "COPY::C-S01-02"
    body_run = body.text_frame.paragraphs[0].add_run()
    body_run.text = body_text
    body_run.font.size = Pt(20)
    if run_font:
        body_run.font.name = run_font
    presentation.save(path)

    _rewrite_text_run_fonts(path, run_font=run_font, east_asian_font=east_asian_font)
    return path


def _rewrite_text_run_fonts(path: Path, *, run_font: str | None, east_asian_font: str | None) -> None:
    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ET.register_namespace("a", namespace)
    source = path.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source), "r") as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in archive.namelist():
            payload = archive.read(name)
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                root = ET.fromstring(payload)
                for run in root.findall(f".//{{{namespace}}}r"):
                    properties = run.find(f"{{{namespace}}}rPr")
                    if properties is None:
                        properties = ET.Element(f"{{{namespace}}}rPr")
                        run.insert(0, properties)
                    for child_tag, value in (("latin", run_font), ("ea", east_asian_font)):
                        old = properties.find(f"{{{namespace}}}{child_tag}")
                        if old is not None:
                            properties.remove(old)
                        if value:
                            properties.append(ET.Element(f"{{{namespace}}}{child_tag}", {"typeface": value}))
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(name, payload)
    path.write_bytes(output.getvalue())


def patch_theme_east_asian_font(path: Path, typeface: str) -> Path:
    """Patch every theme's major/minor East Asian typeface for font tests."""

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ET.register_namespace("a", namespace)
    source = path.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source), "r") as archive, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in archive.namelist():
            payload = archive.read(name)
            if name.startswith("ppt/theme/theme") and name.endswith(".xml"):
                root = ET.fromstring(payload)
                for node in root.findall(f".//{{{namespace}}}majorFont/{{{namespace}}}ea") + root.findall(f".//{{{namespace}}}minorFont/{{{namespace}}}ea"):
                    node.set("typeface", typeface)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(name, payload)
    path.write_bytes(output.getvalue())
    return path


def build_synthetic_pixel_render(path: Path, *, label: str = "synthetic render") -> Path:
    """Write a non-empty PNG with metadata declaring its synthetic origin."""

    from PIL import Image, ImageDraw, PngImagePlugin

    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (640, 360), (247, 248, 252))
    draw = ImageDraw.Draw(image)
    draw.rectangle((32, 32, 608, 328), outline=(91, 91, 214), width=4)
    draw.rectangle((72, 180, 220, 280), fill=(91, 91, 214))
    draw.rectangle((264, 132, 412, 280), fill=(20, 134, 109))
    draw.rectangle((456, 84, 604, 280), fill=(194, 65, 86))
    draw.text((48, 48), label, fill=(23, 26, 43))
    info = PngImagePlugin.PngInfo()
    info.add_text("render_origin", "synthetic-test-image")
    info.add_text("native_render", "false")
    info.add_text("purpose", "calibrated-audit acceptance fixture")
    image.save(path, format="PNG", pnginfo=info)
    return path


def assert_real_png(path: Path) -> None:
    """Validate bytes and pixel variance without claiming native rendering."""

    from PIL import Image

    raw = path.read_bytes()
    if not raw.startswith(b"\x89PNG\r\n\x1a\n"):
        raise AssertionError("fixture is not a PNG")
    with Image.open(path) as image:
        if image.size[0] < 2 or image.size[1] < 2:
            raise AssertionError("fixture PNG has no useful pixels")
        if len(set(image.convert("RGB").getdata())) < 4:
            raise AssertionError("fixture PNG is pixel-empty")
        if image.info.get("native_render") != "false":
            raise AssertionError("fixture must declare that it is not a native render")
