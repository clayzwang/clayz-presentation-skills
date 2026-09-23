"""Observable handoff regressions, using synthetic content and real image bytes."""
from __future__ import annotations
import base64
import copy
import hashlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'packages' / 'validators'))
sys.path.insert(0, str(ROOT / 'scripts'))
from story_handoff import (digest, spec_digest, validate_copy_trace, validate_visual_baseline,
                           validate_output_baseline, validate_design_audit, build_stage_documents,
                           handoff_archive_bytes, embed_final_renders, validate_embedded_documents)
from validate_logic_package import validate_package as validate_logic
from validate_ppt_package import validate_package as validate_copy
from validate_art_direction_plan import validate_plan


def reference(path):
    payload = path.read_bytes()
    return {'path': str(path.resolve()), 'sha256': hashlib.sha256(payload).hexdigest(), 'bytes': len(payload)}


def synthetic_handoff(folder):
    from PIL import Image, ImageDraw
    package = json.loads((ROOT / 'tests/fixtures/synthetic-copy-package.json').read_text(encoding='utf-8'))
    plan = json.loads((ROOT / 'tests/fixtures/synthetic-art-direction-plan.json').read_text(encoding='utf-8'))
    package['contract_version'] = '3.0'
    layer = package['logic_layer']
    story = {key: copy.deepcopy(layer[key]) for key in ('sources','glossary','metric_dictionary','open_items')}
    story.update(title='Synthetic pilot review', thesis=layer['deck_message_tree']['root_claim'],
                 audience='Pilot manager', desired_outcome='Understand the pilot before expansion',
                 opening='A synthetic pilot tests a limited operating change.',
                 conclusion='The observed result supports a bounded follow-up, with uncertainty retained.',
                 invariants=[], chapters=[{'chapter_id':'C1','title':'Pilot evidence','purpose':'Explain the observed result.',
                 'transition':'The observed evidence informs the next bounded decision.',
                 'blocks':[{'story_id':'B1','text':layer['slides'][0]['claim'],
                            'claim_status':'source-fact','source_ids':['SRC-SYN-01'], 'must_preserve':True,
                            'qualifiers':['Synthetic test evidence only.']}]}])
    package['story'] = story
    origin = copy.deepcopy(package)
    origin.update(status='logic-approved', logic_layer=None, copy_layer=None)
    path = folder / 'logic.json'
    path.write_text(json.dumps(origin,ensure_ascii=False),encoding='utf-8')
    package['logic_artifact'] = reference(path)
    package['copy_layer'].update(pagination_owner='copy', story_sha256=digest(story), chapter_order=['C1'],
                                semantic_preservation_review='Synthetic caveats, numbers and relationships remain unchanged.')
    for page in package['copy_layer']['slides']:
        for unit in page['copy_units']+page['speaker_notes']:
            unit['source_story_ids']=['B1']
    image_path = folder/'design.png'
    image = Image.new('RGB',(960,540),'white')
    ImageDraw.Draw(image).text((40,40),'Synthetic pilot review — test fixture',fill='black')
    image.save(image_path)
    slide = package['copy_layer']['slides'][0]
    elements = []
    for i, unit in enumerate(slide['copy_units']):
        elements.append({'element_id':f'E{i}', 'kind':'text','purpose':unit['role'],'copy_ids':[unit['copy_id']],
                         'box':[0.05,0.05,0.9,0.08], 'native_type':'text','group_id':'page','alignment':'left',
                         'typography':{'font_family':'Arial','size_pt':20},'locked_properties':['text','hierarchy'],
                         'allowed_adjustments':{'box_delta_max':0.01}})
    plan.update(contract_version='2.0',package_contract_version='3.0')
    plan['visual_baseline']={'status':'locked','locked_at':'2026-09-20T01:00:00+00:00',
        'copy_package_sha256':digest(package), 'spec_sha256':spec_digest(plan),
        'slides':[{'slide_id':slide['slide_id'],'image':reference(image_path),'first_visual':'Pilot result',
                   'reading_path':'Title, evidence, conclusion.', 'legibility_review':'Fixture image decoded; visual quality is not asserted.',
                   'copy_and_data_review':'Synthetic metadata test; not a production approval.',
                   'spec_consistency_review':'Fixture tags cover all copy IDs.', 'elements':elements}]}
    qa={'visual_baseline_sha256':digest(plan['visual_baseline']),'output_started_at':'2026-09-20T01:01:00+00:00'}
    return origin, package, plan, qa


class StoryVisualHandoffTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder=Path(self.temp.name)
        self.origin,self.package,self.plan,self.qa=synthetic_handoff(self.folder)

    def test_complete_story_before_pagination_and_copy_projection(self):
        self.assertEqual([],validate_logic(self.origin))
        self.assertEqual([],validate_copy(self.package))
        self.assertEqual([],validate_plan(self.package,self.plan))

    def test_logic_cannot_lock_pages(self):
        self.origin['logic_layer']=self.package['logic_layer']
        self.assertTrue(any('not a locked page' in e for e in validate_logic(self.origin)))

    def test_copy_cannot_change_original_story(self):
        self.package['story']['conclusion']='A stronger unsupported conclusion.'
        self.package['copy_layer']['story_sha256']=digest(self.package['story'])
        self.assertTrue(any('immutable Logic story' in e for e in validate_copy_trace(self.package)))

    def test_missing_original_logic_is_rejected(self):
        self.package['logic_artifact']['path']=str(self.folder/'missing.json')
        self.assertTrue(validate_copy_trace(self.package))

    def test_source_trace_and_required_coverage(self):
        for page in self.package['copy_layer']['slides']:
            for unit in page['copy_units']+page['speaker_notes']:
                unit['source_story_ids']=['unknown']
        errors=validate_copy_trace(self.package)
        self.assertTrue(any('required story block' in e for e in errors))

    def test_preview_bytes_and_all_pages_are_required(self):
        (self.folder/'design.png').write_bytes(b'not an image')
        self.assertTrue(validate_visual_baseline(self.package,self.plan))
        self.plan['visual_baseline']['slides']=[]
        self.assertTrue(validate_visual_baseline(self.package,self.plan))

    def test_spec_changes_invalidate_lock(self):
        self.plan['decision_log']['changed']='layout'
        self.assertTrue(any('specification hash' in e for e in validate_visual_baseline(self.package,self.plan)))

    def test_copy_changes_invalidate_preview(self):
        self.package['copy_layer']['slides'][0]['copy_units'][0]['text']='Changed title'
        self.assertTrue(any('stale Copy' in e for e in validate_visual_baseline(self.package,self.plan)))

    def test_missing_element_and_outside_canvas_rejected(self):
        self.plan['visual_baseline']['slides'][0]['elements'].pop()
        self.plan['visual_baseline']['slides'][0]['elements'][0]['box']=[0.9,0,0.9,0.1]
        errors=validate_visual_baseline(self.package,self.plan)
        self.assertTrue(any('every Copy unit' in e for e in errors))
        self.assertTrue(any('fit normalized canvas' in e for e in errors))

    def test_output_before_design_and_stale_baseline_rejected(self):
        self.qa['output_started_at']='2026-09-20T00:59:00+00:00'
        self.qa['visual_baseline_sha256']='0'*64
        errors=validate_output_baseline(self.package,self.plan,self.qa)
        self.assertTrue(any('before' in e for e in errors))
        self.assertTrue(any('exact Art Direction' in e for e in errors))

    def report(self):
        from pptx import Presentation
        pptx=self.folder/'deck.pptx'
        presentation=Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(pptx)
        sha=reference(pptx)['sha256']
        page={'slide_id':self.plan['slides'][0]['slide_id'], 'preview_sha256':reference(self.folder/'design.png')['sha256'],
              'status':'reviewed','final_render':reference(self.folder/'design.png'),'rendered_from_pptx_sha256':sha}
        for key in ('content','visual_fidelity','native_editability','design_quality'):
            page[key]={'status':'pass','observation':'Synthetic binding test only; no production visual claim.'}
        report={'run_status':'clean','issues':[], 'stage_documents':build_stage_documents(self.package,self.plan),
                'design_comparison':{'visual_baseline_sha256':digest(self.plan['visual_baseline']),'pptx_sha256':sha,'slides':[page]}}
        embed_final_renders(report)
        return report,pptx

    def test_audit_embeds_actual_documents_and_comparison(self):
        report,pptx=self.report()
        self.assertEqual([],validate_design_audit(self.package,self.plan,self.qa,report,pptx))
        report['stage_documents']['documents']['logic']['markdown']='Reconstructed summary'
        self.assertTrue(validate_design_audit(self.package,self.plan,self.qa,report,pptx))

    def test_matching_bad_design_is_not_clean(self):
        report,pptx=self.report()
        report['design_comparison']['slides'][0]['design_quality']={'status':'fail','observation':'Misleading scale.',
            'earliest_owner':'art-direction','issue_ids':['I1']}
        report['issues']=[{'issue_id':'I1'}]
        self.assertTrue(any('cannot be reported clean' in e for e in validate_design_audit(self.package,self.plan,self.qa,report,pptx)))
        report['run_status']='issues-found'
        self.assertEqual([],validate_design_audit(self.package,self.plan,self.qa,report,pptx))

    def test_deferred_render_cannot_claim_clean(self):
        report,pptx=self.report()
        report['design_comparison']['slides'][0].update(status='deferred',reason='No render route.')
        self.assertTrue(validate_design_audit(self.package,self.plan,self.qa,report,pptx))
        report['run_status']='complete-with-deferred-acceptance'
        self.assertEqual([],validate_design_audit(self.package,self.plan,self.qa,report,pptx))

    def test_portable_companion_survives_source_removal(self):
        report,_=self.report()
        first=handoff_archive_bytes(report)
        (self.folder/'logic.json').unlink()
        (self.folder/'design.png').unlink()
        self.assertEqual([],validate_embedded_documents(report))
        self.assertEqual(first,handoff_archive_bytes(report))
        with zipfile.ZipFile(io.BytesIO(first)) as archive:
            self.assertIn('comparison.html',archive.namelist())
            self.assertIn('logic.md',archive.namelist())
            self.assertIn('copy.md',archive.namelist())
            self.assertIn('art_direction.md',archive.namelist())
            self.assertIn('images/001-final.png',archive.namelist())

    def test_portable_tampered_preview_rejected(self):
        report,_=self.report()
        report['stage_documents']['previews'][0]['base64']=base64.b64encode(b'bad').decode()
        with self.assertRaises(ValueError):
            handoff_archive_bytes(report)


if __name__=='__main__':
    unittest.main()
