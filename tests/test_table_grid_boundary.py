import tempfile
import unittest
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches
from packages.layout.semantic_table import table_grid_boundary


class ActualTableGrid(unittest.TestCase):
    def table(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(6), Inches(3))
        for c, w in zip(shape.table.columns, [2.4, 1.1, 2.5]): c.width = Inches(w)
        for r, h in zip(shape.table.rows, [.6, 1.1, 1.3]): r.height = Inches(h)
        return prs, slide, shape

    def test_unequal_grid_and_outside_edges_read_without_mutation(self):
        _, _, s = self.table()
        before = etree.tostring(s._element)
        x = int(s.left) + int(s.table.columns[0].width)
        self.assertEqual(table_grid_boundary(s, axis="column", boundary_index=1),
                         ((x, int(s.top)), (x, int(s.top+s.height))))
        self.assertNotEqual(x, int(s.left+s.width//3))
        y = int(s.top) + int(s.table.rows[0].height)
        self.assertEqual(table_grid_boundary(s, axis="row", boundary_index=1),
                         ((int(s.left), y), (int(s.left+s.width), y)))
        self.assertEqual(table_grid_boundary(s, axis="column", boundary_index=0)[0][0], s.left)
        self.assertEqual(table_grid_boundary(s, axis="column", boundary_index=3)[0][0], s.left+s.width)
        self.assertEqual(table_grid_boundary(s, axis="row", boundary_index=3)[0][1], s.top+s.height)
        self.assertEqual(etree.tostring(s._element), before)

    def test_resize_translation_and_save_reopen_use_current_grid(self):
        prs, _, s = self.table()
        old = table_grid_boundary(s, axis="column", boundary_index=2)
        s.left += 271828
        s.top += 314159
        s.table.columns[0].width += 37
        current = table_grid_boundary(s, axis="column", boundary_index=2)
        self.assertEqual(current[0], (old[0][0]+271828+37, old[0][1]+314159))
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)/"grid.pptx"
            prs.save(p)
            reopened = Presentation(p).slides[0].shapes[0]
            self.assertEqual(table_grid_boundary(reopened, axis="column", boundary_index=2), current)

    def test_merged_cell_keeps_grid_without_claiming_visible_border(self):
        _, _, s = self.table()
        old = table_grid_boundary(s, axis="column", boundary_index=1)
        s.table.cell(0, 0).merge(s.table.cell(0, 1))
        self.assertEqual(table_grid_boundary(s, axis="column", boundary_index=1), old)

    def test_invalid_axis_index_and_non_table(self):
        _, slide, s = self.table()
        for axis, i, exc in [("time", 1, ValueError), ("column", -1, ValueError),
                             ("row", 4, ValueError), ("row", True, TypeError), ("row", 1.2, TypeError)]:
            with self.assertRaises(exc): table_grid_boundary(s, axis=axis, boundary_index=i)
        text = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        with self.assertRaisesRegex(ValueError, "native table"):
            table_grid_boundary(text, axis="row", boundary_index=0)

    def test_transformed_grouped_and_inconsistent_frames_are_explicit(self):
        _, slide, s = self.table()
        s.rotation = 10
        with self.assertRaisesRegex(ValueError, "rotated"):
            table_grid_boundary(s, axis="row", boundary_index=0)
        s.rotation = 0
        s._element.xfrm.set("flipH", "1")
        with self.assertRaisesRegex(ValueError, "flipped"):
            table_grid_boundary(s, axis="row", boundary_index=0)
        s._element.xfrm.set("flipH", "0")
        s.width += 1
        with self.assertRaisesRegex(ValueError, "frame and grid"):
            table_grid_boundary(s, axis="row", boundary_index=0)
        _, slide, s = self.table()
        group = slide.shapes.add_group_shape()
        group.shapes._spTree.append(s._element)
        nested = group.shapes[0]
        self.assertTrue(nested.has_table)
        with self.assertRaisesRegex(ValueError, "top-level"):
            table_grid_boundary(nested, axis="row", boundary_index=0)
        _, _, s = self.table()
        s.table.columns[0].width = 0
        with self.assertRaisesRegex(ValueError, "positive"):
            table_grid_boundary(s, axis="column", boundary_index=0)


if __name__ == "__main__": unittest.main()
